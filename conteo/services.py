from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile
import re
import unicodedata

from django.db import IntegrityError, transaction
from django.db.models import Max
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException

from .models import CargaDiferencias, Diferencia
from archivos.models import Linea


REQUIRED_HEADERS = (
    "linea",
    "sublinea",
    "codigo",
    "ultimo_conteo",
    "existencia",
    "diferencia",
    "ubicacion_1",
    "ubicacion_2",
    "ubicacion_3",
    "ubicacion_4",
    "fecha_generacion",
    "hora_generacion",
)

OPTIONAL_HEADERS = ("item", "descripcion")

HEADER_ALIASES = {
    # Variantes encontradas en los reportes de diferencias.
    "fecha_geracion": "fecha_generacion",
    "fecha_de_geracion": "fecha_generacion",
    "fecha_de_generacion": "fecha_generacion",
}


class DifferenceImportError(Exception):
    pass


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    # Algunos generadores guardan los saltos de linea de Excel como texto
    # escapado (_x000A_ o _x000D_) en vez de devolver \n / \r.
    text = re.sub(r"_x(?:000a|000d)_", " ", text, flags=re.IGNORECASE)
    # Excel puede traer espacios no separables, saltos de linea o signos invisibles
    # en los encabezados. Se convierten todos los separadores a un solo guion bajo.
    return re.sub(r"[^a-z0-9]+", "_", text.casefold()).strip("_")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal(value, row_number, column):
    if value is None or str(value).strip() == "":
        raise DifferenceImportError(f"Fila {row_number}: {column} no puede estar vacío.")
    if isinstance(value, Decimal):
        result = value
    else:
        cleaned = str(value).strip().replace(" ", "")
        if "," in cleaned and "." in cleaned:
            cleaned = cleaned.replace(",", "")
        elif "," in cleaned:
            cleaned = cleaned.replace(",", ".")
        try:
            result = Decimal(cleaned)
        except InvalidOperation as error:
            raise DifferenceImportError(
                f"Fila {row_number}: {column} debe contener un número válido."
            ) from error
    if abs(result) >= Decimal("100000000000000"):
        raise DifferenceImportError(f"Fila {row_number}: {column} excede el límite permitido.")
    return result


def _date(value, row_number):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            pass
    text = str(value or "").strip()
    for date_format in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    raise DifferenceImportError(
        f"Fila {row_number}: Fecha Geracion debe ser una fecha válida."
    )


def _time(value, row_number):
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.time().replace(microsecond=0)
            if isinstance(converted, time):
                return converted.replace(microsecond=0)
        except (TypeError, ValueError, OverflowError):
            pass
    text = str(value or "").strip()
    for time_format in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return datetime.strptime(text, time_format).time()
        except ValueError:
            continue
    raise DifferenceImportError(
        f"Fila {row_number}: Hora Generación debe contener una hora válida."
    )


def _same_line(value, selected_line):
    reported = _text(value)
    accepted = {str(selected_line.pk).strip(), str(selected_line.nlinea).strip()}
    if reported in accepted:
        return True
    if reported.isdigit():
        return any(code.isdigit() and int(reported) == int(code) for code in accepted)
    return False


def import_differences(uploaded_file, selected_line, user, as400_pdf=None):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise DifferenceImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        sheet = workbook.active
        # Los reportes pueden colocar los encabezados en la fila 1 o en la 2.
        # Se conserva la candidata que reconozca más columnas obligatorias.
        positions = {}
        header_row_number = 1
        for candidate_row_number, header_row in enumerate(
            sheet.iter_rows(min_row=1, max_row=2, values_only=True),
            start=1,
        ):
            candidate_positions = {}
            for index, value in enumerate(header_row):
                header = _normalize_header(value)
                if header:
                    candidate_positions[HEADER_ALIASES.get(header, header)] = index
            candidate_matches = sum(
                header in candidate_positions for header in REQUIRED_HEADERS
            )
            current_matches = sum(header in positions for header in REQUIRED_HEADERS)
            if candidate_matches > current_matches:
                positions = candidate_positions
                header_row_number = candidate_row_number
            if candidate_matches == len(REQUIRED_HEADERS):
                positions = candidate_positions
                header_row_number = candidate_row_number
                break
        missing = [header for header in REQUIRED_HEADERS if header not in positions]
        if missing:
            readable = ", ".join(header.replace("_", " ") for header in missing)
            raise DifferenceImportError(f"Faltan las columnas obligatorias: {readable}.")

        parsed_rows = []
        first_data_row = header_row_number + 1
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=first_data_row, values_only=True),
            start=first_data_row,
        ):
            values = {
                header: row[positions[header]] if positions[header] < len(row) else None
                for header in REQUIRED_HEADERS
            }
            values.update(
                {
                    header: row[positions[header]]
                    if header in positions and positions[header] < len(row)
                    else None
                    for header in OPTIONAL_HEADERS
                }
            )
            if not any(value not in (None, "") for value in values.values()):
                continue
            if len(parsed_rows) >= 50000:
                raise DifferenceImportError("El archivo supera el máximo de 50,000 registros.")
            if not _same_line(values["linea"], selected_line):
                raise DifferenceImportError(
                    f"Fila {row_number}: la línea {values['linea']} no corresponde a la línea "
                    f"{selected_line.nlinea} seleccionada."
                )

            sublinea = _text(values["sublinea"])
            item = _text(values["item"])
            codigo = _text(values["codigo"])
            descripcion = _text(values["descripcion"])
            if not sublinea or not codigo:
                raise DifferenceImportError(
                    f"Fila {row_number}: sublinea y codigo son obligatorios."
                )
            if len(sublinea) > 50 or len(codigo) > 100:
                raise DifferenceImportError(
                    f"Fila {row_number}: sublinea o codigo supera la longitud permitida."
                )
            if len(item) > 100 or len(descripcion) > 255:
                raise DifferenceImportError(
                    f"Fila {row_number}: item o descripción supera la longitud permitida."
                )

            locations = {}
            for index in range(1, 5):
                field = f"ubicacion_{index}"
                locations[field] = _text(values[field])
                if len(locations[field]) > 100:
                    raise DifferenceImportError(
                        f"Fila {row_number}: Ubicación {index} supera 100 caracteres."
                    )

            ultimo_conteo = _decimal(
                values["ultimo_conteo"], row_number, "Último Conteo"
            )
            existencia = _decimal(
                values["existencia"], row_number, "Existencia"
            )
            # Se valida la columna recibida para conservar el contrato del
            # archivo, pero el signo confiable siempre se calcula internamente.
            _decimal(values["diferencia"], row_number, "Diferencia")
            parsed_rows.append(
                {
                    "sublinea": sublinea,
                    "item": item,
                    "codigo": codigo,
                    "descripcion": descripcion,
                    "ultimo_conteo": ultimo_conteo,
                    "existencia": existencia,
                    "diferencia": ultimo_conteo - existencia,
                    **locations,
                    "fecha_generacion": _date(values["fecha_generacion"], row_number),
                    "hora_generacion": _time(values["hora_generacion"], row_number),
                }
            )

        if not parsed_rows:
            raise DifferenceImportError("El archivo no contiene registros para importar.")
        if not any(record["diferencia"] != 0 for record in parsed_rows):
            raise DifferenceImportError(
                "El archivo no contiene códigos con una diferencia distinta de cero."
            )

        try:
            with transaction.atomic():
                locked_line = Linea.objects.select_for_update().get(pk=selected_line.pk)
                latest_load = CargaDiferencias.objects.filter(linea=locked_line).first()
                if latest_load and latest_load.estado == CargaDiferencias.Estado.FINALIZADO:
                    raise DifferenceImportError(
                        f"La línea {locked_line.nlinea} está finalizada y no admite nuevas cargas."
                    )
                pending = CargaDiferencias.objects.filter(
                    linea=locked_line,
                    estado=CargaDiferencias.Estado.PENDIENTE,
                ).first()
                if pending:
                    raise DifferenceImportError(
                        f"La línea {locked_line.nlinea} ya tiene una carga pendiente de revisión. "
                        "El Gerente debe seleccionar los códigos antes de cargar otro archivo."
                    )
                last_number = (
                    CargaDiferencias.objects.filter(linea=locked_line)
                    .aggregate(max_number=Max("numero_reconteo"))["max_number"]
                    or 0
                )
                uploaded_file.seek(0)
                load = CargaDiferencias.objects.create(
                    linea=locked_line,
                    archivo_nombre=uploaded_file.name[:255],
                    archivo_original=uploaded_file,
                    archivo_as400=as400_pdf or "",
                    numero_reconteo=last_number + 1,
                    estado=CargaDiferencias.Estado.PENDIENTE,
                    total_registros=len(parsed_rows),
                    cargado_por=user,
                )
                Diferencia.objects.bulk_create(
                    [Diferencia(carga=load, **record) for record in parsed_rows],
                    batch_size=1000,
                )
        except IntegrityError as error:
            raise DifferenceImportError(
                "La línea recibió otra carga simultánea. Actualiza la página antes de continuar."
            ) from error
        return load
    finally:
        workbook.close()
