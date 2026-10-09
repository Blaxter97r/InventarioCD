from datetime import date, time
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import shutil
import tempfile
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from archivos.models import (
    CargaBaseCosto,
    CargaConteoBE,
    CargaConteoInventario,
    CargaRGE,
    CostoProducto,
    DetalleConteoBE,
    DetalleConteoInventario,
    DetalleRGE,
    Linea,
    Sublinea,
    Ubicacion,
)

from .models import (
    CargaDiferencias,
    Diferencia,
    InicioReconteoBodega,
    ValidacionUbicacion,
)
from .progress_reports import duplicate_scan_rows
from .recount_reports import (
    EXTRA_LOCATION_ROWS,
    build_recount_pdf,
    _combined_inventory_table,
    _filter_pdf_locations,
    _paginated_inventory_rows,
    _product_rows,
    _report_styles,
)
from .recount_analysis import build_recount_analysis


TEST_MEDIA_ROOT = tempfile.mkdtemp(prefix="inventario-conteo-tests-")


HEADERS = (
    "linea",
    "sublinea",
    "codigo",
    "ultimo conteo",
    "existencia",
    "Diferencia",
    "Ubicación 1",
    "Ubicación 2",
    "Ubicación 3",
    "Ubicación 4",
    "Fecha Geracion",
    "Hora Generación",
)


def make_difference_workbook(rows, headers=HEADERS, rows_before_headers=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Diferencias"
    for leading_row in rows_before_headers:
        sheet.append(leading_row)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "diferencias_linea.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def make_as400_pdf():
    output = BytesIO()
    document = canvas.Canvas(output)
    document.drawString(72, 760, "AS400 DIFERENCIAS PRUEBA")
    document.showPage()
    document.save()
    return SimpleUploadedFile(
        "as400_diferencias.pdf", output.getvalue(), content_type="application/pdf"
    )


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class DifferenceAccessTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser(
            username="noel.rivera", email="noel@empresa.com", password="admin-2026"
        )
        self.user = User.objects.create_user(username="supervisor", password="supervisor-2026")
        self.user.groups.add(Group.objects.get(name="Supervisores"))
        self.manager = User.objects.create_user(username="gerente", password="gerente-2026")
        self.manager.groups.add(Group.objects.get(name="Gerente"))
        self.line = Linea.objects.create(id_linea="L01", nlinea="1", descripcion="Filtros")

    def test_authenticated_user_can_view_differences(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("conteo:differences_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Diferencias por línea")
        self.assertContains(response, reverse("conteo:differences_upload"))
        self.assertNotContains(response, "Permiso de administrador")

    def test_authenticated_user_can_view_progress(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("conteo:progress"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Avance del conteo")
        self.assertContains(response, "Conteo local (desarrollo)")

    def test_progress_filters_preserve_scroll_position(self):
        Ubicacion.objects.create(
            ubicacion="A01I-B01", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress"))

        self.assertContains(response, "data-progress-floor-link")
        self.assertContains(response, "conteo/js/progress_scroll.js")

    def test_progress_explorer_builds_pending_location_hierarchy(self):
        Ubicacion.objects.create(
            ubicacion="A01I-B01", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="A01I-B02", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="A01I-B03", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        load = CargaConteoInventario.objects.create(
            archivo_nombre="snapshot.xlsx", cargado_por=self.admin, total_registros=1
        )
        DetalleConteoInventario.objects.create(
            carga=load, bodega="CENTRAL", ubicacion="A01I-B01", codigo="ABC", cantidad=1
        )
        DetalleConteoInventario.objects.create(
            carga=load, bodega="CENTRAL", ubicacion="A01I-B03", codigo="ABC", cantidad=1
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx",
            total_registros=2,
            total_codigos=2,
            total_ubicaciones=2,
            cargado_por=self.admin,
        )
        DetalleRGE.objects.create(
            carga=rge_load,
            codigo="CONTADO",
            ubicacion="A01I-B01",
        )
        DetalleRGE.objects.create(
            carga=rge_load,
            codigo="ABC",
            ubicacion="A01I-B02",
        )
        DetalleRGE.objects.create(
            carga=rge_load,
            codigo="ABC",
            ubicacion="A01I-B03",
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress"), {"piso": "A"})

        body = response.context["progress_tree"][0]["sides"][0]["bodies"][0]
        pending_location = body["pending_locations"][0]
        locations = {item["code"]: item for item in body["locations"]}
        self.assertEqual(body["name"], "B")
        self.assertEqual(body["pending"], 1)
        self.assertEqual(len(body["locations"]), 3)
        self.assertTrue(locations["A01I-B01"]["counted"])
        self.assertEqual(locations["A01I-B01"]["status"], "counted")
        self.assertFalse(locations["A01I-B02"]["counted"])
        self.assertEqual(pending_location["code"], "A01I-B02")
        self.assertEqual(pending_location["level"], "02")
        self.assertEqual(pending_location["status"], "possible_location_error")
        self.assertEqual(pending_location["expected_products"], 1)
        self.assertEqual(pending_location["found_elsewhere"], 1)
        self.assertEqual(
            pending_location["products"][0]["counted_locations"],
            ["A01I-B01", "A01I-B03"],
        )
        self.assertEqual(
            pending_location["products"][0]["unexpected_counted_locations"],
            ["A01I-B01"],
        )
        self.assertEqual(pending_location["products"][0]["expected_location"], "A01I-B02")
        self.assertContains(response, "data-progress-explorer")
        self.assertContains(response, "data-reconciliation-dialog")
        self.assertContains(response, "conteo/js/progress_explorer.js")
        self.assertContains(response, reverse("conteo:pending_locations_download", args=["A"]))

        download = self.client.get(reverse("conteo:pending_locations_download", args=["A"]))
        self.assertEqual(download.status_code, 200)
        self.assertIn("spreadsheetml", download["Content-Type"])
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(
            rows,
            [
                ("Piso", "Pasillo", "Lado", "Estante", "Nivel", "Ubicación"),
                ("A", "01", "I", "B", "02", "A01I-B02"),
            ],
        )

    def test_reconciliation_metrics_distinguish_elsewhere_from_not_located(self):
        for location in ("A01I-B01", "A01I-B02"):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="1",
                tipo="PRIMARIO",
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-metricas.xlsx",
            cargado_por=self.admin,
        )
        for code in ("CONTADO-AQUI", "EN-OTRO-LUGAR", "NO-LOCALIZADO"):
            DetalleRGE.objects.create(
                carga=rge_load,
                codigo=code,
                ubicacion="A01I-B01",
            )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo-metricas.xlsx",
            cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="A01I-B01",
            codigo="CONTADO-AQUI",
            cantidad=1,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="A01I-B02",
            codigo="EN-OTRO-LUGAR",
            cantidad=1,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress"), {"piso": "A"})

        location = response.context["progress_tree"][0]["sides"][0]["bodies"][0]["locations"][0]
        self.assertEqual(location["expected_products"], 3)
        self.assertEqual(location["found_elsewhere"], 1)
        self.assertEqual(location["not_located_products"], 1)
        self.assertContains(response, "Productos esperados en esta ubicación")
        self.assertContains(response, "Encontrados en otras ubicaciones")
        self.assertContains(response, "No localizados en el conteo")

    def test_validated_location_counts_toward_progress_and_can_be_unchecked(self):
        for location in ("A01I-B01", "A01I-B02"):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="1",
                tipo="PRIMARIO",
            )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="snapshot.xlsx", cargado_por=self.admin, total_registros=1
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="A01I-B01",
            codigo="CONTADO",
            cantidad=1,
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx",
            total_registros=2,
            total_codigos=2,
            total_ubicaciones=2,
            cargado_por=self.admin,
        )
        for code, location in (("CONTADO", "A01I-B01"), ("PENDIENTE", "A01I-B02")):
            DetalleRGE.objects.create(carga=rge_load, codigo=code, ubicacion=location)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("conteo:validate_progress_location"),
            {"ubicacion": "a01i-b02", "validada": "true"},
        )

        self.assertEqual(response.status_code, 200)
        validation = ValidacionUbicacion.objects.get()
        self.assertEqual(validation.ubicacion, "A01I-B02")
        self.assertEqual(validation.validada_por, self.user)

        progress = self.client.get(reverse("conteo:progress"), {"piso": "A"})
        self.assertEqual(progress.context["counted_total"], 1)
        self.assertEqual(progress.context["validated_total"], 1)
        self.assertEqual(progress.context["completed_total"], 2)
        self.assertEqual(progress.context["pending_total"], 0)
        self.assertEqual(progress.context["percentage"], 100)
        location = progress.context["progress_tree"][0]["sides"][0]["bodies"][0]["locations"][1]
        self.assertTrue(location["validated"])
        self.assertEqual(location["status"], "validated")
        self.assertContains(progress, "data-validation-checkbox")

        download = self.client.get(reverse("conteo:pending_locations_download", args=["A"]))
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(rows), 1)

        response = self.client.post(
            reverse("conteo:validate_progress_location"),
            {"ubicacion": "A01I-B02", "validada": "false"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(ValidacionUbicacion.objects.exists())
        progress = self.client.get(reverse("conteo:progress"), {"piso": "A"})
        self.assertEqual(progress.context["completed_total"], 1)
        self.assertEqual(progress.context["pending_total"], 1)

    def test_counted_or_unknown_location_cannot_be_validated(self):
        Ubicacion.objects.create(
            ubicacion="A01I-B01", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="snapshot.xlsx", cargado_por=self.admin, total_registros=1
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="A01I-B01", codigo="ABC", cantidad=1
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin, total_ubicaciones=1
        )
        DetalleRGE.objects.create(carga=rge_load, codigo="ABC", ubicacion="A01I-B01")
        self.client.force_login(self.user)

        counted = self.client.post(
            reverse("conteo:validate_progress_location"),
            {"ubicacion": "A01I-B01", "validada": "true"},
        )
        unknown = self.client.post(
            reverse("conteo:validate_progress_location"),
            {"ubicacion": "A01I-B99", "validada": "true"},
        )

        self.assertEqual(counted.status_code, 409)
        self.assertEqual(unknown.status_code, 404)
        self.assertFalse(ValidacionUbicacion.objects.exists())

    def test_secondary_aisle_pdf_reports_only_same_location_quantity_differences(self):
        for location in ("E08-045Y", "E08-046Y"):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="8",
                tipo="SECUNDARIO",
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin, total_ubicaciones=2
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="DIF-001", ubicacion="E08-045Y", cantidad=10
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="IGUAL-001", ubicacion="E08-046Y", cantidad=4
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="E08-045Y",
            codigo="DIF-001", cantidad=7,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="E08-046Y",
            codigo="IGUAL-001", cantidad=4,
        )
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("conteo:secondary_quantity_differences_pdf"),
            {"bodega": "BODEGA CENTRAL", "piso": "E", "pasillo": "08"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        document = PdfReader(BytesIO(response.content))
        text = "\n".join(page.extract_text() or "" for page in document.pages)
        self.assertIn("DIF-001", text)
        self.assertIn("E08-045Y", text)
        self.assertIn("Cantidad RGE", text)
        self.assertIn("Firma auditor", text)
        self.assertNotIn("IGUAL-001", text)

        page = self.client.get(reverse("conteo:progress"), {
            "bodega": "BODEGA CENTRAL", "tipo": "SECUNDARIO", "piso": "E",
        })
        self.assertNotContains(page, "data-secondary-pdf-download")
        self.assertContains(page, reverse("conteo:rge_modifications_download"))

        differences_page = self.client.get(reverse("conteo:progress_differences"), {
            "bodega": "BODEGA CENTRAL", "tipo": "SECUNDARIO", "piso": "E",
        })
        self.assertContains(differences_page, "data-secondary-pdf-download")
        self.assertContains(
            differences_page, reverse("conteo:secondary_quantity_differences_pdf")
        )

    def test_difference_report_modal_exposes_dynamic_filters(self):
        for location, location_type in (
            ("A01I-A01", "PRIMARIO"),
            ("E08-045Y", "SECUNDARIO"),
        ):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="1",
                tipo=location_type,
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="PRI-001", ubicacion="A01I-A01", cantidad=1
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="SEC-001", ubicacion="E08-045Y", cantidad=1
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress_differences"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "data-difference-report-open")
        self.assertContains(response, "data-difference-report-modal")
        self.assertContains(response, "Posibles escaneos duplicados")
        self.assertContains(response, reverse("conteo:filtered_differences_pdf"))
        self.assertContains(response, "difference_report_modal.js")
        central = response.context["difference_report_filters"][0]
        self.assertEqual(central["value"], "BODEGA CENTRAL")
        self.assertEqual(
            [item["value"] for item in central["types"]],
            ["PRIMARIO", "SECUNDARIO"],
        )
        self.assertEqual(central["types"][0]["scopes"][0]["aisles"][0], {
            "value": "01", "location_count": 1,
        })

    def test_filtered_primary_pdf_lists_missing_codes_for_selected_aisle(self):
        for location in ("A01I-A01", "A02I-A01"):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="1",
                tipo="PRIMARIO",
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-primario.xlsx", cargado_por=self.admin
        )
        for code, location in (
            ("CONTADO-001", "A01I-A01"),
            ("KIBCF0011", "A01I-A01"),
            ("OTRO-PASILLO", "A02I-A01"),
        ):
            DetalleRGE.objects.create(
                carga=rge_load, codigo=code, ubicacion=location, cantidad=1
            )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="A01I-A01",
            codigo="CONTADO-001",
            cantidad=1,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:filtered_differences_pdf"), {
            "bodega": "BODEGA CENTRAL",
            "tipo": "PRIMARIO",
            "piso": "A",
            "pasillo": ["01"],
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("REPORTE DE VALIDACI", text)
        self.assertGreaterEqual(text.count("contados"), 2)
        self.assertIn("A01I-A01", text)
        self.assertIn("1/2", text)
        self.assertIn("KIBCF0011", text)
        self.assertNotIn("OTRO-PASILLO", text)

    def test_filtered_secondary_pdf_supports_multiple_aisles_and_omits_zero_count(self):
        for location in ("E08-045Y", "E09-045Y"):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="8",
                tipo="SECUNDARIO",
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-secundario.xlsx", cargado_por=self.admin
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="SIN-CONTAR", ubicacion="E08-045Y", cantidad=1
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="DIF-009", ubicacion="E09-045Y", cantidad=5
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="E08-045Y",
            codigo="SIN-CONTAR",
            cantidad=0,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="E09-045Y",
            codigo="DIF-009",
            cantidad=3,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:filtered_differences_pdf"), {
            "bodega": "BODEGA CENTRAL",
            "tipo": "SECUNDARIO",
            "piso": "E",
            "pasillo": ["08", "09"],
        })

        self.assertEqual(response.status_code, 200)
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("Cantidad RGE", text)
        self.assertIn("Cantidad contada", text)
        self.assertIn("Firma auditor", text)
        self.assertNotIn("SIN-CONTAR", text)
        self.assertNotIn("E08-045Y", text)
        self.assertIn("DIF-009", text)
        self.assertIn("E09-045Y", text)
        self.assertIn("-2", text)

        invalid = self.client.get(reverse("conteo:filtered_differences_pdf"), {
            "bodega": "BODEGA CENTRAL",
            "tipo": "SECUNDARIO",
            "piso": "E",
            "pasillo": ["99"],
        })
        self.assertEqual(invalid.status_code, 404)

    def test_duplicate_scan_pdf_groups_exact_repeated_scans(self):
        for location in ("A01I-A01", "A02I-A01"):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="1",
                tipo="PRIMARIO",
            )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo-duplicados.xlsx", cargado_por=self.admin
        )
        for code, location, quantity in (
            ("DUP-001", "A01I-A01", 2),
            ("DUP-001", "A01I-A01", 2),
            ("DUP-001", "A01I-A01", 3),
            ("UNICO-001", "A01I-A01", 2),
            ("OTRO-PASILLO", "A02I-A01", 4),
            ("OTRO-PASILLO", "A02I-A01", 4),
        ):
            DetalleConteoInventario.objects.create(
                carga=count_load,
                bodega="CENTRAL",
                ubicacion=location,
                codigo=code,
                cantidad=quantity,
            )
        self.client.force_login(self.user)

        rows = duplicate_scan_rows("BODEGA CENTRAL", "PRIMARIO", "A", ["01"])
        self.assertEqual(rows, [{
            "location": "A01I-A01",
            "code": "DUP-001",
            "counted_quantity": Decimal("2"),
            "scan_count": 2,
            "duplicate_count": 1,
        }])

        response = self.client.get(reverse("conteo:filtered_differences_pdf"), {
            "reporte": "duplicados",
            "bodega": "BODEGA CENTRAL",
            "tipo": "PRIMARIO",
            "piso": "A",
            "pasillo": ["01"],
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn("posibles_duplicados", response["Content-Disposition"])
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("Posibles escaneos duplicados", text)
        self.assertIn("Veces escaneado", text)
        self.assertIn("DUP-001", text)
        self.assertIn("A01I-A01", text)
        self.assertNotIn("UNICO-001", text)
        self.assertNotIn("OTRO-PASILLO", text)

    def test_secondary_modify_rge_flag_exports_rge_codes_at_zero(self):
        source = "E08-045Y"
        actual = "E08-046Y"
        for location in (source, actual):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector="8",
                tipo="SECUNDARIO",
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin, total_ubicaciones=1
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="MOVER-001", ubicacion=source, cantidad=6
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="MOVER-002", ubicacion=source, cantidad=3
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion=actual,
            codigo="MOVER-001", cantidad=6,
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("conteo:validate_progress_location"),
            {
                "ubicacion": source,
                "campo": "modificar_rge",
                "modificar_rge": "true",
            },
        )

        self.assertEqual(response.status_code, 200)
        validation = ValidacionUbicacion.objects.get()
        self.assertFalse(validation.validada)
        self.assertTrue(validation.modificar_rge)
        progress = self.client.get(reverse("conteo:progress"), {
            "bodega": "BODEGA CENTRAL", "tipo": "SECUNDARIO", "piso": "E",
        })
        self.assertEqual(progress.context["validated_total"], 0)
        source_item = next(
            item
            for aisle in progress.context["progress_tree"]
            for shelf in aisle["sides"]
            for level in shelf["bodies"]
            for item in level["locations"]
            if item["code"] == source
        )
        self.assertTrue(source_item["modify_rge"])

        download = self.client.get(reverse("conteo:rge_modifications_download"))
        self.assertEqual(download.status_code, 200)
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(rows, [
            ("Codigo", "Ubicacion", "Cantidad contada"),
            ("MOVER-001", source, 0),
            ("MOVER-002", source, 0),
        ])

        response = self.client.post(
            reverse("conteo:validate_progress_location"),
            {
                "ubicacion": source,
                "campo": "validada",
                "validada": "true",
            },
        )
        self.assertEqual(response.status_code, 200)
        validation.refresh_from_db()
        self.assertTrue(validation.validada)
        self.assertTrue(validation.modificar_rge)

    def test_modify_rge_rejects_primary_and_accepts_unlocated_secondary_product(self):
        Ubicacion.objects.create(
            ubicacion="A01I-B01", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="E08-045Y", bodega="BODEGA CENTRAL", sector="8", tipo="SECUNDARIO"
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="ABC", ubicacion="A01I-B01", cantidad=1
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="NO-LOCALIZADO", ubicacion="E08-045Y", cantidad=1
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("conteo:validate_progress_location"),
            {
                "ubicacion": "A01I-B01",
                "campo": "modificar_rge",
                "modificar_rge": "true",
            },
        )

        self.assertEqual(response.status_code, 409)
        self.assertFalse(ValidacionUbicacion.objects.exists())

        response = self.client.post(
            reverse("conteo:validate_progress_location"),
            {
                "ubicacion": "E08-045Y",
                "campo": "modificar_rge",
                "modificar_rge": "true",
            },
        )

        self.assertEqual(response.status_code, 200)
        validation = ValidacionUbicacion.objects.get(ubicacion="E08-045Y")
        self.assertFalse(validation.validada)
        self.assertTrue(validation.modificar_rge)

        download = self.client.get(reverse("conteo:rge_modifications_download"))
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(rows, [
            ("Codigo", "Ubicacion", "Cantidad contada"),
            ("NO-LOCALIZADO", "E08-045Y", 0),
        ])

    def test_progress_only_includes_central_primary_mezzanine_locations(self):
        Ubicacion.objects.create(
            ubicacion="A01I-B01", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="A01I-B02", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="R-05-01-03", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="A01I-B03", bodega="RANSA", sector="1", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="A01I-B04", bodega="BODEGA CENTRAL", sector="1", tipo="SECUNDARIO"
        )
        Ubicacion.objects.create(
            ubicacion="A01I-B05", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-todas-bodegas.xlsx",
            total_registros=5,
            total_codigos=5,
            total_ubicaciones=5,
            cargado_por=self.admin,
        )
        for code, location in (
            ("CENTRAL-1", "A01I-B01"),
            ("CENTRAL-2", "A01I-B02"),
            ("IRREGULAR", "R-05-01-03"),
            ("OTRA-BODEGA", "A01I-B03"),
            ("SECUNDARIA", "A01I-B04"),
        ):
            DetalleRGE.objects.create(carga=rge_load, codigo=code, ubicacion=location)
        load = CargaConteoInventario.objects.create(
            archivo_nombre="snapshot.xlsx", cargado_por=self.admin, total_registros=1
        )
        DetalleConteoInventario.objects.create(
            carga=load, bodega="CENTRAL", ubicacion="A01I-B01", codigo="ABC", cantidad=1
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("conteo:progress"), {"piso": "A"})
        self.assertContains(response, "50%")
        self.assertContains(response, "1 contadas · 1 pendientes")
        self.assertEqual(response.context["expected_total"], 2)
        expected_codes = {
            item["code"]
            for aisle in response.context["progress_tree"]
            for side in aisle["sides"]
            for body in side["bodies"]
            for item in body["locations"]
        }
        self.assertEqual(expected_codes, {"A01I-B01", "A01I-B02"})

    def test_progress_supports_central_secondary_with_side_at_end(self):
        Ubicacion.objects.create(
            ubicacion="E08-045Y", bodega="BODEGA CENTRAL", sector="8", tipo="SECUNDARIO"
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-secundario.xlsx",
            total_registros=1,
            total_codigos=1,
            total_ubicaciones=1,
            cargado_por=self.admin,
        )
        DetalleRGE.objects.create(carga=rge_load, codigo="SEC-1", ubicacion="E08-045Y")
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress"), {
            "bodega": "BODEGA CENTRAL", "tipo": "SECUNDARIO", "piso": "E",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["group_labels"], ["Pasillo", "Estante", "Nivel"])
        aisle = response.context["progress_tree"][0]
        shelf = aisle["sides"][0]
        level = shelf["bodies"][0]
        location = level["locations"][0]
        self.assertEqual((aisle["name"], shelf["name"], level["name"]), ("08", "04", "5"))
        self.assertEqual(location["leaf_value"], "Y")
        self.assertEqual(location["path"][-1], {"label": "Lado", "value": "Y"})

    def test_central_progress_accepts_requested_floors_with_both_location_formats(self):
        locations = (
            ("A01I-B01", "PRIMARIO"),
            ("E08-045Y", "PRIMARIO"),
            ("B09I-C01", "SECUNDARIO"),
            ("D11I-G01", "SECUNDARIO"),
            ("F12D-A02", "SECUNDARIO"),
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-pisos-central.xlsx",
            total_registros=len(locations),
            total_codigos=len(locations),
            total_ubicaciones=len(locations),
            cargado_por=self.admin,
        )
        for index, (location, location_type) in enumerate(locations, start=1):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector=location[0],
                tipo=location_type,
            )
            DetalleRGE.objects.create(
                carga=rge_load,
                codigo=f"PISO-{index}",
                ubicacion=location,
            )
        self.client.force_login(self.user)

        primary = self.client.get(reverse("conteo:progress"), {
            "bodega": "BODEGA CENTRAL", "tipo": "PRIMARIO",
        })
        secondary = self.client.get(reverse("conteo:progress"), {
            "bodega": "BODEGA CENTRAL", "tipo": "SECUNDARIO",
        })

        self.assertEqual(
            [scope["name"] for scope in primary.context["scopes"]],
            ["A", "E"],
        )
        self.assertEqual(
            [scope["name"] for scope in secondary.context["scopes"]],
            ["B", "D", "F"],
        )
        self.assertEqual(primary.context["invalid_total"], 0)
        self.assertEqual(secondary.context["invalid_total"], 0)

    def test_central_primary_floors_e_and_r_use_secondary_hierarchy(self):
        locations = (
            ("E02-063M", "E", "02", "06", "3", "M"),
            ("R03-074X", "R", "03", "07", "4", "X"),
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-primario-e-r.xlsx",
            total_registros=len(locations),
            total_codigos=len(locations),
            total_ubicaciones=len(locations),
            cargado_por=self.admin,
        )
        for index, (location, floor, *_dimensions) in enumerate(locations, start=1):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector=floor,
                tipo="PRIMARIO",
            )
            DetalleRGE.objects.create(
                carga=rge_load,
                codigo=f"PRIMARIO-{floor}-{index}",
                ubicacion=location,
            )
        self.client.force_login(self.user)

        for location, floor, aisle_name, shelf_name, level_name, side_name in locations:
            response = self.client.get(reverse("conteo:progress"), {
                "bodega": "BODEGA CENTRAL",
                "tipo": "PRIMARIO",
                "piso": floor,
            })

            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response.context["group_labels"],
                ["Pasillo", "Estante", "Nivel"],
            )
            self.assertEqual(response.context["leaf_label"], "Lado")
            aisle = response.context["progress_tree"][0]
            shelf = aisle["sides"][0]
            level = shelf["bodies"][0]
            location_row = level["locations"][0]
            self.assertEqual(
                (aisle["name"], shelf["name"], level["name"]),
                (aisle_name, shelf_name, level_name),
            )
            self.assertEqual(location_row["code"], location)
            self.assertEqual(location_row["leaf_value"], side_name)
            self.assertEqual(
                location_row["path"][-1],
                {"label": "Lado", "value": side_name},
            )

        download = self.client.get(
            reverse("conteo:pending_locations_download", args=["E"]),
            {"bodega": "BODEGA CENTRAL", "tipo": "PRIMARIO"},
        )
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(rows, [
            ("Piso", "Pasillo", "Estante", "Nivel", "Lado", "Ubicación"),
            ("E", "02", "06", "3", "M", "E02-063M"),
        ])

    def test_central_secondary_floors_b_d_and_f_use_primary_hierarchy(self):
        locations = (
            ("B09I-C01", "B", "09", "I", "C", "01"),
            ("D11D-G02", "D", "11", "D", "G", "02"),
            ("F12I-A03", "F", "12", "I", "A", "03"),
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-secundario-b-d-f.xlsx",
            total_registros=len(locations),
            total_codigos=len(locations),
            total_ubicaciones=len(locations),
            cargado_por=self.admin,
        )
        for index, (location, floor, *_dimensions) in enumerate(locations, start=1):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="BODEGA CENTRAL",
                sector=floor,
                tipo="SECUNDARIO",
            )
            DetalleRGE.objects.create(
                carga=rge_load,
                codigo=f"SECUNDARIO-{floor}-{index}",
                ubicacion=location,
            )
        self.client.force_login(self.user)

        for location, floor, aisle_name, side_name, shelf_name, level_name in locations:
            response = self.client.get(reverse("conteo:progress"), {
                "bodega": "BODEGA CENTRAL",
                "tipo": "SECUNDARIO",
                "piso": floor,
            })

            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response.context["group_labels"],
                ["Pasillo", "Lado", "Estante"],
            )
            self.assertEqual(response.context["leaf_label"], "Nivel")
            aisle = response.context["progress_tree"][0]
            side = aisle["sides"][0]
            shelf = side["bodies"][0]
            location_row = shelf["locations"][0]
            self.assertEqual(
                (aisle["name"], side["name"], shelf["name"]),
                (aisle_name, side_name, shelf_name),
            )
            self.assertEqual(location_row["code"], location)
            self.assertEqual(location_row["leaf_value"], level_name)
            self.assertEqual(
                location_row["path"][-1],
                {"label": "Nivel", "value": level_name},
            )

        download = self.client.get(
            reverse("conteo:pending_locations_download", args=["B"]),
            {"bodega": "BODEGA CENTRAL", "tipo": "SECUNDARIO"},
        )
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(rows, [
            ("Piso", "Pasillo", "Lado", "Estante", "Nivel", "Ubicación"),
            ("B", "09", "I", "C", "01", "B09I-C01"),
        ])

    def test_progress_supports_both_ransa_types_and_global_summary(self):
        for code, location_type, location in (
            ("R-P", "PRIMARIO", "R-08-18-03"),
            ("R-S", "SECUNDARIO", "R-05-19-04"),
        ):
            Ubicacion.objects.create(
                ubicacion=location, bodega="RANSA", sector="R", tipo=location_type
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-ransa.xlsx",
            total_registros=2,
            total_codigos=2,
            total_ubicaciones=2,
            cargado_por=self.admin,
        )
        DetalleRGE.objects.create(carga=rge_load, codigo="R-P", ubicacion="R-08-18-03")
        DetalleRGE.objects.create(carga=rge_load, codigo="R-S", ubicacion="R-05-19-04")
        load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo-ransa.xlsx", cargado_por=self.admin, total_registros=1
        )
        DetalleConteoInventario.objects.create(
            carga=load, bodega="RANSA", ubicacion="R-08-18-03", codigo="R-P", cantidad=1
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress"), {
            "bodega": "RANSA", "tipo": "SECUNDARIO", "piso": "R",
        })

        self.assertEqual(response.context["global_expected_total"], 2)
        self.assertEqual(response.context["global_counted_total"], 1)
        self.assertEqual(response.context["global_percentage"], 50)
        self.assertEqual(response.context["scope_label"], "Zona")
        self.assertEqual(response.context["download_headers"], (
            "Zona", "Pasillo", "Estante", "Nivel", "Ubicación",
        ))
        location = response.context["progress_tree"][0]["sides"][0]["bodies"][0]["locations"][0]
        self.assertEqual(location["code"], "R-05-19-04")
        self.assertEqual(location["leaf_value"], "")

    def test_general_progress_shows_warehouse_and_type_summary(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress_general"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "conteo/progress_general.html")
        self.assertContains(response, "Avance general")
        self.assertContains(response, "Avance detallado")
        self.assertContains(response, "Primario")
        self.assertContains(response, "Secundario")

    def test_general_progress_skips_product_reconciliation(self):
        for location in ("A01I-B01", "A01I-B02"):
            Ubicacion.objects.create(
                ubicacion=location, bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
            )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin,
        )
        for location in ("A01I-B01", "A01I-B02"):
            DetalleRGE.objects.create(carga=rge_load, codigo=location, ubicacion=location)
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="A01I-B01",
            codigo="A01I-B01", cantidad=1,
        )
        self.client.force_login(self.user)

        with patch("conteo.progress._build_reconciliation_data", side_effect=AssertionError):
            response = self.client.get(reverse("conteo:progress_general"))

        self.assertEqual(response.status_code, 200)
        primary = response.context["warehouse_cards"][0]["types"][0]
        self.assertEqual((primary["expected"], primary["counted"], primary["percentage"]), (2, 1, 50))

    def test_general_progress_includes_new_warehouse_from_map_and_rge(self):
        Ubicacion.objects.create(
            ubicacion="P-01", bodega="BODEGA PEDREGAL", sector="P", tipo="PRIMARIO"
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge-pedregal.xlsx", cargado_por=self.admin,
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="PED-001", ubicacion="P-01"
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo-pedregal.xlsx", cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="PEDREGAL", ubicacion="P-01",
            codigo="PED-001", cantidad=1,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("conteo:progress_general"))

        warehouses = {item["value"]: item for item in response.context["warehouse_cards"]}
        self.assertIn("BODEGA PEDREGAL", warehouses)
        pedregal = warehouses["BODEGA PEDREGAL"]
        self.assertEqual(
            (pedregal["expected"], pedregal["counted"], pedregal["percentage"]),
            (1, 1, 100),
        )
        self.assertEqual(pedregal["types"][0]["name"], "Primario")
        self.assertContains(response, "Bodega Pedregal")

    def test_progress_differences_only_counts_same_warehouse_and_location_type(self):
        for location in ("A01I-B01", "A01I-B02"):
            Ubicacion.objects.create(
                ubicacion=location, bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO"
            )
        Ubicacion.objects.create(
            ubicacion="A02I-B01", bodega="BODEGA CENTRAL", sector="2", tipo="PRIMARIO"
        )
        Ubicacion.objects.create(
            ubicacion="E08-045Y", bodega="BODEGA CENTRAL", sector="8", tipo="SECUNDARIO"
        )
        Ubicacion.objects.create(
            ubicacion="R-08-18-03", bodega="RANSA", sector="R", tipo="PRIMARIO"
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin,
        )
        for code, location in (
            ("CONTADO", "A01I-B01"),
            ("FALTA-UNO", "A01I-B01"),
            ("FALTA-DOS", "A01I-B02"),
        ):
            DetalleRGE.objects.create(carga=rge_load, codigo=code, ubicacion=location)
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin,
        )
        for warehouse, location, code in (
            ("CENTRAL", "A02I-B01", "CONTADO"),
            ("CENTRAL", "E08-045Y", "FALTA-DOS"),
            ("RANSA", "R-08-18-03", "FALTA-UNO"),
        ):
            DetalleConteoInventario.objects.create(
                carga=count_load, bodega=warehouse, ubicacion=location,
                codigo=code, cantidad=1,
            )
        self.client.force_login(self.user)

        selector = self.client.get(reverse("conteo:progress_differences"))
        self.assertEqual(selector.context["selected_scope"], "")
        self.assertNotContains(selector, "data-code-difference-explorer")
        self.assertContains(selector, "Selecciona un piso")

        response = self.client.get(
            reverse("conteo:progress_differences"), {"piso": "A"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "conteo/progress_differences.html")
        self.assertEqual(response.context["total_codes"], 3)
        self.assertEqual(response.context["missing_codes"], 2)
        self.assertEqual(response.context["selected_scope"], "A")
        aisle = response.context["scopes"][0]["groups"][0]
        shelf = aisle["children"][0]["children"][0]
        self.assertEqual(aisle["missing_codes"], 2)
        self.assertEqual(shelf["missing_codes"], 2)
        locations = {item["code"]: item for item in shelf["locations"]}
        self.assertEqual(locations["A01I-B01"]["missing_code_list"], ["FALTA-UNO"])
        self.assertEqual(locations["A01I-B02"]["missing_code_list"], ["FALTA-DOS"])
        self.assertContains(response, "2</strong> códigos sin contar")
        self.assertContains(response, "Diferencias")
        self.assertContains(response, "data-code-difference-explorer")
        self.assertContains(response, "conteo/js/progress_differences.js")

    def test_supervisor_can_open_upload_and_download_template(self):
        self.client.force_login(self.user)
        index = self.client.get(reverse("conteo:differences_index"))
        self.assertContains(index, reverse("conteo:differences_upload"))
        self.assertNotContains(index, "data-admin-permission-open")
        self.assertEqual(self.client.get(reverse("conteo:differences_upload")).status_code, 200)
        self.assertEqual(self.client.get(reverse("conteo:differences_template")).status_code, 200)

        workbook = make_difference_workbook([
            ("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30)),
        ])
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(CargaDiferencias.objects.get().cargado_por, self.user)

    def test_manager_cannot_open_upload(self):
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(reverse("conteo:differences_upload")).status_code, 403)
        self.assertEqual(self.client.get(reverse("conteo:differences_template")).status_code, 403)

    def test_admin_sees_upload_confirmation_markup(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("conteo:differences_upload"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "data-differences-confirm-modal")
        self.assertContains(response, "data-confirm-file")
        self.assertContains(response, "data-confirm-line")

    def test_manager_can_open_pending_queue_but_supervisor_cannot(self):
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(reverse("conteo:pending_queue")).status_code, 200)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("conteo:pending_queue")).status_code, 403)

    def test_supervisor_and_manager_can_view_count_differences_recounts_summary_and_progress(self):
        routes = (
            "archivos:inventory_count_index",
            "conteo:differences_index",
            "conteo:recounts",
            "conteo:line_summary",
            "conteo:recount_analysis",
            "conteo:progress",
            "conteo:progress_general",
        )
        for user in (self.user, self.manager):
            self.client.force_login(user)
            for route in routes:
                with self.subTest(role=user.username, route=route):
                    self.assertEqual(self.client.get(reverse(route)).status_code, 200)

    def test_only_manager_and_admin_can_open_code_selection_tab(self):
        self.client.force_login(self.user)
        supervisor_page = self.client.get(reverse("conteo:differences_index"))
        self.assertNotContains(supervisor_page, reverse("conteo:pending_queue"))
        self.assertEqual(self.client.get(reverse("conteo:pending_queue")).status_code, 403)
        for user in (self.manager, self.admin):
            self.client.force_login(user)
            page = self.client.get(reverse("conteo:pending_queue"))
            self.assertEqual(page.status_code, 200)
            self.assertContains(page, "Archivos por revisar")


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class DifferenceImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="noel.rivera", email="noel@empresa.com", password="admin-2026"
        )
        self.line = Linea.objects.create(id_linea="L01", nlinea="1", descripcion="Filtros")
        self.client.force_login(self.admin)

    def test_valid_file_creates_load_and_differences(self):
        workbook = make_difference_workbook(
            [
                ("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30)),
                ("1", "2", "ABC-002", 5.5, 7, 1.5, "A-02", "B-01", "", "", date(2026, 8, 8), time(14, 31)),
            ]
        )
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        self.assertEqual(response.status_code, 302)
        load = CargaDiferencias.objects.get()
        self.assertEqual(load.linea, self.line)
        self.assertEqual(load.total_registros, 2)
        self.assertEqual(load.numero_reconteo, 1)
        self.assertEqual(load.estado, CargaDiferencias.Estado.PENDIENTE)
        self.assertTrue(load.archivo_original.name)
        self.assertTrue(Path(load.archivo_original.path).exists())
        self.assertEqual(Diferencia.objects.count(), 2)
        first = Diferencia.objects.get(codigo="ABC-001")
        self.assertEqual(first.ubicacion_1, "A-01")
        self.assertEqual(first.diferencia, -2)
        page = self.client.get(reverse("conteo:differences_index"), {"linea": self.line.pk})
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Detalle de diferencias")
        self.assertNotContains(page, "Buscar producto")
        queue = self.client.get(reverse("conteo:pending_queue"))
        self.assertEqual(queue.status_code, 200)
        self.assertContains(queue, "diferencias_linea.xlsx")
        self.assertContains(queue, "#1")
        self.assertContains(queue, "data-review-sort")
        self.assertContains(queue, "data-review-min-total")
        self.assertContains(queue, "data-review-max-total")
        self.assertContains(queue, "Seleccionar visibles")
        cost_load = CargaBaseCosto.objects.create(
            archivo_nombre="costos.xlsx",
            total_registros=1,
            cargado_por=self.admin,
        )
        CostoProducto.objects.create(
            carga=cost_load,
            codigo="ABC-001",
            costo=Decimal("12.3456"),
        )
        codes_response = self.client.get(reverse("conteo:pending_codes", args=[load.pk]))
        self.assertEqual(codes_response.status_code, 200)
        codes = codes_response.json()["codes"]
        self.assertEqual(len(codes), 2)
        costs = {item["codigo"]: (item["costo"], item["costo_total"]) for item in codes}
        self.assertEqual(costs["ABC-001"], ("12.35", "-24.69"))
        self.assertEqual(costs["ABC-002"], (None, None))
        self.assertEqual(
            codes_response.json()["totales"],
            {"faltante": "-24.69", "sobrante": "0.00", "diferencia_neta": "-24.69"},
        )
        first_code = next(item for item in codes if item["codigo"] == "ABC-001")
        second_code = next(item for item in codes if item["codigo"] == "ABC-002")
        self.assertEqual(
            (first_code["ultimo_conteo"], first_code["existencia"], first_code["diferencia"]),
            ("10", "8", "-2"),
        )
        self.assertEqual(
            (second_code["ultimo_conteo"], second_code["existencia"], second_code["diferencia"]),
            ("6", "7", "2"),
        )
        CostoProducto.objects.create(
            carga=cost_load,
            codigo="ABC-002",
            costo=Decimal("3.333"),
        )
        totals = self.client.get(reverse("conteo:pending_codes", args=[load.pk])).json()["totales"]
        self.assertEqual(
            totals,
            {"faltante": "-24.69", "sobrante": "5.00", "diferencia_neta": "-19.69"},
        )

    def test_recount_analysis_applies_warehouse_matrix_and_audit_rule(self):
        ransa_line = Linea.objects.create(
            id_linea="L02", nlinea="2", descripcion="Línea mixta"
        )
        same_line = Linea.objects.create(
            id_linea="L03", nlinea="3", descripcion="Línea solo CD"
        )
        pedregal_line = Linea.objects.create(
            id_linea="L09", nlinea="9", descripcion="Línea Kensei"
        )
        first_load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="linea-1.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
            fecha_seleccion=timezone.now(),
        )
        mixed_load = CargaDiferencias.objects.create(
            linea=ransa_line,
            archivo_nombre="linea-2.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
            fecha_seleccion=timezone.now(),
        )
        same_load = CargaDiferencias.objects.create(
            linea=same_line,
            archivo_nombre="linea-3.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
            fecha_seleccion=timezone.now(),
        )
        pedregal_load = CargaDiferencias.objects.create(
            linea=pedregal_line,
            archivo_nombre="linea-9.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
            fecha_seleccion=timezone.now(),
        )
        for load, code in (
            (first_load, "SOLO-CD"),
            (first_load, "EXCEDENTE-RANSA"),
            (first_load, "EXCEDENTE-CD-RANSA"),
            (mixed_load, "PICKING-CD"),
            (mixed_load, "PICKING-RANSA"),
            (mixed_load, "PICKING-PEDREGAL"),
            (same_load, "SOLO-MISMA-BODEGA"),
            (pedregal_load, "PEDREGAL-CON-PRIMARIO"),
            (pedregal_load, "PEDREGAL-SIN-PRIMARIO"),
        ):
            Diferencia.objects.create(
                carga=load,
                sublinea="1",
                codigo=code,
                ultimo_conteo=2,
                existencia=1,
                diferencia=1,
                fecha_generacion=date(2026, 8, 17),
                hora_generacion=time(9, 30),
                seleccionada=True,
            )

        locations = (
            ("C-P", "CENTRAL", "Primario"),
            ("C-S", "CENTRAL", "Secundario"),
            ("R-P", "RANSA", "Primario"),
            ("R-S", "RANSA", "Secundario"),
            ("P-P", "PEDREGAL", "Primario"),
            ("P-S", "PEDREGAL", "Secundario"),
        )
        for location, warehouse, location_type in locations:
            Ubicacion.objects.create(
                ubicacion=location,
                bodega=warehouse,
                sector="Piso",
                tipo=location_type,
            )
        rge = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin
        )
        for code, location in (
            ("SOLO-CD", "C-P"),
            ("SOLO-CD", "C-S"),
            ("EXCEDENTE-RANSA", "C-P"),
            ("EXCEDENTE-RANSA", "R-S"),
            ("EXCEDENTE-CD-RANSA", "C-P"),
            ("EXCEDENTE-CD-RANSA", "C-S"),
            ("EXCEDENTE-CD-RANSA", "R-S"),
            ("PICKING-CD", "C-P"),
            ("PICKING-CD", "C-S"),
            ("PICKING-RANSA", "R-P"),
            ("PICKING-RANSA", "R-S"),
            ("PICKING-PEDREGAL", "P-P"),
            ("PICKING-PEDREGAL", "P-S"),
            ("SOLO-MISMA-BODEGA", "C-P"),
            ("SOLO-MISMA-BODEGA", "C-S"),
            ("PEDREGAL-CON-PRIMARIO", "P-P"),
            ("PEDREGAL-CON-PRIMARIO", "P-S"),
        ):
            DetalleRGE.objects.create(
                carga=rge, codigo=code, ubicacion=location, cantidad=1
            )

        analysis = build_recount_analysis()

        self.assertEqual(
            [item["line"].nlinea for item in analysis],
            ["1", "2", "3", "9"],
        )
        self.assertEqual(analysis[0]["primary_label"], "CD")
        self.assertEqual(analysis[0]["secondary_label"], "CD Y RANSA")
        self.assertEqual(analysis[0]["condition"], "surplus")
        self.assertEqual(analysis[0]["comment"], "Recibirás solo excedente de RANSA.")
        self.assertEqual(analysis[1]["audit_warehouses"], {"ransa", "pedregal"})
        self.assertEqual(analysis[1]["condition"], "audit")
        self.assertEqual(
            analysis[1]["comment"],
            "Esperar folleto de auditoría de RANSA y PEDREGAL.",
        )
        self.assertEqual(analysis[2]["condition"], "same_warehouse")
        self.assertEqual(
            analysis[2]["comment"],
            "No recibirás nada de otra bodega.",
        )
        self.assertEqual(analysis[3]["condition"], "external_line")
        self.assertEqual(analysis[3]["condition_label"], "Línea PEDREGAL")
        self.assertEqual(
            analysis[3]["comment"],
            "LÍNEA PEDREGAL. Ellos deben enviar el reconteo al CD.",
        )

        page = self.client.get(reverse("conteo:recount_analysis"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Análisis por línea")
        self.assertContains(page, "Conclusión de la línea")
        self.assertContains(page, "Esperar folleto")
        self.assertNotContains(page, "EXCEDENTE-RANSA")

    def test_recount_analysis_keeps_finished_recount_until_next_selection(self):
        finished_load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="reconteo-finalizado.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
            fecha_seleccion=timezone.now(),
            fecha_finalizacion=timezone.now(),
            finalizado_por=self.admin,
        )
        Diferencia.objects.create(
            carga=finished_load,
            sublinea="1",
            codigo="ANTERIOR-001",
            ultimo_conteo=2,
            existencia=1,
            diferencia=1,
            fecha_generacion=date(2026, 8, 17),
            hora_generacion=time(9, 30),
            seleccionada=True,
        )

        finished_analysis = build_recount_analysis()
        self.assertEqual([item["load"] for item in finished_analysis], [finished_load])
        self.assertContains(
            self.client.get(reverse("conteo:recount_analysis")),
            "Reconteo #1 · Finalizado",
        )

        pending_load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="nueva-carga.xlsx",
            numero_reconteo=2,
            estado=CargaDiferencias.Estado.PENDIENTE,
            cargado_por=self.admin,
        )
        new_difference = Diferencia.objects.create(
            carga=pending_load,
            sublinea="1",
            codigo="NUEVO-002",
            ultimo_conteo=3,
            existencia=1,
            diferencia=2,
            fecha_generacion=date(2026, 8, 18),
            hora_generacion=time(10, 0),
        )

        waiting_analysis = build_recount_analysis()
        self.assertEqual([item["load"] for item in waiting_analysis], [finished_load])

        new_difference.seleccionada = True
        new_difference.save(update_fields=("seleccionada",))
        pending_load.estado = CargaDiferencias.Estado.RECONTEO
        pending_load.seleccionado_por = self.admin
        pending_load.fecha_seleccion = timezone.now()
        pending_load.save(update_fields=("estado", "seleccionado_por", "fecha_seleccion"))

        refreshed_analysis = build_recount_analysis()
        self.assertEqual([item["load"] for item in refreshed_analysis], [pending_load])
        self.assertEqual(refreshed_analysis[0]["code_count"], 1)

    def test_as400_pdf_is_saved_and_appended_to_each_warehouse_report(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "C-01", "R-01", "", "", date(2026, 8, 8), time(14, 30))]
        )
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook, "archivo_as400": make_as400_pdf()},
        )
        self.assertEqual(response.status_code, 302)
        load = CargaDiferencias.objects.get()
        self.assertTrue(load.archivo_as400.name)
        self.assertTrue(Path(load.archivo_as400.path).exists())
        self.client.post(
            reverse("conteo:select_recount_codes", args=[load.pk]),
            {"select_all": "1"},
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx",
            total_registros=2,
            total_codigos=1,
            total_bodegas=2,
            cargado_por=self.admin,
        )
        for warehouse, location in (("CENTRAL", "C-01"), ("RANSA", "R-01")):
            DetalleConteoInventario.objects.create(
                carga=count_load,
                bodega=warehouse,
                ubicacion=location,
                codigo="ABC-001",
                cantidad=Decimal("1"),
            )

        page = self.client.get(reverse("conteo:recounts"))
        self.assertEqual(
            {option["key"]: option["count"] for option in page.context["recounts"][0].warehouse_options},
            {"central": 1, "ransa": 1, "pedregal": 0},
        )
        for warehouse in ("central", "ransa"):
            report = self.client.get(
                reverse("conteo:recount_warehouse_report", args=[load.pk, warehouse])
            )
            self.assertEqual(report.status_code, 200)
            pages = PdfReader(BytesIO(report.content)).pages
            self.assertGreaterEqual(len(pages), 2)
            self.assertIn("AS400 DIFERENCIAS PRUEBA", pages[-1].extract_text())

    def test_report_survives_lost_as400_file_with_clear_warning(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "C-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook, "archivo_as400": make_as400_pdf()},
        )
        load = CargaDiferencias.objects.get()
        self.client.post(reverse("conteo:select_recount_codes", args=[load.pk]), {"select_all": "1"})
        report_url = reverse("conteo:recount_warehouse_report", args=[load.pk, "central"])
        load.archivo_as400.storage.delete(load.archivo_as400.name)

        with self.assertLogs("conteo.recount_reports", level="ERROR"):
            legacy_report = self.client.get(report_url)
        self.assertEqual(legacy_report.status_code, 200)
        self.assertEqual(legacy_report["X-AS400-Status"], "unavailable")
        self.assertIn("SIN_AS400", legacy_report["Content-Disposition"])
        self.assertTrue(legacy_report.content.startswith(b"%PDF-"))
        self.assertNotIn("AS400 DIFERENCIAS PRUEBA", "\n".join(
            page.extract_text() for page in PdfReader(BytesIO(legacy_report.content)).pages
        ))

    def test_editor_quantity_and_location_changes_keep_recount_pdf_downloadable(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "C-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(reverse("conteo:differences_upload"), {"linea": self.line.pk, "archivo": workbook})
        load = CargaDiferencias.objects.get()
        self.client.post(reverse("conteo:select_recount_codes", args=[load.pk]), {"select_all": "1"})
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", total_registros=1, total_codigos=1,
            total_bodegas=1, cargado_por=self.admin,
        )
        detail = DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="C-01",
            codigo="ABC-001", cantidad=Decimal("5"),
        )
        editor = get_user_model().objects.create_user(username="editor-pdf", password="test-2026")
        editor.groups.add(Group.objects.get(name="Editores"))
        self.client.force_login(editor)
        self.assertEqual(self.client.post(reverse("archivos:inventory_count_bulk_quantity"), {
            "detalle_ids": [detail.pk],
            f"original_{detail.pk}": "5.0000",
            f"cantidad_{detail.pk}": "7",
        }).status_code, 302)
        self.assertEqual(self.client.post(reverse("archivos:inventory_count_location"), {
            "carga_id": count_load.pk, "codigo": "ABC-001", "bodega": "CENTRAL",
            "ubicacion": "NUEVA-01", "cantidad": "2",
        }).status_code, 302)

        self.client.force_login(self.admin)
        report_url = reverse("conteo:recount_warehouse_report", args=[load.pk, "central"])
        report = self.client.get(report_url)
        self.assertEqual(report.status_code, 200)
        text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(report.content)).pages)
        self.assertIn("C-01", text)
        self.assertIn("NUEVA-01", text)

        added = DetalleConteoInventario.objects.get(ubicacion="NUEVA-01")
        self.client.post(reverse("archivos:inventory_count_detail_delete", args=[added.pk]), {
            f"original_{added.pk}": "2.0000",
        })
        self.assertEqual(self.client.get(report_url).status_code, 200)

    def test_admin_can_restore_as400_for_existing_recount(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "C-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(reverse("conteo:differences_upload"), {"linea": self.line.pk, "archivo": workbook})
        load = CargaDiferencias.objects.get()
        self.client.post(reverse("conteo:select_recount_codes", args=[load.pk]), {"select_all": "1"})
        upload_url = reverse("conteo:replace_recount_as400", args=[load.pk])
        report_url = reverse("conteo:recount_warehouse_report", args=[load.pk, "central"])
        self.assertContains(self.client.get(reverse("conteo:recounts")), upload_url)

        supervisor = get_user_model().objects.create_user(username="supervisor-as400", password="test-2026")
        supervisor.groups.add(Group.objects.get(name="Supervisores"))
        self.client.force_login(supervisor)
        self.assertEqual(self.client.post(upload_url, {"archivo_as400": make_as400_pdf()}).status_code, 403)
        self.client.force_login(self.admin)
        invalid = SimpleUploadedFile("as400.pdf", b"invalido", content_type="application/pdf")
        self.assertRedirects(self.client.post(upload_url, {"archivo_as400": invalid}), reverse("conteo:recounts"))
        load.refresh_from_db()
        self.assertFalse(load.archivo_as400.name)

        self.assertRedirects(
            self.client.post(upload_url, {"archivo_as400": make_as400_pdf()}),
            reverse("conteo:recounts"),
        )
        load.refresh_from_db()
        self.assertTrue(Path(load.archivo_as400.path).exists())
        report = self.client.get(report_url)
        self.assertEqual(report.status_code, 200)
        self.assertIn("AS400 DIFERENCIAS PRUEBA", PdfReader(BytesIO(report.content)).pages[-1].extract_text())

    def test_invalid_as400_pdf_is_rejected_before_saving(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "C-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        invalid_pdf = SimpleUploadedFile(
            "as400.pdf", b"esto no es un PDF", content_type="application/pdf"
        )
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook, "archivo_as400": invalid_pdf},
        )
        self.assertContains(response, "El archivo AS400 no es un PDF válido.")
        self.assertFalse(CargaDiferencias.objects.exists())

    def test_pdf_groups_rge_counted_and_external_rows_by_selected_code(self):
        load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="diferencias.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
        )
        for code, selected in (("ABC-001", True), ("SIN-BE", True), ("NO-ELEGIDO", False)):
            Diferencia.objects.create(
                carga=load,
                sublinea="1",
                item="66",
                codigo=code,
                descripcion="Filtro de aceite",
                ultimo_conteo=0,
                existencia=1,
                diferencia=1,
                fecha_generacion=date(2026, 8, 17),
                hora_generacion=time(9, 30),
                seleccionada=selected,
            )
        be_load = CargaConteoBE.objects.create(archivo_nombre="be.xlsx", cargado_por=self.admin)
        for warehouse, code, quantity in (
            ("AGDOSA", "ABC-001", 3),
            ("Lourdes", "ABC-001", 4),
            ("Miramar", "ABC-001", 5),
            ("Miramar", "NO-ELEGIDO", 100),
        ):
            DetalleConteoBE.objects.create(
                carga=be_load, bodega=warehouse, codigo=code, cantidad=quantity
            )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin
        )
        for location, location_type in (
            ("C-01", "Primario"),
            ("C-02", "Secundario"),
            ("C-03", "Secundario"),
            ("C-04", "Secundario"),
        ):
            Ubicacion.objects.create(
                ubicacion=location,
                bodega="CENTRAL",
                sector="Piso",
                tipo=location_type,
            )
        for warehouse, location, quantity in (
            ("BODEGA CENTRAL", "C-01", 4),
            ("Central", "C-02", 2),
            ("RANSA", "R-01", 3),
            ("BODEGA PEDREGAL", "P-01", 7),
        ):
            DetalleConteoInventario.objects.create(
                carga=count_load, bodega=warehouse, ubicacion=location,
                codigo="abc-001", cantidad=quantity,
            )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="C-03",
            codigo="SIN-BE", cantidad=Decimal("2"),
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="ABC-001", ubicacion="C-01", cantidad=Decimal("10")
        )
        DetalleRGE.objects.create(
            carga=rge_load, codigo="ABC-001", ubicacion="C-04", cantidad=Decimal("8")
        )

        rows = _product_rows(load, "central")
        abc_row = next(row for row in rows if row["base"][3] == "ABC-001")
        self.assertEqual(
            abc_row["locations"],
            [
                {
                    "warehouse": "CENTRAL", "type": "PRIMARIO", "location": "C-01",
                    "rge_quantity": Decimal("10"), "counted_quantity": Decimal("4"),
                },
                {
                    "warehouse": "CENTRAL", "type": "SECUNDARIO", "location": "C-02",
                    "rge_quantity": None, "counted_quantity": Decimal("2"),
                },
                {
                    "warehouse": "CENTRAL", "type": "SECUNDARIO", "location": "C-04",
                    "rge_quantity": Decimal("8"), "counted_quantity": None,
                },
            ],
        )
        self.assertEqual(abc_row["last_count"], Decimal("0"))
        self.assertEqual(abc_row["stock"], Decimal("1"))
        self.assertEqual(abc_row["calculated_difference"], Decimal("-1"))
        self.assertEqual(
            [(row["warehouse"], row["counted_quantity"]) for row in abc_row["external"]],
            [("LOURDES", Decimal("4")), ("MIRAMAR", Decimal("5")), ("AGDOSA", Decimal("3"))],
        )
        sin_be_row = next(row for row in rows if row["base"][3] == "SIN-BE")
        self.assertEqual(sin_be_row["external"], [])
        table = _combined_inventory_table(rows, _report_styles())
        self.assertEqual(EXTRA_LOCATION_ROWS, 1)
        self.assertEqual(table._cellvalues[0][-1].getPlainText(), "Firma auditor")
        first_product_row_count = (
            len(abc_row["locations"])
            + len(abc_row["external"])
            + EXTRA_LOCATION_ROWS
        )
        first_product_rows = table._cellvalues[3:3 + first_product_row_count]
        self.assertEqual(
            [row[5].getPlainText() for row in first_product_rows[-4:]],
            ["LOURDES", "MIRAMAR", "AGDOSA", ""],
        )
        response = self.client.get(
            reverse("conteo:recount_warehouse_report", args=[load.pk, "central"])
        )
        self.assertEqual(response.status_code, 200)
        pdf_reader = PdfReader(BytesIO(response.content))
        first_page = pdf_reader.pages[0]
        self.assertLess(float(first_page.mediabox.width), float(first_page.mediabox.height))
        pdf_text = "\n".join(page.extract_text() for page in pdf_reader.pages)
        for label in ("Tipo ubicación", "RGE", "Contado", "Firma auditor", "Último conteo", "Existencia", "Diferencia", "AGDOSA", "LOURDES", "MIRAMAR", "ABC-001", "SIN-BE", "C-04"):
            self.assertIn(label, pdf_text)
        self.assertEqual(pdf_text.count("Último conteo"), 2)
        self.assertEqual(pdf_text.count("Existencia"), 2)
        self.assertEqual(pdf_text.count("Diferencia"), 2)
        self.assertIn("Reconteo #1", pdf_text)
        self.assertIn("Bodega Central", pdf_text)
        self.assertNotIn("RESUMEN DE CANTIDADES CONTADAS POR BODEGA", pdf_text)
        self.assertNotIn("FORMATO DE RECONTEO DE INVENTARIO", pdf_text)
        self.assertNotIn("CANTIDADES POR UBICACIÓN Y BODEGA", pdf_text)
        self.assertNotIn("PEDREGAL", pdf_text)
        self.assertNotIn("R-01", pdf_text)
        self.assertNotIn("NO-ELEGIDO", pdf_text)

    def test_recount_pdf_orders_surplus_codes_before_other_differences(self):
        load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="diferencias.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", cargado_por=self.admin
        )
        Ubicacion.objects.create(
            ubicacion="C-01", bodega="CENTRAL", sector="Piso", tipo="Primario"
        )
        for code, subline, last_count, stock in (
            ("FALTANTE-PRIMERO-ALFABETICAMENTE", "01", 2, 5),
            ("EXCEDENTE-B", "03", 8, 5),
            ("EXCEDENTE-A", "02", 7, 5),
        ):
            Diferencia.objects.create(
                carga=load,
                sublinea=subline,
                codigo=code,
                ultimo_conteo=last_count,
                existencia=stock,
                diferencia=last_count - stock,
                fecha_generacion=date(2026, 8, 17),
                hora_generacion=time(9, 30),
                seleccionada=True,
            )
            DetalleConteoInventario.objects.create(
                carga=count_load,
                bodega="CENTRAL",
                ubicacion="C-01",
                codigo=code,
                cantidad=last_count,
            )

        rows = _product_rows(load, "central")

        self.assertEqual(
            [row["base"][3] for row in rows],
            ["EXCEDENTE-A", "EXCEDENTE-B", "FALTANTE-PRIMERO-ALFABETICAMENTE"],
        )

    def test_recount_pdf_location_decision_matrix(self):
        warehouses = ("central", "ransa", "pedregal")
        secondary_sets = (
            ("central",),
            ("ransa",),
            ("pedregal",),
            ("central", "ransa"),
            ("central", "pedregal"),
            ("ransa", "pedregal"),
            ("central", "ransa", "pedregal"),
        )

        for primary_warehouse in warehouses:
            for secondary_warehouses in secondary_sets:
                primary_location = f"P-{primary_warehouse}"
                location_rows = [(
                    primary_warehouse,
                    {"type": "PRIMARIO", "location": primary_location},
                )]
                location_rows.extend(
                    (
                        secondary_warehouse,
                        {
                            "type": "SECUNDARIO",
                            "location": f"S-{secondary_warehouse}",
                        },
                    )
                    for secondary_warehouse in secondary_warehouses
                )

                for report_warehouse in warehouses:
                    visible = _filter_pdf_locations(location_rows, report_warehouse)
                    if report_warehouse == primary_warehouse:
                        expected = {
                            primary_location,
                            *(f"S-{warehouse}" for warehouse in secondary_warehouses),
                        }
                    elif report_warehouse in secondary_warehouses:
                        expected = {f"S-{report_warehouse}"}
                    else:
                        expected = set()
                    self.assertEqual(
                        {row["location"] for row in visible},
                        expected,
                        msg=(primary_warehouse, secondary_warehouses, report_warehouse),
                    )

    def test_central_primary_report_distributes_secondary_locations_by_warehouse(self):
        load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="matriz.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
        )
        Diferencia.objects.create(
            carga=load,
            sublinea="1",
            codigo="MATRIZ-001",
            ultimo_conteo=Decimal("4"),
            existencia=Decimal("5"),
            diferencia=Decimal("1"),
            fecha_generacion=date(2026, 8, 17),
            hora_generacion=time(9, 30),
            seleccionada=True,
        )
        locations = (
            ("C-P", "BODEGA CENTRAL", "PRIMARIO"),
            ("C-S", "BODEGA CENTRAL", "SECUNDARIO"),
            ("R-S", "RANSA", "SECUNDARIO"),
            ("P-S", "BODEGA PEDREGAL", "SECUNDARIO"),
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx",
            total_registros=len(locations),
            total_codigos=1,
            total_bodegas=3,
            cargado_por=self.admin,
        )
        for location, warehouse, location_type in locations:
            Ubicacion.objects.create(
                ubicacion=location,
                bodega=warehouse,
                sector="Piso",
                tipo=location_type,
            )
            DetalleConteoInventario.objects.create(
                carga=count_load,
                bodega=warehouse,
                ubicacion=location,
                codigo="MATRIZ-001",
                cantidad=Decimal("1"),
            )

        page = self.client.get(reverse("conteo:recounts"))
        self.assertEqual(
            {
                option["key"]: option["count"]
                for option in page.context["recounts"][0].warehouse_options
            },
            {"central": 1, "ransa": 1, "pedregal": 1},
        )
        self.assertEqual(
            {row["location"] for row in _product_rows(load, "central")[0]["locations"]},
            {"C-P", "C-S", "R-S", "P-S"},
        )
        self.assertEqual(
            [row["location"] for row in _product_rows(load, "ransa")[0]["locations"]],
            ["R-S"],
        )
        self.assertEqual(
            [row["location"] for row in _product_rows(load, "pedregal")[0]["locations"]],
            ["P-S"],
        )

        pedregal_report = self.client.get(
            reverse("conteo:recount_warehouse_report", args=[load.pk, "pedregal"])
        )
        self.assertEqual(pedregal_report.status_code, 200)
        report_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(BytesIO(pedregal_report.content)).pages
        )
        self.assertIn("P-S", report_text)
        self.assertNotIn("C-P", report_text)
        self.assertNotIn("R-S", report_text)

    def test_recount_pdf_keeps_each_code_on_a_single_page(self):
        rows = []
        for product_index in range(12):
            code = f"Z{product_index:03}X"
            rows.append({
                "base": ["1", "1", str(product_index), code, "Producto para reconteo"],
                "last_count": Decimal("10"),
                "stock": Decimal("8"),
                "calculated_difference": Decimal("2"),
                "scanned": [],
                "rge": [],
                "locations": [
                    {
                        "warehouse": "CENTRAL",
                        "type": "PRIMARIO",
                        "location": f"A{location_index:02}I-B01",
                        "rge_quantity": Decimal("1"),
                        "counted_quantity": Decimal("1"),
                    }
                    for location_index in range(5)
                ],
                "external": [],
            })

        pages = _paginated_inventory_rows(
            rows,
            _report_styles(),
            A4[0] - 10 * mm,
            A4[1] - 25 * mm,
        )
        self.assertGreater(len(pages), 1)
        self.assertEqual(
            [row["base"][3] for page in pages for row in page],
            [row["base"][3] for row in rows],
        )

        load = SimpleNamespace(
            linea=SimpleNamespace(nlinea="1"),
            numero_reconteo=1,
            archivo_as400=None,
        )
        with patch("conteo.recount_reports._product_rows", return_value=rows):
            report, missing_as400 = build_recount_pdf(load, "central")

        self.assertFalse(missing_as400)
        page_texts = [page.extract_text() or "" for page in PdfReader(BytesIO(report)).pages]
        self.assertGreater(len(page_texts), 1)
        total_pages = len(page_texts)
        for page_number, page_text in enumerate(page_texts, start=1):
            self.assertIn(f"Página {page_number}/{total_pages}", page_text)
        for row in rows:
            code = row["base"][3]
            self.assertEqual(sum(code in page_text for page_text in page_texts), 1)

    def test_recount_warehouse_buttons_and_report_use_count_warehouse(self):
        Sublinea.objects.create(
            id_sublinea="SL02",
            linea=self.line,
            numero="2",
            descripcion="Filtros de aceite",
        )
        load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="linea_1.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            total_registros=1,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
        )
        Diferencia.objects.create(
            carga=load,
            sublinea="2",
            item="15",
            codigo="ABC-001",
            descripcion="Filtro premium",
            ultimo_conteo=7,
            existencia=10,
            diferencia=-3,
            fecha_generacion=date(2026, 8, 17),
            hora_generacion=time(9, 30),
            seleccionada=True,
        )
        Diferencia.objects.create(
            carga=load,
            sublinea="2",
            codigo="XYZ-002",
            ultimo_conteo=1,
            existencia=2,
            diferencia=-1,
            fecha_generacion=date(2026, 8, 17),
            hora_generacion=time(9, 30),
            seleccionada=True,
        )
        Ubicacion.objects.create(
            ubicacion="C-01",
            bodega="RANSA",
            sector="Piso",
            tipo="Primario",
        )
        Ubicacion.objects.create(
            ubicacion="R-01",
            bodega="CENTRAL",
            sector="Piso",
            tipo="Primario",
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx",
            total_registros=3,
            total_codigos=2,
            total_bodegas=2,
            cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="BODEGA CENTRAL",
            ubicacion="C-01",
            codigo="abc-001",
            cantidad=Decimal("4"),
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="RANSA",
            ubicacion="R-01",
            codigo="ABC-001",
            cantidad=Decimal("3"),
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="SIN-MAPA",
            codigo="XYZ-002",
            cantidad=Decimal("1"),
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx",
            total_registros=1,
            total_codigos=1,
            total_ubicaciones=1,
            cargado_por=self.admin,
        )
        DetalleRGE.objects.create(
            carga=rge_load,
            codigo="ABC-001",
            ubicacion="C-01",
            cantidad=Decimal("10"),
        )

        page = self.client.get(reverse("conteo:recounts"))
        warehouse_counts = {
            option["key"]: option["count"]
            for option in page.context["recounts"][0].warehouse_options
        }
        self.assertEqual(warehouse_counts, {"central": 2, "ransa": 1, "pedregal": 0})
        self.assertEqual(page.context["recounts"][0].unclassified_codes, [])
        central_rows = _product_rows(load, "central")
        ransa_rows = _product_rows(load, "ransa")
        self.assertEqual(
            [(row["base"][3], row["scanned"]) for row in central_rows],
            [
                ("ABC-001", [("C-01", Decimal("4"))]),
                ("XYZ-002", [("SIN-MAPA", Decimal("1"))]),
            ],
        )
        self.assertEqual(
            [(row["base"][3], row["scanned"]) for row in ransa_rows],
            [("ABC-001", [("R-01", Decimal("3"))])],
        )
        self.assertContains(page, "Bodega Central")
        self.assertContains(page, "RANSA")
        self.assertContains(page, "Generando PDF")
        self.assertContains(
            page,
            reverse("conteo:recount_warehouse_report", args=[load.pk, "central"]),
        )
        response = self.client.get(
            reverse("conteo:recount_warehouse_report", args=[load.pk, "central"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn(".pdf", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreater(len(response.content), 1000)
        self.assertGreaterEqual(len(PdfReader(BytesIO(response.content)).pages), 1)
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
        self.assertNotIn("FORMATO DE RECONTEO DE INVENTARIO", pdf_text)
        self.assertNotIn("CANTIDADES POR UBICACIÓN Y BODEGA", pdf_text)
        self.assertNotIn("RESUMEN DE CANTIDADES CONTADAS POR BODEGA", pdf_text)
        self.assertNotIn("Conteo BE pendiente de carga", pdf_text)

    def test_recount_start_is_manual_and_independent_by_warehouse(self):
        load = CargaDiferencias.objects.create(
            linea=self.line, archivo_nombre="linea_1.xlsx", numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO, total_registros=2,
            cargado_por=self.admin, seleccionado_por=self.admin,
        )
        for code in ("CENTRAL-1", "RANSA-1"):
            Diferencia.objects.create(
                carga=load, sublinea="2", codigo=code, ultimo_conteo=1,
                existencia=2, diferencia=1, fecha_generacion=date(2026, 8, 17),
                hora_generacion=time(9, 30), seleccionada=True,
            )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", total_registros=2, total_codigos=2,
            total_bodegas=2, cargado_por=self.admin,
        )
        for warehouse, code in (("CENTRAL", "CENTRAL-1"), ("RANSA", "RANSA-1")):
            DetalleConteoInventario.objects.create(
                carga=count_load, bodega=warehouse, ubicacion="A-01",
                codigo=code, cantidad=Decimal("1"),
            )

        start_central = reverse("conteo:start_recount_warehouse", args=[load.pk, "central"])
        start_ransa = reverse("conteo:start_recount_warehouse", args=[load.pk, "ransa"])
        finish_central = reverse("conteo:finish_recount_warehouse", args=[load.pk, "central"])
        finish_ransa = reverse("conteo:finish_recount_warehouse", args=[load.pk, "ransa"])
        self.assertContains(self.client.get(reverse("conteo:recounts")), "Iniciar reconteo", count=2)
        self.client.get(reverse("conteo:recount_warehouse_report", args=[load.pk, "central"]))
        self.assertFalse(InicioReconteoBodega.objects.exists())
        self.assertEqual(self.client.get(start_central).status_code, 404)
        self.assertEqual(self.client.get(finish_central).status_code, 404)
        self.assertContains(self.client.post(finish_central, follow=True), "Primero inicia el reconteo")

        self.assertRedirects(self.client.post(start_central), reverse("conteo:recounts"))
        first = InicioReconteoBodega.objects.get(carga=load, bodega="central")
        self.assertEqual(first.iniciado_por, self.admin)
        summary = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(summary.context["status_counts"]["EN_PROCESO"], 1)
        self.assertContains(summary, "En proceso · Bodega Central")
        self.assertEqual(
            self.client.get(reverse("conteo:line_summary"), {"estado": "EN_PROCESO"}).context["filtered_total"], 1
        )
        self.assertContains(self.client.get(reverse("conteo:recounts")), "En reconteo")

        self.client.post(start_central)
        self.assertEqual(InicioReconteoBodega.objects.count(), 1)
        self.client.post(start_ransa)
        self.assertEqual(InicioReconteoBodega.objects.count(), 2)
        self.assertContains(self.client.get(reverse("conteo:line_summary")), "En proceso · Central y RANSA")

        self.assertRedirects(self.client.post(finish_central), reverse("conteo:recounts"))
        first.refresh_from_db()
        self.assertEqual(first.finalizado_por, self.admin)
        self.assertIsNotNone(first.fecha_fin)
        partial = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(partial.context["status_counts"]["EN_PROCESO"], 1)
        self.assertContains(partial, "Finalizó Bodega Central · falta finalizar RANSA")
        partial_recounts = self.client.get(reverse("conteo:recounts"))
        self.assertContains(partial_recounts, "Falta finalizar RANSA")
        self.assertEqual(partial_recounts.context["selected_recount_lines"], 1)

        self.client.post(finish_central)
        first.refresh_from_db()
        self.assertEqual(first.finalizado_por, self.admin)
        self.client.post(finish_ransa)
        load.refresh_from_db()
        self.assertIsNotNone(load.fecha_finalizacion)
        done = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(done.context["status_counts"]["RECONTEO_TERMINADO"], 1)
        self.assertContains(done, "Reconteo terminado · pendiente nueva carga")
        self.assertEqual(
            self.client.get(reverse("conteo:line_summary"), {"estado": "RECONTEO_TERMINADO"}).context["filtered_total"], 1
        )
        recount_page = self.client.get(reverse("conteo:recounts"))
        self.assertEqual(recount_page.context["recount_total"], 0)
        self.assertEqual(recount_page.context["selected_recount_lines"], 0)
        self.assertContains(recount_page, "No hay reconteos pendientes")

    def test_cannot_start_warehouse_without_selected_codes(self):
        load = CargaDiferencias.objects.create(
            linea=self.line, archivo_nombre="linea_1.xlsx", numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO, total_registros=1,
            cargado_por=self.admin, seleccionado_por=self.admin,
        )
        response = self.client.post(
            reverse("conteo:start_recount_warehouse", args=[load.pk, "central"]), follow=True
        )
        self.assertContains(response, "No hay códigos asignados")
        self.assertFalse(InicioReconteoBodega.objects.exists())

    def test_one_assigned_warehouse_can_finish_without_other_warehouse(self):
        load = CargaDiferencias.objects.create(
            linea=self.line, archivo_nombre="linea_1.xlsx", numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO, total_registros=1,
            cargado_por=self.admin, seleccionado_por=self.admin,
        )
        Diferencia.objects.create(
            carga=load, sublinea="2", codigo="SOLO-CENTRAL", ultimo_conteo=1,
            existencia=2, diferencia=1, fecha_generacion=date(2026, 8, 17),
            hora_generacion=time(9, 30), seleccionada=True,
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", total_registros=1, total_codigos=1,
            total_bodegas=1, cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="A-01",
            codigo="SOLO-CENTRAL", cantidad=Decimal("1"),
        )
        self.client.post(reverse("conteo:start_recount_warehouse", args=[load.pk, "central"]))
        self.client.post(reverse("conteo:finish_recount_warehouse", args=[load.pk, "central"]))
        load.refresh_from_db()
        self.assertIsNotNone(load.fecha_finalizacion)
        summary = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(summary.context["status_counts"]["RECONTEO_TERMINADO"], 1)
        recount_page = self.client.get(reverse("conteo:recounts"))
        self.assertEqual(recount_page.context["recount_total"], 0)
        self.assertEqual(recount_page.context["selected_recount_lines"], 0)

        # Una carga terminada antes de este cambio también se reconoce al migrar.
        CargaDiferencias.objects.filter(pk=load.pk).update(
            fecha_finalizacion=None, finalizado_por=None
        )
        from django.apps import apps
        from importlib import import_module

        import_module("conteo.migrations.0009_mark_finished_recounts").mark_finished_recounts(apps, None)
        load.refresh_from_db()
        self.assertIsNotNone(load.fecha_finalizacion)
        self.assertEqual(self.client.get(reverse("conteo:recounts")).context["selected_recount_lines"], 0)

    def test_shared_code_appears_in_each_warehouse_and_unique_code_only_in_its_warehouse(self):
        load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="linea_1.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            total_registros=2,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
        )
        for code in ("CONTADO", "SIN-CONTEO", "SOLO-CENTRAL"):
            Diferencia.objects.create(
                carga=load,
                sublinea="2",
                codigo=code,
                ultimo_conteo=0,
                existencia=5,
                diferencia=5,
                fecha_generacion=date(2026, 8, 17),
                hora_generacion=time(9, 30),
                seleccionada=True,
            )
        Ubicacion.objects.create(
            ubicacion="C-01", bodega="CENTRAL", sector="Piso", tipo="Primario"
        )
        Ubicacion.objects.create(
            ubicacion="R-01", bodega="RANSA", sector="Piso", tipo="Primario"
        )
        Ubicacion.objects.create(
            ubicacion="R-02", bodega="RANSA", sector="Piso", tipo="Secundario"
        )
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx",
            total_registros=3,
            total_codigos=2,
            total_bodegas=2,
            cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="C-01",
            codigo="CONTADO",
            cantidad=Decimal("3"),
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="CENTRAL",
            ubicacion="C-01",
            codigo="SOLO-CENTRAL",
            cantidad=Decimal("2"),
        )
        DetalleConteoInventario.objects.create(
            carga=count_load,
            bodega="RANSA",
            ubicacion="R-02",
            codigo="CONTADO",
            cantidad=Decimal("7"),
        )
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx",
            total_registros=3,
            total_codigos=2,
            total_ubicaciones=2,
            cargado_por=self.admin,
        )
        for code, location, quantity in (
            ("SIN-CONTEO", "C-01", Decimal("5")),
            ("SIN-CONTEO", "R-01", Decimal("6")),
            ("CONTADO", "R-01", Decimal("9")),
        ):
            DetalleRGE.objects.create(
                carga=rge_load, codigo=code, ubicacion=location, cantidad=quantity
            )

        page = self.client.get(reverse("conteo:recounts"))
        warehouse_counts = {
            option["key"]: option["count"]
            for option in page.context["recounts"][0].warehouse_options
        }
        self.assertEqual(warehouse_counts, {"central": 3, "ransa": 2, "pedregal": 0})
        self.assertEqual(
            [(row["base"][3], row["scanned"]) for row in _product_rows(load, "central")],
            [
                ("CONTADO", [("C-01", Decimal("3"))]),
                ("SIN-CONTEO", [("C-01", Decimal("0"))]),
                ("SOLO-CENTRAL", [("C-01", Decimal("2"))]),
            ],
        )
        self.assertEqual(
            [(row["base"][3], row["scanned"]) for row in _product_rows(load, "ransa")],
            [
                ("CONTADO", [("R-02", Decimal("7"))]),
                ("SIN-CONTEO", [("R-01", Decimal("0"))]),
            ],
        )
        for report_warehouse in ("central", "ransa"):
            shared_row = next(
                row for row in _product_rows(load, report_warehouse)
                if row["base"][3] == "CONTADO"
            )
            locations = {
                (row["warehouse"], row["location"]): row
                for row in shared_row["locations"]
            }
            self.assertEqual(
                locations[("CENTRAL", "C-01")]["counted_quantity"], Decimal("3")
            )
            self.assertEqual(
                locations[("RANSA", "R-02")]["counted_quantity"], Decimal("7")
            )
            self.assertEqual(
                locations[("RANSA", "R-01")]["rge_quantity"], Decimal("9")
            )
        for warehouse in ("central", "ransa"):
            response = self.client.get(
                reverse("conteo:recount_warehouse_report", args=[load.pk, warehouse])
            )
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.content.startswith(b"%PDF-"))
            pdf_text = "\n".join(
                page.extract_text() for page in PdfReader(BytesIO(response.content)).pages
            )
            self.assertIn("CONTADO", pdf_text)
            self.assertIn("C-01", pdf_text)
            self.assertIn("R-02", pdf_text)
            if warehouse == "central":
                self.assertIn("SOLO-CENTRAL", pdf_text)
            else:
                self.assertNotIn("SOLO-CENTRAL", pdf_text)

    def test_selected_code_without_source_warehouse_is_visible_for_review(self):
        load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="linea_1.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            total_registros=1,
            cargado_por=self.admin,
            seleccionado_por=self.admin,
        )
        Diferencia.objects.create(
            carga=load,
            sublinea="2",
            codigo="L4",
            ultimo_conteo=32993,
            existencia=32994,
            diferencia=1,
            fecha_generacion=date(2026, 8, 17),
            hora_generacion=time(9, 30),
            seleccionada=True,
        )

        page = self.client.get(reverse("conteo:recounts"))
        recount = page.context["recounts"][0]
        self.assertEqual(recount.unclassified_codes, ["L4"])
        self.assertEqual(
            {option["key"]: option["count"] for option in recount.warehouse_options},
            {"central": 0, "ransa": 0, "pedregal": 0},
        )
        self.assertContains(page, "1 sin bodega")
        self.assertContains(page, "Sin bodega asignada:")

    def test_file_line_must_match_selected_line(self):
        workbook = make_difference_workbook(
            [("99", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no corresponde")
        self.assertEqual(CargaDiferencias.objects.count(), 0)
        self.assertEqual(Diferencia.objects.count(), 0)

    def test_accepts_fecha_geracion_header_from_report(self):
        headers = list(HEADERS)
        headers[10] = "Fecha\r\nGeracion"
        workbook = make_difference_workbook(
            [
                (
                    "1", "2", "ABC-001", 10, 8, -2,
                    "A-01", "", "", "", date(2026, 8, 8), time(14, 30),
                )
            ],
            headers=headers,
        )

        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(CargaDiferencias.objects.get().total_registros, 1)

    def test_accepts_headers_in_second_row(self):
        workbook = make_difference_workbook(
            [
                (
                    "1", "2", "ABC-001", 10, 8, -2,
                    "A-01", "", "", "", date(2026, 8, 8), time(14, 30),
                )
            ],
            rows_before_headers=(("REPORTE DE DIFERENCIAS",),),
        )

        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(CargaDiferencias.objects.get().total_registros, 1)

    def test_normalizes_excel_escaped_line_break_in_fecha_geracion(self):
        from .services import HEADER_ALIASES, _normalize_header

        header = _normalize_header("Fecha_x000D_\nGeracion")

        self.assertEqual(HEADER_ALIASES.get(header, header), "fecha_generacion")

    def test_missing_column_rejects_file(self):
        workbook = make_difference_workbook([], headers=HEADERS[:-1])
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Faltan las columnas obligatorias")
        self.assertEqual(CargaDiferencias.objects.count(), 0)

    def test_admin_can_download_template(self):
        response = self.client.get(reverse("conteo:differences_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

    def test_pending_recount_blocks_next_upload(self):
        first_file = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": first_file},
        )
        second_file = make_difference_workbook(
            [("1", "2", "ABC-002", 10, 7, -3, "A-02", "", "", "", date(2026, 8, 9), time(9, 0))]
        )
        response = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": second_file},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "pendiente de revisión")
        self.assertEqual(CargaDiferencias.objects.count(), 1)

    def test_selection_sends_codes_to_recounts_and_enables_next_upload(self):
        initial_summary = self.client.get(reverse("conteo:line_summary"))
        self.assertContains(initial_summary, "Pte carga")
        self.assertNotContains(initial_summary, "#1")
        self.assertEqual(
            initial_summary.context["status_counts"],
            {"PTE_CARGA": 1, "PTE_GERENCIA": 0, "PTE_INICIAR": 0, "EN_PROCESO": 0, "RECONTEO_TERMINADO": 0, "FINALIZADA": 0},
        )
        first_file = make_difference_workbook(
            [
                ("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30)),
                ("1", "2", "ABC-002", 10, 7, -3, "A-02", "", "", "", date(2026, 8, 8), time(14, 30)),
            ]
        )
        self.client.post(reverse("conteo:differences_upload"), {"linea": self.line.pk, "archivo": first_file})
        manager_review_summary = self.client.get(reverse("conteo:line_summary"))
        self.assertContains(manager_review_summary, "Pte revisión Gerencia")
        self.assertIsNone(manager_review_summary.context["recount_summary"][0]["recount_number"])
        self.assertContains(manager_review_summary, "Carga #1")
        self.assertEqual(manager_review_summary.context["status_counts"]["PTE_GERENCIA"], 1)
        load = CargaDiferencias.objects.get()
        selected_code = load.diferencias.get(codigo="ABC-001")
        response = self.client.post(
            reverse("conteo:select_recount_codes", args=[load.pk]),
            {"codes": [selected_code.pk], "select_all": "0"},
        )
        self.assertRedirects(response, reverse("conteo:pending_queue"))
        load.refresh_from_db()
        self.assertEqual(load.estado, CargaDiferencias.Estado.RECONTEO)
        self.assertEqual(load.seleccionado_por, self.admin)
        self.assertIsNotNone(load.fecha_seleccion)
        self.assertTrue(Diferencia.objects.get(pk=selected_code.pk).seleccionada)
        self.assertEqual(Diferencia.objects.filter(seleccionada=True).count(), 1)
        recounts_page = self.client.get(reverse("conteo:recounts"))
        self.assertContains(recounts_page, "ABC-001")
        self.assertNotContains(recounts_page, "ABC-002")
        self.assertContains(recounts_page, "#1")
        selected_summary = self.client.get(reverse("conteo:line_summary"))
        self.assertContains(selected_summary, "Pte iniciar reconteo")
        self.assertContains(selected_summary, "#1")
        self.assertEqual(selected_summary.context["status_counts"]["PTE_INICIAR"], 1)

        second_file = make_difference_workbook(
            [("1", "2", "ABC-003", 5, 4, -1, "A-03", "", "", "", date(2026, 8, 9), time(9, 0))]
        )
        self.client.post(reverse("conteo:differences_upload"), {"linea": self.line.pk, "archivo": second_file})
        waiting_recounts = self.client.get(reverse("conteo:recounts"))
        self.assertEqual(waiting_recounts.context["recount_total"], 0)
        self.assertEqual(waiting_recounts.context["selected_recount_lines"], 0)
        self.assertEqual(
            list(CargaDiferencias.objects.order_by("numero_reconteo").values_list("numero_reconteo", flat=True)),
            [1, 2],
        )
        second_review_summary = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(second_review_summary.context["recount_summary"][0]["recount_number"], 1)
        self.assertEqual(second_review_summary.context["recount_summary"][0]["status_code"], "PTE_GERENCIA")
        second_load = CargaDiferencias.objects.get(numero_reconteo=2)
        second_code = second_load.diferencias.get(codigo="ABC-003")
        self.client.post(
            reverse("conteo:select_recount_codes", args=[second_load.pk]),
            {"codes": [second_code.pk], "select_all": "0"},
        )
        latest_recounts_page = self.client.get(reverse("conteo:recounts"))
        self.assertEqual(latest_recounts_page.content.count(b'class="recount-card"'), 1)
        self.assertEqual(latest_recounts_page.context["selected_recount_lines"], 1)
        self.assertContains(latest_recounts_page, "ABC-003")
        self.assertNotContains(latest_recounts_page, "ABC-001")
        latest_summary = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(latest_summary.context["recount_summary"][0]["recount_number"], 2)

    def test_line_summary_filters_by_status_without_changing_totals(self):
        pending_line = Linea.objects.create(id_linea="L02", nlinea="2", descripcion="Pendiente Gerencia")
        final_line = Linea.objects.create(id_linea="L03", nlinea="3", descripcion="Línea terminada")
        for line, state in (
            (pending_line, CargaDiferencias.Estado.PENDIENTE),
            (final_line, CargaDiferencias.Estado.FINALIZADO),
        ):
            CargaDiferencias.objects.create(
                linea=line, archivo_nombre="diferencias.xlsx", numero_reconteo=1,
                estado=state, total_registros=1, cargado_por=self.admin,
            )

        all_rows = self.client.get(reverse("conteo:line_summary"))
        self.assertEqual(all_rows.context["line_total"], 3)
        self.assertEqual(all_rows.context["filtered_total"], 3)
        self.assertContains(all_rows, "Filtrar por estado")
        self.assertContains(all_rows, "Todas las líneas")
        self.assertContains(all_rows, 'href="?estado=EN_PROCESO"')
        self.assertNotContains(all_rows, 'class="recount-summary-filters"')
        for status, expected_line in (
            ("PTE_CARGA", self.line),
            ("PTE_GERENCIA", pending_line),
            ("FINALIZADA", final_line),
        ):
            response = self.client.get(reverse("conteo:line_summary"), {"estado": status})
            self.assertEqual(response.context["filtered_total"], 1)
            self.assertEqual(response.context["recount_summary"][0]["line"], expected_line)
            self.assertEqual(response.context["status_counts"], all_rows.context["status_counts"])
            self.assertEqual(response.context["selected_status_label"], dict(response.context["status_filters"])[status])
            self.assertContains(response, 'aria-current="page"')

        empty = self.client.get(reverse("conteo:line_summary"), {"estado": "EN_PROCESO"})
        self.assertEqual(empty.context["filtered_total"], 0)
        self.assertContains(empty, "No hay líneas con este estado")
        invalid = self.client.get(reverse("conteo:line_summary"), {"estado": "desconocido"})
        self.assertEqual(invalid.context["filtered_total"], 3)
        self.assertEqual(invalid.context["selected_status"], "")

    def test_line_summary_shows_cost_totals_from_each_latest_load(self):
        cost_load = CargaBaseCosto.objects.create(
            archivo_nombre="costos.xlsx",
            total_registros=2,
            cargado_por=self.admin,
        )
        CostoProducto.objects.bulk_create([
            CostoProducto(carga=cost_load, codigo="ABC-001", costo=Decimal("12.3456")),
            CostoProducto(carga=cost_load, codigo="ABC-002", costo=Decimal("3.333")),
        ])
        old_load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="anterior.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.FINALIZADO,
            total_registros=1,
            cargado_por=self.admin,
        )
        Diferencia.objects.create(
            carga=old_load, sublinea="2", codigo="ABC-001",
            ultimo_conteo=0, existencia=100, diferencia=100,
            fecha_generacion=date(2026, 8, 7), hora_generacion=time(8, 0),
        )
        latest_load = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="actual.xlsx",
            numero_reconteo=2,
            estado=CargaDiferencias.Estado.PENDIENTE,
            total_registros=2,
            cargado_por=self.admin,
        )
        Diferencia.objects.bulk_create([
            Diferencia(
                carga=latest_load, sublinea="2", codigo="ABC-001",
                ultimo_conteo=10, existencia=8, diferencia=-2,
                fecha_generacion=date(2026, 8, 8), hora_generacion=time(9, 0),
            ),
            Diferencia(
                carga=latest_load, sublinea="2", codigo="ABC-002",
                ultimo_conteo=5.5, existencia=7, diferencia=1.5,
                fecha_generacion=date(2026, 8, 8), hora_generacion=time(9, 1),
            ),
        ])

        response = self.client.get(reverse("conteo:line_summary"))

        item = response.context["recount_summary"][0]
        self.assertEqual(item["load"], latest_load)
        self.assertEqual(item["shortage_total"], "-$24.69")
        self.assertEqual(item["overage_total"], "$5.00")
        self.assertEqual(item["net_total"], "-$19.69")
        self.assertTrue(item["net_is_negative"])
        for heading in ("Faltante", "Sobrante", "Neto"):
            self.assertContains(response, f"<span>{heading}</span>")
        self.assertContains(response, "-$24.69")
        self.assertContains(response, "$5.00")
        self.assertContains(response, "-$19.69")

    def test_line_summary_excludes_retired_catalog_lines(self):
        retired = Linea.objects.create(
            id_linea="ANTIGUA",
            nlinea=self.line.nlinea,
            descripcion="Versión anterior",
            activa=False,
        )
        CargaDiferencias.objects.create(
            linea=retired,
            archivo_nombre="historico.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.FINALIZADO,
            total_registros=1,
            cargado_por=self.admin,
        )

        response = self.client.get(reverse("conteo:line_summary"))

        self.assertEqual(response.context["line_total"], 1)
        self.assertEqual(
            [item["line"].pk for item in response.context["recount_summary"]],
            [self.line.pk],
        )
        self.assertNotContains(response, "Versión anterior")

    def test_manager_can_download_original_file(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(reverse("conteo:differences_upload"), {"linea": self.line.pk, "archivo": workbook})
        load = CargaDiferencias.objects.get()
        response = self.client.get(reverse("conteo:download_recount", args=[load.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])

    def test_line_summary_expands_uploaded_files_by_recount(self):
        first_file = make_difference_workbook([
            ("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30)),
        ])
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": first_file, "archivo_as400": make_as400_pdf()},
        )
        first_load = CargaDiferencias.objects.get()
        self.client.post(
            reverse("conteo:select_recount_codes", args=[first_load.pk]),
            {"select_all": "1"},
        )
        first_start = InicioReconteoBodega.objects.create(
            carga=first_load, bodega="central", iniciado_por=self.admin,
        )
        first_start.fecha_fin = timezone.now()
        first_start.finalizado_por = self.admin
        first_start.save(update_fields=("fecha_fin", "finalizado_por"))
        first_load.fecha_finalizacion = first_start.fecha_fin
        first_load.finalizado_por = self.admin
        first_load.save(update_fields=("fecha_finalizacion", "finalizado_por"))

        second_file = make_difference_workbook([
            ("1", "2", "ABC-002", 10, 7, -3, "A-02", "", "", "", date(2026, 8, 9), time(9, 0)),
        ])
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": second_file},
        )
        second_load = CargaDiferencias.objects.get(numero_reconteo=2)

        page = self.client.get(reverse("conteo:line_summary"))
        history = page.context["recount_summary"][0]["history"]
        self.assertEqual([row["load"].pk for row in history], [second_load.pk, first_load.pk])
        self.assertEqual(history[0]["status_code"], "PTE_GERENCIA")
        self.assertEqual(history[1]["status_code"], "RECONTEO_TERMINADO")
        self.assertIsNotNone(history[1]["started_at"])
        self.assertEqual(history[1]["finished_at"], first_start.fecha_fin)
        self.assertContains(page, '<details class="recount-summary-item">')
        self.assertContains(page, reverse("conteo:download_recount", args=[first_load.pk]))
        self.assertContains(page, reverse("conteo:download_recount", args=[second_load.pk]))
        self.assertContains(page, reverse("conteo:download_recount_as400", args=[first_load.pk]))
        self.assertNotContains(page, reverse("conteo:download_recount_as400", args=[second_load.pk]))

        supervisor = get_user_model().objects.create_user(username="supervisor-descarga")
        supervisor.groups.add(Group.objects.get(name="Supervisores"))
        self.client.force_login(supervisor)
        excel = self.client.get(reverse("conteo:download_recount", args=[first_load.pk]))
        self.assertEqual(excel.status_code, 200)
        self.assertIn("attachment", excel["Content-Disposition"])
        pdf = self.client.get(reverse("conteo:download_recount_as400", args=[first_load.pk]))
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertIn("AS400_linea_1_reconteo_1.pdf", pdf["Content-Disposition"])
        self.assertTrue(b"".join(pdf.streaming_content).startswith(b"%PDF-"))
        self.assertEqual(
            self.client.get(reverse("conteo:download_recount_as400", args=[second_load.pk])).status_code,
            404,
        )
        self.client.logout()
        self.assertEqual(
            self.client.get(reverse("conteo:download_recount", args=[first_load.pk])).status_code,
            302,
        )
        self.assertEqual(
            self.client.get(reverse("conteo:download_recount_as400", args=[first_load.pk])).status_code,
            302,
        )

    def test_manager_can_select_all_codes(self):
        workbook = make_difference_workbook(
            [
                ("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30)),
                ("1", "2", "ABC-002", 10, 7, -3, "A-02", "", "", "", date(2026, 8, 8), time(14, 30)),
            ]
        )
        self.client.post(reverse("conteo:differences_upload"), {"linea": self.line.pk, "archivo": workbook})
        load = CargaDiferencias.objects.get()
        self.client.post(
            reverse("conteo:select_recount_codes", args=[load.pk]),
            {"select_all": "1"},
        )
        self.assertEqual(Diferencia.objects.filter(seleccionada=True).count(), 2)

    def test_manager_can_finalize_line_without_selecting_codes(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        load = CargaDiferencias.objects.get()
        manager = get_user_model().objects.create_user(username="gerente-finaliza", password="test")
        manager.groups.add(Group.objects.get(name="Gerente"))
        self.client.force_login(manager)

        response = self.client.post(reverse("conteo:finalize_line", args=[load.pk]))

        self.assertRedirects(response, reverse("conteo:line_summary"))
        load.refresh_from_db()
        self.assertEqual(load.estado, CargaDiferencias.Estado.FINALIZADO)
        self.assertEqual(load.finalizado_por, manager)
        self.assertIsNotNone(load.fecha_finalizacion)
        self.assertFalse(load.diferencias.filter(seleccionada=True).exists())
        summary = self.client.get(reverse("conteo:line_summary"))
        self.assertContains(summary, "Finalizada")
        self.assertEqual(summary.context["status_counts"]["FINALIZADA"], 1)
        self.assertIsNone(summary.context["recount_summary"][0]["recount_number"])
        self.assertNotContains(self.client.get(reverse("conteo:recounts")), "ABC-001")

        self.client.force_login(self.admin)
        next_workbook = make_difference_workbook(
            [("1", "2", "ABC-002", 10, 7, -3, "A-02", "", "", "", date(2026, 8, 9), time(9, 0))]
        )
        blocked = self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": next_workbook},
        )
        self.assertContains(blocked, "está finalizada y no admite nuevas cargas")
        self.assertEqual(CargaDiferencias.objects.count(), 1)

    def test_admin_can_delete_pending_recount_and_uploaded_file(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook, "archivo_as400": make_as400_pdf()},
        )
        load = CargaDiferencias.objects.get()
        uploaded_path = Path(load.archivo_original.path)
        as400_path = Path(load.archivo_as400.path)

        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse("conteo:delete_pending_recount", args=[load.pk])
            )

        self.assertRedirects(response, reverse("conteo:pending_queue"))
        self.assertFalse(CargaDiferencias.objects.filter(pk=load.pk).exists())
        self.assertEqual(Diferencia.objects.count(), 0)
        self.assertFalse(uploaded_path.exists())
        self.assertFalse(as400_path.exists())

    def test_admin_can_reset_difference_tests_without_erasing_master_data(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook, "archivo_as400": make_as400_pdf()},
        )
        load = CargaDiferencias.objects.get()
        original_path = Path(load.archivo_original.path)
        as400_path = Path(load.archivo_as400.path)
        count_load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx", total_registros=1, total_codigos=1,
            total_bodegas=1, cargado_por=self.admin,
        )
        DetalleConteoInventario.objects.create(
            carga=count_load, bodega="CENTRAL", ubicacion="A-01",
            codigo="ABC-001", cantidad=Decimal("8"),
        )
        Ubicacion.objects.create(ubicacion="A-01", bodega="BODEGA CENTRAL", sector="1", tipo="PRIMARIO")
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", total_registros=1, total_codigos=1,
            total_ubicaciones=1, cargado_por=self.admin,
        )
        DetalleRGE.objects.create(carga=rge_load, codigo="ABC-001", ubicacion="A-01")
        cost_load = CargaBaseCosto.objects.create(
            archivo_nombre="costos.xlsx", total_registros=1, cargado_por=self.admin,
        )
        CostoProducto.objects.create(carga=cost_load, codigo="ABC-001", costo=Decimal("1"))
        self.client.post(reverse("conteo:select_recount_codes", args=[load.pk]), {"select_all": "1"})
        self.client.post(reverse("conteo:start_recount_warehouse", args=[load.pk, "central"]))
        self.client.post(reverse("conteo:finish_recount_warehouse", args=[load.pk, "central"]))

        self.assertContains(self.client.get(reverse("usuarios:inicio")), "Reiniciar pruebas de diferencias")
        preview = self.client.get(reverse("conteo:reset_differences"))
        self.assertEqual(preview.context["load_count"], 1)
        self.assertEqual(preview.context["difference_count"], 1)
        self.assertEqual(preview.context["warehouse_progress_count"], 1)
        self.client.post(reverse("conteo:reset_differences"), {"confirmacion": "incorrecto"})
        self.assertTrue(CargaDiferencias.objects.exists())
        self.assertTrue(original_path.exists())
        self.assertTrue(as400_path.exists())

        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse("conteo:reset_differences"),
                {"confirmacion": "REINICIAR DIFERENCIAS"},
            )
        self.assertRedirects(response, reverse("usuarios:inicio"))
        self.assertFalse(CargaDiferencias.objects.exists())
        self.assertFalse(Diferencia.objects.exists())
        self.assertFalse(InicioReconteoBodega.objects.exists())
        self.assertFalse(original_path.exists())
        self.assertFalse(as400_path.exists())
        self.assertTrue(Linea.objects.filter(pk=self.line.pk).exists())
        self.assertEqual(DetalleConteoInventario.objects.count(), 1)
        self.assertEqual(DetalleRGE.objects.count(), 1)
        self.assertEqual(Ubicacion.objects.count(), 1)
        self.assertEqual(CostoProducto.objects.count(), 1)
        self.assertEqual(self.client.get(reverse("conteo:line_summary")).context["status_counts"]["PTE_CARGA"], 1)

    def test_non_admin_cannot_reset_differences_or_see_dashboard_tool(self):
        viewer = get_user_model().objects.create_user(username="visor-reset", password="clave-segura")
        self.client.force_login(viewer)
        self.assertNotContains(self.client.get(reverse("usuarios:inicio")), "Reiniciar pruebas de diferencias")
        self.assertEqual(self.client.get(reverse("conteo:reset_differences")).status_code, 403)
        self.assertEqual(
            self.client.post(
                reverse("conteo:reset_differences"),
                {"confirmacion": "REINICIAR DIFERENCIAS"},
            ).status_code,
            403,
        )

    def test_admin_can_reset_progress_statuses_without_erasing_sources(self):
        rge_load = CargaRGE.objects.create(
            archivo_nombre="rge.xlsx", cargado_por=self.admin
        )
        for location in ("E08-045Y", "E08-046Y"):
            Ubicacion.objects.create(
                ubicacion=location, bodega="BODEGA CENTRAL",
                sector="8", tipo="SECUNDARIO",
            )
            DetalleRGE.objects.create(
                carga=rge_load, codigo=location, ubicacion=location
            )
        ValidacionUbicacion.objects.create(
            carga_rge=rge_load, ubicacion="E08-045Y",
            validada_por=self.admin, validada=True, modificar_rge=True,
        )
        ValidacionUbicacion.objects.create(
            carga_rge=rge_load, ubicacion="E08-046Y",
            validada_por=self.admin, validada=False, modificar_rge=True,
        )

        home = self.client.get(reverse("usuarios:inicio"))
        self.assertContains(home, "Reiniciar estados de avance")
        preview = self.client.get(reverse("conteo:reset_progress_statuses"))
        self.assertEqual(preview.context["status_count"], 2)
        self.assertEqual(preview.context["validated_count"], 1)
        self.assertEqual(preview.context["modify_rge_count"], 2)

        self.client.post(
            reverse("conteo:reset_progress_statuses"),
            {"confirmacion": "incorrecto"},
        )
        self.assertEqual(ValidacionUbicacion.objects.count(), 2)

        response = self.client.post(
            reverse("conteo:reset_progress_statuses"),
            {"confirmacion": "REINICIAR ESTADOS DE AVANCE"},
        )
        self.assertRedirects(response, reverse("usuarios:inicio"))
        self.assertFalse(ValidacionUbicacion.objects.exists())
        self.assertEqual(DetalleRGE.objects.count(), 2)
        self.assertEqual(Ubicacion.objects.count(), 2)

    def test_non_admin_cannot_reset_progress_statuses(self):
        viewer = get_user_model().objects.create_user(
            username="visor-estados", password="clave-segura"
        )
        self.client.force_login(viewer)
        self.assertNotContains(
            self.client.get(reverse("usuarios:inicio")),
            "Reiniciar estados de avance",
        )
        self.assertEqual(
            self.client.get(reverse("conteo:reset_progress_statuses")).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                reverse("conteo:reset_progress_statuses"),
                {"confirmacion": "REINICIAR ESTADOS DE AVANCE"},
            ).status_code,
            403,
        )

    def test_manager_cannot_delete_pending_recount(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        load = CargaDiferencias.objects.get()
        manager = get_user_model().objects.create_user(username="gerente-delete", password="test")
        manager.groups.add(Group.objects.get(name="Gerente"))
        self.client.force_login(manager)

        response = self.client.post(
            reverse("conteo:delete_pending_recount", args=[load.pk])
        )

        self.assertEqual(response.status_code, 403)
        self.assertTrue(CargaDiferencias.objects.filter(pk=load.pk).exists())

    def test_admin_cannot_delete_recount_selection(self):
        workbook = make_difference_workbook(
            [("1", "2", "ABC-001", 10, 8, -2, "A-01", "", "", "", date(2026, 8, 8), time(14, 30))]
        )
        self.client.post(
            reverse("conteo:differences_upload"),
            {"linea": self.line.pk, "archivo": workbook},
        )
        load = CargaDiferencias.objects.get()
        load.estado = CargaDiferencias.Estado.RECONTEO
        load.save(update_fields=("estado",))

        response = self.client.post(
            reverse("conteo:delete_pending_recount", args=[load.pk]),
            follow=True,
        )

        self.assertContains(response, "Solo se pueden eliminar reconteos pendientes")
        self.assertTrue(CargaDiferencias.objects.filter(pk=load.pk).exists())


def tearDownModule():
    shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)
