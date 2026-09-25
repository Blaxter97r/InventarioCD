from functools import wraps
from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP, localcontext
import logging
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Count, IntegerField, OuterRef, Prefetch, Q, Subquery
from django.db.models.functions import Cast, Upper
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.text import slugify
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from django.utils import timezone

from archivos.models import (
    CargaConteoInventario,
    CargaRGE,
    CostoProducto,
    DetalleConteoInventario,
    DetalleRGE,
    Linea,
    Ubicacion,
)

from .forms import AS400ReplacementForm, DiferenciasUploadForm
from .models import (
    CargaDiferencias,
    Diferencia,
    InicioReconteoBodega,
    ValidacionUbicacion,
)
from .services import DifferenceImportError, import_differences
from .progress import (
    LOCATION_CONFIGS,
    _normalize,
    _parse_location,
    build_code_differences_dashboard,
    build_progress_dashboard,
)
from .progress_reports import (
    build_secondary_quantity_difference_pdf,
    rge_modification_rows,
    secondary_quantity_difference_rows,
)
from .recount_reports import WAREHOUSES, attach_warehouse_options, build_recount_pdf


logger = logging.getLogger(__name__)


def admin_required(view):
    @login_required
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not (
            request.user.is_superuser
            or request.user.groups.filter(name="Admin").exists()
        ):
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return wrapped


def differences_uploader_required(view):
    @login_required
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not (
            request.user.is_superuser
            or request.user.groups.filter(name__in=("Admin", "Supervisores")).exists()
        ):
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return wrapped


def manager_required(view):
    @login_required
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not (
            request.user.is_superuser
            or request.user.groups.filter(name__in=("Admin", "Gerente")).exists()
        ):
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return wrapped


def _ordered_lines():
    return Linea.objects.annotate(
        orden_numerico=Cast("nlinea", IntegerField())
    ).order_by("orden_numerico", "nlinea")


@login_required
def differences_index_view(request):
    selected_line = request.GET.get("linea", "").strip()
    loads = CargaDiferencias.objects.select_related("linea", "cargado_por")
    if selected_line:
        loads = loads.filter(linea_id=selected_line)
    latest_load = loads.first()

    context = {
        "selected_line": selected_line,
        "line_choices": _ordered_lines(),
        "latest_load": latest_load,
        "recent_loads": CargaDiferencias.objects.select_related("linea", "cargado_por")[:6],
        "total_loads": CargaDiferencias.objects.count(),
        "total_differences": Diferencia.objects.count(),
    }
    return render(request, "conteo/differences_index.html", context)


@differences_uploader_required
def differences_upload_view(request):
    form = DiferenciasUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            load = import_differences(
                form.cleaned_data["archivo"],
                form.cleaned_data["linea"],
                request.user,
                form.cleaned_data["archivo_as400"],
            )
        except DifferenceImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(
                request,
                f"Se cargaron {load.total_registros} diferencias para la línea "
                f"{load.linea.nlinea} · {load.linea.descripcion}.",
            )
            target = reverse("conteo:differences_index")
            return redirect(f"{target}?{urlencode({'linea': load.linea_id})}")
    return render(request, "conteo/differences_upload.html", {"form": form})


@differences_uploader_required
def differences_template_view(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Diferencias"
    headers = (
        "linea",
        "sublinea",
        "item",
        "codigo",
        "descripcion",
        "ultimo conteo",
        "existencia",
        "Diferencia",
        "Ubicación 1",
        "Ubicación 2",
        "Ubicación 3",
        "Ubicación 4",
        "Fecha Geracion",
        "Hora Generación",
    )
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:N1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    widths = (14, 15, 14, 22, 38, 18, 15, 15, 20, 20, 20, 20, 20, 18)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    instructions = workbook.create_sheet("Instrucciones", 0)
    instructions.append(["FORMATO PARA CARGAR DIFERENCIAS POR LÍNEA"])
    instructions.append(["Selecciona en la plataforma la misma línea indicada en todas las filas."])
    instructions.append(["No cambies ni elimines los encabezados de la hoja Diferencias."])
    instructions.append(["Ubicación 1 a Ubicación 4 pueden quedar vacías."])
    instructions.append(["Fecha sugerida: DD/MM/AAAA. Hora sugerida: HH:MM:SS."])
    instructions.column_dimensions["A"].width = 90

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_diferencias_por_linea.xlsx"'
    return response


@manager_required
def pending_queue_view(request):
    pending_loads = (
        CargaDiferencias.objects.filter(estado=CargaDiferencias.Estado.PENDIENTE)
        .select_related("linea", "cargado_por")
        .annotate(
            candidate_count=Count(
                "diferencias",
                filter=~Q(diferencias__diferencia=0),
            )
        )
        .order_by("fecha_carga")
    )
    return render(
        request,
        "conteo/pending_queue.html",
        {
            "pending_loads": pending_loads,
            "pending_total": pending_loads.count(),
        },
    )


@manager_required
def pending_codes_view(request, pk):
    load = get_object_or_404(
        CargaDiferencias.objects.select_related("linea"),
        pk=pk,
        estado=CargaDiferencias.Estado.PENDIENTE,
    )
    codes = list(
        load.diferencias.exclude(diferencia=0)
        .order_by("sublinea", "codigo")
        .values(
            "id",
            "sublinea",
            "codigo",
            "ultimo_conteo",
            "existencia",
            "diferencia",
            "ubicacion_1",
            "ubicacion_2",
            "ubicacion_3",
            "ubicacion_4",
        )
    )
    costs_by_code = {
        codigo.upper(): costo
        for codigo, costo in CostoProducto.objects.filter(
            codigo__in={code["codigo"].upper() for code in codes}
        ).values_list("codigo", "costo")
    }
    shortage_total = Decimal("0.00")
    overage_total = Decimal("0.00")
    for code in codes:
        difference = code["diferencia"]
        for field in ("ultimo_conteo", "existencia", "diferencia"):
            code[field] = format(
                code[field].quantize(Decimal("1"), rounding=ROUND_HALF_UP),
                "f",
            )
        cost = costs_by_code.get(code["codigo"].upper())
        if cost is None:
            code["costo"] = None
            code["costo_total"] = None
        else:
            with localcontext() as context:
                context.prec = 50
                total = (cost * difference).quantize(
                    Decimal("0.01"),
                    rounding=ROUND_HALF_UP,
                )
            code["costo"] = format(cost, ".2f")
            code["costo_total"] = format(total, ".2f")
            if total < 0:
                shortage_total += total
            elif total > 0:
                overage_total += total
    net_total = shortage_total + overage_total
    return JsonResponse(
        {
            "id": load.pk,
            "linea": f"Línea {load.linea.nlinea} · {load.linea.descripcion}",
            "reconteo": load.numero_reconteo,
            "archivo": load.archivo_nombre,
            "totales": {
                "faltante": format(shortage_total, ".2f"),
                "sobrante": format(overage_total, ".2f"),
                "diferencia_neta": format(net_total, ".2f"),
            },
            "codes": codes,
        }
    )


@manager_required
def select_recount_codes_view(request, pk):
    if request.method != "POST":
        raise Http404
    with transaction.atomic():
        load = get_object_or_404(
            CargaDiferencias.objects.select_for_update().select_related("linea"),
            pk=pk,
        )
        if load.estado != CargaDiferencias.Estado.PENDIENTE:
            messages.info(request, "Los códigos de este archivo ya fueron seleccionados.")
            return redirect("conteo:pending_queue")

        candidates = load.diferencias.exclude(diferencia=0)
        if request.POST.get("select_all") == "1":
            selected = candidates
        else:
            selected_ids = request.POST.getlist("codes")
            selected = candidates.filter(pk__in=selected_ids)
        if not selected.exists():
            messages.error(request, "Selecciona al menos un código antes de guardar.")
            return redirect("conteo:pending_queue")

        load.diferencias.update(seleccionada=False)
        selected.update(seleccionada=True)
        load.estado = CargaDiferencias.Estado.RECONTEO
        load.seleccionado_por = request.user
        load.fecha_seleccion = timezone.now()
        load.save(update_fields=("estado", "seleccionado_por", "fecha_seleccion"))

    messages.success(
        request,
        f"Se enviaron {selected.count()} códigos de la línea {load.linea.nlinea} "
        "a la bandeja de Reconteos.",
    )
    return redirect("conteo:recounts")


@manager_required
def finalize_line_view(request, pk):
    if request.method != "POST":
        raise Http404

    with transaction.atomic():
        load = get_object_or_404(
            CargaDiferencias.objects.select_for_update().select_related("linea"),
            pk=pk,
        )
        if load.estado != CargaDiferencias.Estado.PENDIENTE:
            messages.info(request, "Este archivo ya fue procesado anteriormente.")
            return redirect("conteo:pending_queue")

        load.diferencias.update(seleccionada=False)
        load.estado = CargaDiferencias.Estado.FINALIZADO
        load.finalizado_por = request.user
        load.fecha_finalizacion = timezone.now()
        load.save(update_fields=("estado", "finalizado_por", "fecha_finalizacion"))

    messages.success(
        request,
        f"La línea {load.linea.nlinea} fue finalizada sin códigos para reconteo.",
    )
    return redirect("conteo:line_summary")


@admin_required
def delete_pending_recount_view(request, pk):
    if request.method != "POST":
        raise Http404

    with transaction.atomic():
        load = get_object_or_404(
            CargaDiferencias.objects.select_for_update().select_related("linea"),
            pk=pk,
        )
        if load.estado != CargaDiferencias.Estado.PENDIENTE:
            messages.error(request, "Solo se pueden eliminar reconteos pendientes.")
            return redirect("conteo:pending_queue")

        line_number = load.linea.nlinea
        recount_number = load.numero_reconteo
        file_name = load.archivo_original.name
        file_storage = load.archivo_original.storage
        as400_file_name = load.archivo_as400.name
        as400_file_storage = load.archivo_as400.storage
        load.delete()
        if file_name:
            transaction.on_commit(lambda: file_storage.delete(file_name))
        if as400_file_name:
            transaction.on_commit(lambda: as400_file_storage.delete(as400_file_name))

    messages.success(
        request,
        f"Reconteo #{recount_number} de la línea {line_number} eliminado correctamente.",
    )
    return redirect("conteo:pending_queue")


@login_required
def download_recount_view(request, pk):
    load = get_object_or_404(CargaDiferencias, pk=pk)
    if not load.archivo_original:
        raise Http404("El archivo original no está disponible.")
    try:
        file_handle = load.archivo_original.open("rb")
    except (FileNotFoundError, OSError) as error:
        raise Http404("El archivo original no está disponible.") from error
    return FileResponse(file_handle, as_attachment=True, filename=load.archivo_nombre)


@login_required
def download_recount_as400_view(request, pk):
    load = get_object_or_404(CargaDiferencias.objects.select_related("linea"), pk=pk)
    if not load.archivo_as400:
        raise Http404("El archivo PDF no está disponible.")
    try:
        file_handle = load.archivo_as400.open("rb")
    except (FileNotFoundError, OSError) as error:
        raise Http404("El archivo PDF no está disponible.") from error
    return FileResponse(
        file_handle,
        as_attachment=True,
        filename=f"AS400_linea_{slugify(load.linea.nlinea)}_reconteo_{load.numero_reconteo}.pdf",
        content_type="application/pdf",
    )


@login_required
def recounts_view(request):
    selected_codes = Diferencia.objects.filter(seleccionada=True).order_by("sublinea", "codigo")
    latest_load_id = CargaDiferencias.objects.filter(
        linea_id=OuterRef("linea_id")
    ).order_by("-fecha_carga", "-pk").values("pk")[:1]
    loads = (
        CargaDiferencias.objects.filter(
            pk=Subquery(latest_load_id),
            estado=CargaDiferencias.Estado.RECONTEO,
            fecha_finalizacion__isnull=True,
        )
        .select_related("linea", "seleccionado_por")
        .prefetch_related(Prefetch("diferencias", queryset=selected_codes, to_attr="codigos_reconteo"))
        .prefetch_related("inicios_bodega")
        .annotate(selected_count=Count("diferencias", filter=Q(diferencias__seleccionada=True)))
        .order_by("linea_id", "-fecha_seleccion", "-pk")
    )
    recounts = sorted(
        loads,
        key=lambda load: (
            int(load.linea.nlinea) if load.linea.nlinea.isdigit() else 10**9,
            load.linea.nlinea,
        ),
    )
    attach_warehouse_options(recounts)
    active_recounts = []
    for load in recounts:
        started = {item.bodega: item for item in load.inicios_bodega.all()}
        load.has_started = bool(started)
        for option in load.warehouse_options:
            option["inicio"] = started.get(option["key"])
        load.finished_warehouse_labels = [
            option["label"] for option in load.warehouse_options
            if option["count"] and option["inicio"] and option["inicio"].fecha_fin
        ]
        load.remaining_warehouse_labels = [
            option["label"] for option in load.warehouse_options
            if option["count"] and not (option["inicio"] and option["inicio"].fecha_fin)
        ]
        load.all_assigned_finished = bool(load.finished_warehouse_labels) and not (
            load.remaining_warehouse_labels or load.unclassified_codes
        )
        if not load.all_assigned_finished:
            active_recounts.append(load)
    return render(
        request,
        "conteo/recounts.html",
        {"recounts": active_recounts, "recount_total": len(active_recounts)},
    )


@admin_required
def replace_recount_as400_view(request, pk):
    if request.method != "POST":
        raise Http404
    load = get_object_or_404(CargaDiferencias, pk=pk, estado=CargaDiferencias.Estado.RECONTEO)
    form = AS400ReplacementForm(request.POST, request.FILES)
    if not form.is_valid():
        errors = " ".join(str(error) for field_errors in form.errors.values() for error in field_errors)
        messages.error(request, "No se adjuntó el PDF AS400. " + errors)
        return redirect("conteo:recounts")

    uploaded = form.cleaned_data["archivo_as400"]
    uploaded.seek(0)
    old_name = load.archivo_as400.name
    load.archivo_as400.save(uploaded.name, uploaded, save=True)
    if old_name and old_name != load.archivo_as400.name:
        load.archivo_as400.storage.delete(old_name)
    messages.success(request, "PDF AS400 adjuntado al reconteo y guardado en el disco persistente.")
    return redirect("conteo:recounts")


@login_required
def start_recount_warehouse_view(request, pk, warehouse):
    if request.method != "POST" or warehouse not in WAREHOUSES:
        raise Http404
    with transaction.atomic():
        load = get_object_or_404(
            CargaDiferencias.objects.select_for_update().select_related("linea"),
            pk=pk,
            estado=CargaDiferencias.Estado.RECONTEO,
        )
        latest = CargaDiferencias.objects.filter(linea=load.linea).first()
        if latest.pk != load.pk:
            messages.error(request, "Este reconteo ya no es el vigente de la línea.")
            return redirect("conteo:recounts")
        load.codigos_reconteo = list(load.diferencias.filter(seleccionada=True))
        attach_warehouse_options([load])
        option = next(item for item in load.warehouse_options if item["key"] == warehouse)
        if not option["count"]:
            messages.error(request, "No hay códigos asignados a esa bodega para este reconteo.")
            return redirect("conteo:recounts")
        _, created = InicioReconteoBodega.objects.get_or_create(
            carga=load, bodega=warehouse, defaults={"iniciado_por": request.user}
        )
    if created:
        messages.success(request, f"Reconteo de {WAREHOUSES[warehouse]} iniciado.")
    else:
        messages.info(request, f"El reconteo de {WAREHOUSES[warehouse]} ya estaba iniciado.")
    return redirect("conteo:recounts")


@login_required
def finish_recount_warehouse_view(request, pk, warehouse):
    if request.method != "POST" or warehouse not in WAREHOUSES:
        raise Http404
    with transaction.atomic():
        load = get_object_or_404(
            CargaDiferencias.objects.select_for_update().select_related("linea"),
            pk=pk, estado=CargaDiferencias.Estado.RECONTEO,
        )
        latest = CargaDiferencias.objects.filter(linea=load.linea).first()
        if latest.pk != load.pk:
            messages.error(request, "Este reconteo ya no es el vigente de la línea.")
            return redirect("conteo:recounts")
        progress = InicioReconteoBodega.objects.select_for_update().filter(
            carga=load, bodega=warehouse
        ).first()
        if progress is None:
            messages.error(request, f"Primero inicia el reconteo de {WAREHOUSES[warehouse]}.")
            return redirect("conteo:recounts")
        if progress.fecha_fin:
            messages.info(request, f"El reconteo de {WAREHOUSES[warehouse]} ya estaba finalizado.")
            return redirect("conteo:recounts")
        progress.finalizado_por = request.user
        progress.fecha_fin = timezone.now()
        progress.save(update_fields=("finalizado_por", "fecha_fin"))
        load.codigos_reconteo = list(load.diferencias.filter(seleccionada=True))
        attach_warehouse_options([load])
        finished_by_warehouse = {
            item.bodega: item for item in load.inicios_bodega.all() if item.fecha_fin
        }
        assigned_warehouses = [
            option["key"] for option in load.warehouse_options if option["count"]
        ]
        if (
            assigned_warehouses
            and not load.unclassified_codes
            and all(key in finished_by_warehouse for key in assigned_warehouses)
        ):
            load.finalizado_por = request.user
            load.fecha_finalizacion = progress.fecha_fin
            load.save(update_fields=("finalizado_por", "fecha_finalizacion"))
    messages.success(request, f"Reconteo de {WAREHOUSES[warehouse]} finalizado.")
    return redirect("conteo:recounts")


@login_required
def recount_warehouse_report_view(request, pk, warehouse):
    if warehouse not in WAREHOUSES:
        raise Http404("La bodega solicitada no existe.")
    load = get_object_or_404(
        CargaDiferencias.objects.select_related("linea", "seleccionado_por"),
        pk=pk,
        estado=CargaDiferencias.Estado.RECONTEO,
    )
    try:
        report, missing_as400 = build_recount_pdf(load, warehouse)
    except Exception:
        logger.exception("Falló la generación del PDF de reconteo: carga=%s bodega=%s", load.pk, warehouse)
        raise
    filename = (
        f"reconteo_linea_{load.linea.nlinea}_{warehouse}_"
        f"{load.numero_reconteo}{'_SIN_AS400' if missing_as400 else ''}.pdf"
    )
    response = HttpResponse(
        report,
        content_type="application/pdf",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    if missing_as400:
        response["X-AS400-Status"] = "unavailable"
    return response


@login_required
def line_summary_view(request):
    status_filters = (
        ("PTE_CARGA", "Pte carga"),
        ("PTE_GERENCIA", "Pte revisión Gerencia"),
        ("PTE_INICIAR", "Pte iniciar reconteo"),
        ("EN_PROCESO", "En proceso de reconteo"),
        ("RECONTEO_TERMINADO", "Reconteo terminado"),
        ("FINALIZADA", "Finalizada"),
    )
    selected_status = request.GET.get("estado", "").strip().upper()
    if selected_status not in dict(status_filters):
        selected_status = ""
    recount_history = (
        CargaDiferencias.objects.all()
        .select_related("linea", "seleccionado_por")
        .prefetch_related("inicios_bodega")
        .order_by("-fecha_carga", "-pk")
    )
    lines = list(_ordered_lines().prefetch_related(
        Prefetch("cargas_diferencias", queryset=recount_history, to_attr="historial_reconteos")
    ))
    current_recounts = [
        line.historial_reconteos[0] for line in lines
        if line.historial_reconteos and line.historial_reconteos[0].estado == CargaDiferencias.Estado.RECONTEO
    ]
    codes_by_load = {load.pk: [] for load in current_recounts}
    for difference in Diferencia.objects.filter(
        carga_id__in=codes_by_load, seleccionada=True
    ).only("carga_id", "codigo"):
        codes_by_load[difference.carga_id].append(difference)
    for load in current_recounts:
        load.codigos_reconteo = codes_by_load[load.pk]
    attach_warehouse_options(current_recounts)
    summary = []
    status_counts = {
        "PTE_CARGA": 0,
        "PTE_GERENCIA": 0,
        "PTE_INICIAR": 0,
        "EN_PROCESO": 0,
        "RECONTEO_TERMINADO": 0,
        "FINALIZADA": 0,
    }
    for line in lines:
        latest = line.historial_reconteos[0] if line.historial_reconteos else None
        started = {}
        if latest is None:
            recount_number = None
            status_code = "PTE_CARGA"
            status_label = "Pte carga"
        elif latest.estado == CargaDiferencias.Estado.PENDIENTE:
            # La carga está esperando la decisión de Gerencia; todavía no
            # existe un nuevo reconteo que numerar.
            recount_number = latest.numero_reconteo - 1 if latest.numero_reconteo > 1 else None
            status_code = "PTE_GERENCIA"
            status_label = "Pte revisión Gerencia"
        elif latest.estado == CargaDiferencias.Estado.RECONTEO:
            # El número nace cuando Gerencia decide enviar códigos a reconteo.
            recount_number = latest.numero_reconteo
            started = {item.bodega: item for item in latest.inicios_bodega.all()}
            assigned = [option for option in latest.warehouse_options if option["count"]]
            finished = [option for option in assigned if started.get(option["key"]) and started[option["key"]].fecha_fin]
            remaining = [option for option in assigned if option not in finished]
            if latest.fecha_finalizacion or (
                assigned and len(finished) == len(assigned) and not latest.unclassified_codes
            ):
                status_code = "RECONTEO_TERMINADO"
                status_label = "Reconteo terminado · pendiente nueva carga"
            elif finished:
                status_code = "EN_PROCESO"
                done_names = " y ".join(option["label"] for option in finished)
                remaining_names = " y ".join(option["label"] for option in remaining)
                status_label = f"Finalizó {done_names} · falta finalizar {remaining_names}" if remaining else f"Finalizó {done_names} · códigos sin bodega"
            elif started:
                status_code = "EN_PROCESO"
                if len(started) == 2:
                    status_label = "En proceso · Central y RANSA"
                else:
                    status_label = f"En proceso · {WAREHOUSES[next(iter(started))]}"
            else:
                status_code = "PTE_INICIAR"
                status_label = "Pte iniciar reconteo"
        else:
            # Si se finalizó la primera carga nunca existió un reconteo. En
            # ciclos posteriores se conserva el último que sí fue ejecutado.
            recount_number = latest.numero_reconteo - 1 if latest.numero_reconteo > 1 else None
            status_code = "FINALIZADA"
            status_label = "Finalizada"
        history = []
        for history_load in line.historial_reconteos:
            warehouse_starts = list(history_load.inicios_bodega.all())
            if history_load.pk == latest.pk:
                history_status_code = status_code
                history_status_label = status_label
            elif history_load.estado == CargaDiferencias.Estado.PENDIENTE:
                history_status_code = "PTE_GERENCIA"
                history_status_label = "Pte revisión Gerencia"
            elif history_load.estado == CargaDiferencias.Estado.FINALIZADO:
                history_status_code = "FINALIZADA"
                history_status_label = "Finalizada sin reconteo"
            elif history_load.fecha_finalizacion:
                history_status_code = "RECONTEO_TERMINADO"
                history_status_label = "Reconteo terminado"
            elif warehouse_starts:
                history_status_code = "EN_PROCESO"
                history_status_label = "En proceso de reconteo"
            else:
                history_status_code = "PTE_INICIAR"
                history_status_label = "Pte iniciar reconteo"
            finished_at = history_load.fecha_finalizacion
            if history_status_code == "RECONTEO_TERMINADO" and not finished_at:
                finished_at = max(
                    (item.fecha_fin for item in warehouse_starts if item.fecha_fin),
                    default=None,
                )
            history.append({
                "load": history_load,
                "status_code": history_status_code,
                "status_label": history_status_label,
                "started_at": min(
                    (item.fecha_inicio for item in warehouse_starts), default=None
                ),
                "finished_at": finished_at,
            })
        summary.append(
            {
                "line": line,
                "load": latest,
                "recount_number": recount_number,
                "status_code": status_code,
                "status_label": status_label,
                "last_progress_at": max((item.fecha_fin or item.fecha_inicio for item in started.values()), default=None),
                "history": history,
            }
        )
        status_counts[status_code] += 1
    visible_summary = [item for item in summary if item["status_code"] == selected_status] if selected_status else summary
    return render(
        request,
        "conteo/line_summary.html",
        {
            "recount_summary": visible_summary,
            "line_total": len(summary),
            "filtered_total": len(visible_summary),
            "selected_status": selected_status,
            "selected_status_label": dict(status_filters).get(selected_status, "Todas las líneas"),
            "status_filters": status_filters,
            "status_counts": status_counts,
        },
    )


@login_required
def selected_codes_view(request):
    return redirect("conteo:recounts")


@admin_required
def reset_differences_view(request):
    confirmation_phrase = "REINICIAR DIFERENCIAS"
    if request.method == "POST":
        if request.POST.get("confirmacion", "").strip() != confirmation_phrase:
            messages.error(request, "Escribe REINICIAR DIFERENCIAS para confirmar.")
        else:
            with transaction.atomic():
                loads = list(
                    CargaDiferencias.objects.select_for_update().values_list(
                        "archivo_original", "archivo_as400"
                    )
                )
                CargaDiferencias.objects.all().delete()
                original_storage = CargaDiferencias._meta.get_field("archivo_original").storage
                as400_storage = CargaDiferencias._meta.get_field("archivo_as400").storage
                for original_name, as400_name in loads:
                    if original_name:
                        transaction.on_commit(
                            lambda name=original_name: original_storage.delete(name), robust=True
                        )
                    if as400_name:
                        transaction.on_commit(
                            lambda name=as400_name: as400_storage.delete(name), robust=True
                        )
            messages.success(
                request,
                f"Se reiniciaron {len(loads)} cargas de diferencias y sus estados de reconteo.",
            )
            return redirect("usuarios:inicio")
    return render(
        request,
        "conteo/reset_differences.html",
        {
            "load_count": CargaDiferencias.objects.count(),
            "difference_count": Diferencia.objects.count(),
            "warehouse_progress_count": InicioReconteoBodega.objects.count(),
            "confirmation_phrase": confirmation_phrase,
        },
    )


@admin_required
def reset_progress_statuses_view(request):
    confirmation_phrase = "REINICIAR ESTADOS DE AVANCE"
    if request.method == "POST":
        if request.POST.get("confirmacion", "").strip() != confirmation_phrase:
            messages.error(
                request, "Escribe REINICIAR ESTADOS DE AVANCE para confirmar."
            )
        else:
            status_count = ValidacionUbicacion.objects.count()
            ValidacionUbicacion.objects.all().delete()
            messages.success(
                request,
                f"Se reiniciaron {status_count} estados de validación del avance.",
            )
            return redirect("usuarios:inicio")
    return render(
        request,
        "conteo/reset_progress_statuses.html",
        {
            "status_count": ValidacionUbicacion.objects.count(),
            "validated_count": ValidacionUbicacion.objects.filter(validada=True).count(),
            "modify_rge_count": ValidacionUbicacion.objects.filter(
                modificar_rge=True
            ).count(),
            "confirmation_phrase": confirmation_phrase,
        },
    )


@login_required
def progress_general_view(request):
    context = build_progress_dashboard(summary_only=True)
    return render(request, "conteo/progress_general.html", context)


@login_required
def progress_view(request):
    context = build_progress_dashboard(
        request.GET.get("piso", "").strip(),
        request.GET.get("bodega", "").strip(),
        request.GET.get("tipo", "").strip(),
    )
    return render(request, "conteo/progress.html", context)


@login_required
def progress_differences_view(request):
    context = build_code_differences_dashboard(
        request.GET.get("bodega", "").strip(),
        request.GET.get("tipo", "").strip(),
        request.GET.get("piso", "").strip(),
    )
    return render(request, "conteo/progress_differences.html", context)


@login_required
def secondary_quantity_differences_pdf_view(request):
    warehouse = _normalize(request.GET.get("bodega", ""))
    scope = _normalize(request.GET.get("piso", ""))
    aisle = _normalize(request.GET.get("pasillo", ""))
    config = LOCATION_CONFIGS.get((warehouse, "SECUNDARIO"))
    if not config or not scope or not aisle:
        raise Http404("Selecciona una bodega, un piso o zona y un pasillo secundario.")
    rows = secondary_quantity_difference_rows(warehouse, scope, aisle)
    pdf = build_secondary_quantity_difference_pdf(rows, warehouse, scope, aisle)
    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="diferencias_secundario_{slugify(warehouse)}_{scope}_{aisle}.pdf"'
    )
    return response


@login_required
def rge_modifications_download_view(request):
    rows = rge_modification_rows()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Modificar RGE"
    sheet.append(("Codigo", "Ubicacion", "Cantidad contada"))
    for row in rows:
        sheet.append((row["code"], row["location"], row["counted_quantity"]))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:C{max(sheet.max_row, 1)}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    for column, width in (("A", 24), ("B", 28), ("C", 20)):
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="ubicaciones_modificar_rge.xlsx"'
    return response


@login_required
@require_POST
@transaction.atomic
def validate_progress_location_view(request):
    location = _normalize(request.POST.get("ubicacion", ""))
    field = request.POST.get("campo", "validada")
    if field not in {"validada", "modificar_rge"}:
        return JsonResponse({"error": "La acción solicitada no es válida."}, status=400)
    requested_value = request.POST.get(field, request.POST.get("valor", "false")) == "true"
    current_rge = CargaRGE.objects.select_for_update().first()
    if not current_rge or not location:
        return JsonResponse({"error": "No hay un RGE vigente para validar."}, status=400)

    map_location = Ubicacion.objects.filter(ubicacion__iexact=location).first()
    config = None
    if map_location:
        config = LOCATION_CONFIGS.get(
            (_normalize(map_location.bodega), _normalize(map_location.tipo))
        )
    belongs_to_progress = bool(
        config
        and _parse_location(location, config)
        and DetalleRGE.objects.filter(
            carga=current_rge, ubicacion__iexact=location
        ).exists()
    )
    if not belongs_to_progress:
        raise Http404("La ubicación no pertenece al avance vigente.")

    current_count = CargaConteoInventario.objects.first()
    is_counted = bool(
        current_count
        and DetalleConteoInventario.objects.filter(
            carga=current_count, ubicacion__iexact=location
        ).exists()
    )
    if field == "validada" and requested_value and is_counted:
        return JsonResponse(
            {"error": "La ubicación ya está contada y no necesita validación."},
            status=409,
        )

    if field == "modificar_rge" and requested_value:
        if _normalize(map_location.tipo) != "SECUNDARIO" or is_counted:
            return JsonResponse(
                {"error": "Solo se puede marcar cuando una ubicación secundaria está pendiente de conteo."},
                status=409,
            )

    validation = ValidacionUbicacion.objects.filter(
        carga_rge=current_rge, ubicacion=location
    ).first()
    created = False
    if requested_value:
        if validation is None:
            validation = ValidacionUbicacion.objects.create(
                carga_rge=current_rge,
                ubicacion=location,
                validada_por=request.user,
                validada=field == "validada",
                modificar_rge=field == "modificar_rge",
            )
            created = True
        else:
            setattr(validation, field, True)
            validation.validada_por = request.user
            validation.save(update_fields=(field, "validada_por"))
        return JsonResponse({
            field: True,
            "created": created,
            "registrado_por": validation.validada_por.get_full_name()
            or validation.validada_por.username,
        })

    if validation is not None:
        setattr(validation, field, False)
        if not validation.validada and not validation.modificar_rge:
            validation.delete()
        else:
            validation.validada_por = request.user
            validation.save(update_fields=(field, "validada_por"))
    return JsonResponse({field: False})


@login_required
def pending_locations_download_view(request, floor):
    context = build_progress_dashboard(
        floor.strip(),
        request.GET.get("bodega", "").strip(),
        request.GET.get("tipo", "").strip(),
    )
    if context["selected_floor"] != floor.strip().upper():
        raise Http404("El piso solicitado no existe en el mapa del mezzanine.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Pendientes {context['selected_floor']}"
    sheet.append(context["download_headers"])
    for row in context["pending_rows"]:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    last_column = chr(64 + len(context["download_headers"]))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip(("A", "B", "C", "D", "E", "F"), (12, 14, 12, 14, 12, 24)):
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="ubicaciones_pendientes_{context["selected_type"].lower()}_{context["selected_floor"]}.xlsx"'
    )
    return response
