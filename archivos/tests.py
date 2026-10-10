from io import BytesIO
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from conteo.models import CargaDiferencias, Diferencia

from .models import (
    AjusteCantidadConteoInventario,
    CargaBaseCosto,
    CargaConteoBE,
    CargaConteoInventario,
    CargaRGE,
    CostoProducto,
    DetalleConteoBE,
    DetalleConteoInventario,
    DetalleRGE,
    Linea,
    Sublinea,
    Ubicacion,
)


def make_workbook(line_rows=None, subline_rows=None):
    workbook = Workbook()
    line_sheet = workbook.active
    line_sheet.title = "Lineas"
    line_sheet.append(("id_linea", "nlinea", "descripcion"))
    for row in line_rows or []:
        line_sheet.append(row)

    subline_sheet = workbook.create_sheet("Sublineas")
    subline_sheet.append(("id_sublinea", "id_linea", "numero", "descripcion"))
    for row in subline_rows or []:
        subline_sheet.append(row)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "catalogo.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def make_location_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ubicaciones"
    sheet.append(("ubicacion", "bodega", "sector", "tipo"))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "ubicaciones.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def make_inventory_count_workbook(rows, headers=None, title_row=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos conteo"
    if title_row:
        sheet.append(("REPORTE DE CONTEO DE INVENTARIO",))
    sheet.append(headers or ("Bodega", "Ubicacion", "Codigo", "Cantidad"))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "conteo_inventario.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def make_be_count_workbook(rows, headers=None, title_rows=0):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conteo BE"
    for index in range(title_rows):
        sheet.append((f"REPORTE BE {index + 1}",))
    sheet.append(headers or ("Bodega", "Codigo", "Cantidad"))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "conteo_be.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def make_rge_workbook(rows, headers=None, title_rows=0):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte general"
    for index in range(title_rows):
        sheet.append((f"REPORTE RGE {index + 1}",))
    sheet.append(headers or ("Descripción", "Código", "Ubicación", "Existencia"))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "rge.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def make_cost_base_workbook(rows, headers=None, title_rows=0):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte costos"
    for index in range(title_rows):
        sheet.append((f"REPORTE DE COSTOS {index + 1}",))
    sheet.append(headers or ("Descripción", "Código", "MCOSUC", "Existencia"))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(
        "base_costo.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class ArchivoAccessTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser(
            username="noel.rivera",
            email="noel@empresa.com",
            password="admin-seguro-2026",
        )
        self.other_user = User.objects.create_user(
            username="supervisor",
            password="supervisor-2026",
        )

    def test_admin_can_view_catalog(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("archivos:index")).status_code, 200)

    def test_other_user_can_view_catalog(self):
        self.client.force_login(self.other_user)
        response = self.client.get(reverse("archivos:index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Permiso de administrador")
        self.assertContains(response, "data-admin-permission-open")

    def test_admin_can_download_template(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("archivos:template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

    def test_admin_can_view_location_map(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("archivos:location_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mapa de ubicaciones")

    def test_other_user_can_view_location_map(self):
        self.client.force_login(self.other_user)
        response = self.client.get(reverse("archivos:location_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Permiso de administrador")

    def test_other_user_cannot_open_upload_pages(self):
        manager = get_user_model().objects.create_user(username="gerente-archivos", password="test-2026")
        manager.groups.add(Group.objects.get(name="Gerente"))
        supervisor = get_user_model().objects.create_user(username="supervisor-archivos", password="test-2026")
        supervisor.groups.add(Group.objects.get(name="Supervisores"))
        routes = (
            "archivos:upload",
            "archivos:location_upload",
            "archivos:inventory_count_upload",
            "archivos:be_count_upload",
            "archivos:rge_upload",
            "archivos:cost_base_upload",
        )
        for user in (self.other_user, supervisor, manager):
            self.client.force_login(user)
            for route in routes:
                with self.subTest(user=user.username, route=route):
                    self.assertEqual(self.client.get(reverse(route)).status_code, 403)

    def test_authenticated_user_can_view_inventory_count(self):
        self.client.force_login(self.other_user)
        response = self.client.get(reverse("archivos:inventory_count_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Conteo de inventario")
        self.assertContains(response, "Permiso de administrador")

    def test_authenticated_user_can_view_rge(self):
        self.client.force_login(self.other_user)
        response = self.client.get(reverse("archivos:rge_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Registro general de existencias")
        self.assertContains(response, "Permiso de administrador")

    def test_authenticated_user_can_view_cost_base(self):
        self.client.force_login(self.other_user)
        response = self.client.get(reverse("archivos:cost_base_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Base de costo")
        self.assertContains(response, "Permiso de administrador")

    def test_catalog_uses_numeric_order_and_groups_sublines(self):
        line_10 = Linea.objects.create(id_linea="A", nlinea="10", descripcion="Diez")
        line_3 = Linea.objects.create(id_linea="B", nlinea="3", descripcion="Tres")
        line_1 = Linea.objects.create(id_linea="C", nlinea="1", descripcion="Uno")
        Sublinea.objects.create(id_sublinea="S12", linea=line_1, numero="12", descripcion="Doce")
        Sublinea.objects.create(id_sublinea="S2", linea=line_1, numero="2", descripcion="Dos")
        Sublinea.objects.create(id_sublinea="S1", linea=line_1, numero="1", descripcion="Uno")

        self.client.force_login(self.admin)
        response = self.client.get(reverse("archivos:index"))

        ordered_lines = list(response.context["lineas"].object_list)
        self.assertEqual([line.nlinea for line in ordered_lines], ["1", "3", "10"])
        self.assertEqual(
            [subline.numero for subline in ordered_lines[0].sublineas.all()],
            ["1", "2", "12"],
        )
        self.assertContains(response, "Catálogo por líneas")


class ExcelImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="noel.rivera",
            email="noel@empresa.com",
            password="admin-seguro-2026",
        )
        self.client.force_login(self.admin)

    def test_valid_workbook_creates_lines_and_sublines(self):
        workbook = make_workbook(
            line_rows=[("L01", "01", "Motores")],
            subline_rows=[("SL001", "L01", "001", "Filtros")],
        )
        response = self.client.post(reverse("archivos:upload"), {"archivo": workbook})
        self.assertRedirects(response, reverse("archivos:index"))
        self.assertTrue(Linea.objects.filter(pk="L01", nlinea="01").exists())
        self.assertTrue(Sublinea.objects.filter(pk="SL001", linea_id="L01").exists())

    def test_existing_ids_are_updated(self):
        Linea.objects.create(id_linea="L01", nlinea="01", descripcion="Anterior")
        workbook = make_workbook(line_rows=[("L01", "01", "Descripción actualizada")])
        response = self.client.post(reverse("archivos:upload"), {"archivo": workbook})
        self.assertRedirects(response, reverse("archivos:index"))
        self.assertEqual(Linea.objects.get(pk="L01").descripcion, "Descripción actualizada")

    def test_new_workbook_replaces_the_visible_catalog(self):
        previous_line = Linea.objects.create(
            id_linea="ANT", nlinea="99", descripcion="Línea anterior"
        )
        Sublinea.objects.create(
            id_sublinea="SUB-ANT",
            linea=previous_line,
            numero="999",
            descripcion="Sublínea anterior",
        )
        workbook = make_workbook(
            line_rows=[("NUEVA", "01", "Línea vigente")],
            subline_rows=[("SUB-NUEVA", "NUEVA", "001", "Sublínea vigente")],
        )

        response = self.client.post(reverse("archivos:upload"), {"archivo": workbook})

        self.assertRedirects(response, reverse("archivos:index"))
        self.assertFalse(Linea.objects.get(pk="ANT").activa)
        self.assertTrue(Linea.objects.get(pk="NUEVA").activa)
        self.assertFalse(Sublinea.objects.filter(pk="SUB-ANT").exists())
        self.assertTrue(Sublinea.objects.filter(pk="SUB-NUEVA").exists())
        page = self.client.get(reverse("archivos:index"))
        self.assertEqual(page.context["total_lineas"], 1)
        self.assertNotContains(page, "Línea anterior")

    def test_unknown_line_rejects_complete_workbook(self):
        workbook = make_workbook(
            line_rows=[("L01", "01", "Motores")],
            subline_rows=[("SL001", "NO-EXISTE", "001", "Filtros")],
        )
        response = self.client.post(reverse("archivos:upload"), {"archivo": workbook})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no existe")
        self.assertEqual(Linea.objects.count(), 0)
        self.assertEqual(Sublinea.objects.count(), 0)

    def test_subline_belongs_to_exactly_one_line(self):
        line = Linea.objects.create(id_linea="L01", nlinea="01", descripcion="Motores")
        subline = Sublinea.objects.create(
            id_sublinea="SL001",
            linea=line,
            numero="001",
            descripcion="Filtros",
        )
        self.assertEqual(subline.linea, line)
        self.assertEqual(list(line.sublineas.all()), [subline])


class LocationImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="noel.rivera",
            email="noel@empresa.com",
            password="admin-seguro-2026",
        )
        self.client.force_login(self.admin)

    def test_valid_workbook_creates_locations(self):
        workbook = make_location_workbook(
            [
                ("A-01-01", "Central", "A", "Picking"),
                ("B-02-03", "Repuestos", "B", "Reserva"),
            ]
        )
        response = self.client.post(reverse("archivos:location_upload"), {"archivo": workbook})
        self.assertRedirects(response, reverse("archivos:location_index"))
        self.assertEqual(Ubicacion.objects.count(), 2)
        self.assertEqual(Ubicacion.objects.get(pk="A-01-01").sector, "A")

    def test_new_map_replaces_previous_locations(self):
        Ubicacion.objects.create(
            ubicacion="A-01-01", bodega="Anterior", sector="A", tipo="Reserva"
        )
        Ubicacion.objects.create(
            ubicacion="OBSOLETA-01", bodega="Anterior", sector="Z", tipo="Reserva"
        )
        workbook = make_location_workbook(
            [
                ("A-01-01", "Central", "A", "Picking"),
                ("NUEVA-01", "Central", "B", "Picking"),
            ]
        )
        response = self.client.post(reverse("archivos:location_upload"), {"archivo": workbook})
        self.assertRedirects(response, reverse("archivos:location_index"))
        self.assertEqual(Ubicacion.objects.count(), 2)
        self.assertFalse(Ubicacion.objects.filter(pk="OBSOLETA-01").exists())
        self.assertTrue(Ubicacion.objects.filter(pk="NUEVA-01").exists())
        location = Ubicacion.objects.get(pk="A-01-01")
        self.assertEqual(location.bodega, "Central")
        self.assertEqual(location.tipo, "Picking")

    def test_duplicate_location_rejects_complete_file(self):
        workbook = make_location_workbook(
            [
                ("A-01-01", "Central", "A", "Picking"),
                ("A-01-01", "Central", "B", "Reserva"),
            ]
        )
        response = self.client.post(reverse("archivos:location_upload"), {"archivo": workbook})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "repetida")
        self.assertEqual(Ubicacion.objects.count(), 0)

    def test_invalid_replacement_preserves_current_map(self):
        Ubicacion.objects.create(
            ubicacion="VIGENTE-01", bodega="Central", sector="A", tipo="Picking"
        )
        workbook = make_location_workbook(
            [
                ("NUEVA-01", "Central", "A", "Picking"),
                ("NUEVA-01", "Central", "B", "Reserva"),
            ]
        )

        response = self.client.post(reverse("archivos:location_upload"), {"archivo": workbook})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "repetida")
        self.assertEqual(Ubicacion.objects.count(), 1)
        self.assertTrue(Ubicacion.objects.filter(pk="VIGENTE-01").exists())
        self.assertFalse(Ubicacion.objects.filter(pk="NUEVA-01").exists())

    def test_admin_can_download_location_template(self):
        response = self.client.get(reverse("archivos:location_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])


class InventoryCountImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="admin-conteo",
            email="conteo@empresa.com",
            password="admin-seguro-2026",
        )
        self.client.force_login(self.admin)

    def test_import_keeps_only_required_columns_and_accepts_header_row_two(self):
        workbook = make_inventory_count_workbook(
            [
                ("ignorar", "Central", "ABC-001", "A-01", "1,5", "extra"),
                ("ignorar", "Central", "ABC-002", "A-02", 0, "extra"),
            ],
            headers=("Descripcion", "Bodega", "Codigo", "Ubicacion", "Cantidad", "Otra"),
            title_row=True,
        )

        response = self.client.post(
            reverse("archivos:inventory_count_upload"),
            {"archivo": workbook},
        )

        self.assertRedirects(response, reverse("archivos:inventory_count_index"))
        load = CargaConteoInventario.objects.get()
        self.assertEqual(load.total_registros, 2)
        self.assertEqual(load.total_codigos, 2)
        self.assertEqual(load.total_bodegas, 1)
        self.assertEqual(load.cargado_por, self.admin)
        self.assertEqual(DetalleConteoInventario.objects.count(), 2)
        first = DetalleConteoInventario.objects.get(codigo="ABC-001")
        second = DetalleConteoInventario.objects.get(codigo="ABC-002")
        self.assertEqual(first.bodega, "Central")
        self.assertEqual(first.ubicacion, "A-01")
        self.assertEqual(first.cantidad, Decimal("1.5"))
        self.assertEqual(second.cantidad, Decimal("0"))

    def test_import_sums_duplicate_code_at_same_warehouse_and_location(self):
        workbook = make_inventory_count_workbook([
            ("Central", "A-01", "ABC-001", 2),
            ("central", "a-01", "abc-001", "3,5"),
            ("Central", "A-02", "ABC-001", 4),
        ])

        response = self.client.post(
            reverse("archivos:inventory_count_upload"),
            {"archivo": workbook},
        )

        self.assertRedirects(response, reverse("archivos:inventory_count_index"))
        load = CargaConteoInventario.objects.get()
        self.assertEqual(load.total_registros, 2)
        self.assertEqual(load.total_codigos, 1)
        self.assertEqual(load.total_bodegas, 1)
        self.assertEqual(DetalleConteoInventario.objects.count(), 2)
        combined = DetalleConteoInventario.objects.get(ubicacion="A-01")
        self.assertEqual(combined.codigo, "ABC-001")
        self.assertEqual(combined.cantidad, Decimal("5.5"))
        separate = DetalleConteoInventario.objects.get(ubicacion="A-02")
        self.assertEqual(separate.cantidad, Decimal("4"))

    def test_missing_required_column_rejects_complete_file(self):
        workbook = make_inventory_count_workbook(
            [("Central", "A-01", "ABC-001")],
            headers=("Bodega", "Ubicacion", "Codigo"),
        )

        response = self.client.post(
            reverse("archivos:inventory_count_upload"),
            {"archivo": workbook},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Faltan las columnas obligatorias: cantidad")
        self.assertEqual(CargaConteoInventario.objects.count(), 0)
        self.assertEqual(DetalleConteoInventario.objects.count(), 0)

    def test_history_keeps_uploads_and_supports_selecting_each_count(self):
        first = make_inventory_count_workbook([("Central", "A-01", "OLD-001", 2)])
        second = make_inventory_count_workbook([("Repuestos", "B-01", "NEW-001", 7)])
        self.client.post(reverse("archivos:inventory_count_upload"), {"archivo": first})
        self.client.post(reverse("archivos:inventory_count_upload"), {"archivo": second})

        response = self.client.get(
            reverse("archivos:inventory_count_index"),
            {"q": "NEW-001"},
        )

        self.assertEqual(CargaConteoInventario.objects.count(), 2)
        self.assertEqual(DetalleConteoInventario.objects.count(), 2)
        self.assertContains(response, "NEW-001")
        self.assertNotContains(response, "OLD-001")
        self.assertEqual(response.context["current_load"], CargaConteoInventario.objects.first())

        older_load = CargaConteoInventario.objects.last()
        older_response = self.client.get(
            reverse("archivos:inventory_count_index"),
            {"carga": older_load.pk},
        )
        self.assertContains(older_response, "OLD-001")
        self.assertNotContains(older_response, "NEW-001")

    def test_admin_can_delete_duplicate_inventory_count(self):
        first = make_inventory_count_workbook([("Central", "A-01", "DUP-001", 2)])
        second = make_inventory_count_workbook([("Central", "A-01", "DUP-001", 2)])
        self.client.post(reverse("archivos:inventory_count_upload"), {"archivo": first})
        self.client.post(reverse("archivos:inventory_count_upload"), {"archivo": second})
        duplicate = CargaConteoInventario.objects.last()

        response = self.client.post(
            reverse("archivos:inventory_count_delete", args=[duplicate.pk])
        )

        self.assertRedirects(response, reverse("archivos:inventory_count_index"))
        self.assertFalse(CargaConteoInventario.objects.filter(pk=duplicate.pk).exists())
        self.assertEqual(CargaConteoInventario.objects.count(), 1)
        self.assertEqual(DetalleConteoInventario.objects.count(), 1)

    def test_admin_can_download_inventory_count_template(self):
        response = self.client.get(reverse("archivos:inventory_count_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])


class InventoryCountQuantityEditTests(TestCase):
    def setUp(self):
        self.supervisor = get_user_model().objects.create_user(
            username="supervisor-conteo", password="test-2026"
        )
        self.supervisor.groups.add(Group.objects.get(name="Supervisores"))
        self.client.force_login(self.supervisor)
        self.load = CargaConteoInventario.objects.create(
            archivo_nombre="conteo.xlsx",
            total_registros=1,
            total_codigos=1,
            total_bodegas=1,
            cargado_por=self.supervisor,
        )
        self.detail = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="A-01",
            codigo="ABC-001",
            cantidad=Decimal("5.0000"),
        )
        self.line = Linea.objects.create(
            id_linea="L-EDIT", nlinea="10", descripcion="Línea de edición"
        )
        self.recount = CargaDiferencias.objects.create(
            linea=self.line,
            archivo_nombre="diferencias.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.RECONTEO,
            cargado_por=self.supervisor,
            seleccionado_por=self.supervisor,
        )
        self.selected_difference = Diferencia.objects.create(
            carga=self.recount,
            sublinea="01",
            codigo="ABC-001",
            ultimo_conteo=5,
            existencia=4,
            diferencia=1,
            fecha_generacion=timezone.localdate(),
            hora_generacion=timezone.localtime().time(),
            seleccionada=True,
        )
        self.url = reverse("archivos:inventory_count_quantity", args=[self.detail.pk])

    def post_quantity(self, quantity, original="5.0000"):
        return self.client.post(self.url, {
            "cantidad": quantity,
            "cantidad_original": original,
            "return_to": reverse("archivos:inventory_count_index") + "?q=ABC-001",
        })

    def test_supervisor_can_increase_and_decrease_current_count_used_by_recount(self):
        from conteo.recount_reports import _group_scanned_locations

        page = self.client.get(reverse("archivos:inventory_count_index"))
        self.assertTrue(page.context["can_edit_count"])
        self.assertContains(page, "Editar cantidades")

        response = self.post_quantity("7")
        self.assertRedirects(response, reverse("archivos:inventory_count_index") + "?q=ABC-001")
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("7"))
        grouped, _ = _group_scanned_locations({"ABC-001"}, "central")
        self.assertEqual(grouped["ABC-001"]["A-01"], Decimal("7"))

        self.post_quantity("2", original="7")
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("2"))
        self.assertEqual(
            list(AjusteCantidadConteoInventario.objects.values_list(
                "cantidad_anterior", "cantidad_nueva", "modificado_por_id"
            )),
            [
                (Decimal("7"), Decimal("2"), self.supervisor.pk),
                (Decimal("5"), Decimal("7"), self.supervisor.pk),
            ],
        )

    def test_invalid_and_stale_quantities_are_not_saved(self):
        for quantity in ("-100000000000000", "1.5", "8,5", "hola", "100000000000000"):
            self.post_quantity(quantity)
            self.detail.refresh_from_db()
            self.assertEqual(self.detail.cantidad, Decimal("5"))
        self.post_quantity("8", original="4")
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("5"))
        self.assertFalse(AjusteCantidadConteoInventario.objects.exists())

    def test_negative_quantity_is_saved_and_used_by_recount_reports(self):
        from conteo.recount_reports import _group_scanned_locations

        response = self.post_quantity("-3")

        self.assertRedirects(
            response,
            reverse("archivos:inventory_count_index") + "?q=ABC-001",
        )
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("-3"))
        adjustment = AjusteCantidadConteoInventario.objects.get()
        self.assertEqual(adjustment.cantidad_anterior, Decimal("5"))
        self.assertEqual(adjustment.cantidad_nueva, Decimal("-3"))
        grouped, _ = _group_scanned_locations({"ABC-001"}, "central")
        self.assertEqual(grouped["ABC-001"]["A-01"], Decimal("-3"))

    def test_historical_count_is_read_only(self):
        newer = CargaConteoInventario.objects.create(
            archivo_nombre="nuevo.xlsx", cargado_por=self.supervisor
        )
        page = self.client.get(reverse("archivos:inventory_count_index"), {"carga": self.load.pk})
        self.assertFalse(page.context["can_edit_count"])
        self.assertNotContains(page, "Editar cantidades")
        self.assertNotEqual(newer.pk, self.load.pk)
        response = self.post_quantity("9")
        self.assertEqual(response.status_code, 403)
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("5"))

    def test_general_edit_mode_searches_codes_and_saves_visible_quantities_together(self):
        second = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="A-02",
            codigo="ABC-001",
            cantidad=Decimal("3"),
        )
        other = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="B-01",
            codigo="XYZ-999",
            cantidad=Decimal("4"),
        )
        index_url = reverse("archivos:inventory_count_index")
        page = self.client.get(index_url, {
            "editar": "1", "reconteo": self.recount.pk, "q": "ABC-001",
        })
        self.assertTrue(page.context["edit_mode"])
        self.assertContains(page, 'name="editar" value="1"')
        self.assertContains(page, f'name="cantidad_{self.detail.pk}"')
        self.assertContains(page, f'name="cantidad_{second.pk}"')
        self.assertNotContains(page, f'name="cantidad_{other.pk}"')
        self.assertContains(page, 'type="number" inputmode="numeric" min="-99999999999999" max="99999999999999" step="1"')
        self.assertNotContains(page, 'inputmode="decimal"')

        response = self.client.post(reverse("archivos:inventory_count_bulk_quantity"), {
            "detalle_ids": [self.detail.pk, second.pk],
            f"original_{self.detail.pk}": "5.0000",
            f"cantidad_{self.detail.pk}": "8",
            f"original_{second.pk}": "3.0000",
            f"cantidad_{second.pk}": "-1",
            "return_to": index_url + "?editar=1&q=ABC-001",
        })
        self.assertRedirects(response, index_url + "?editar=1&q=ABC-001")
        self.detail.refresh_from_db()
        second.refresh_from_db()
        other.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("8"))
        self.assertEqual(second.cantidad, Decimal("-1"))
        self.assertEqual(other.cantidad, Decimal("4"))
        self.assertEqual(AjusteCantidadConteoInventario.objects.count(), 2)

    def test_edit_flow_lists_recount_lines_and_orders_only_selected_codes_like_pdf(self):
        early = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="A-03",
            codigo="ZZZ-001",
            cantidad=Decimal("2"),
        )
        late = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="A-02",
            codigo="AAA-001",
            cantidad=Decimal("3"),
        )
        unselected = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="B-01",
            codigo="NO-RECONTEO",
            cantidad=Decimal("4"),
        )
        Diferencia.objects.create(
            carga=self.recount,
            sublinea="00",
            codigo="ZZZ-001",
            ultimo_conteo=2,
            existencia=1,
            diferencia=1,
            fecha_generacion=timezone.localdate(),
            hora_generacion=timezone.localtime().time(),
            seleccionada=True,
        )
        Diferencia.objects.create(
            carga=self.recount,
            sublinea="02",
            codigo="AAA-001",
            ultimo_conteo=3,
            existencia=1,
            diferencia=2,
            fecha_generacion=timezone.localdate(),
            hora_generacion=timezone.localtime().time(),
            seleccionada=True,
        )
        other_line = Linea.objects.create(
            id_linea="L-02", nlinea="2", descripcion="Segunda línea"
        )
        other_recount = CargaDiferencias.objects.create(
            linea=other_line,
            archivo_nombre="otra.xlsx",
            numero_reconteo=1,
            estado=CargaDiferencias.Estado.FINALIZADO,
            cargado_por=self.supervisor,
            seleccionado_por=self.supervisor,
            finalizado_por=self.supervisor,
            fecha_finalizacion=timezone.now(),
        )
        Diferencia.objects.create(
            carga=other_recount,
            sublinea="01",
            codigo="OTRO-001",
            ultimo_conteo=1,
            existencia=0,
            diferencia=1,
            fecha_generacion=timezone.localdate(),
            hora_generacion=timezone.localtime().time(),
            seleccionada=True,
        )
        index_url = reverse("archivos:inventory_count_index")

        chooser = self.client.get(index_url, {"editar": "1"})

        self.assertEqual(chooser.context["details"].paginator.count, 0)
        self.assertEqual(
            [item.linea.nlinea for item in chooser.context["recount_choices"]],
            ["2", "10"],
        )
        self.assertContains(chooser, "Selecciona la línea del reconteo")
        self.assertContains(chooser, "Segunda línea")

        page = self.client.get(
            index_url,
            {"editar": "1", "reconteo": self.recount.pk},
        )

        displayed = list(page.context["details"].object_list)
        self.assertEqual(
            [item.pk for item in displayed],
            [early.pk, self.detail.pk, late.pk],
        )
        self.assertNotIn(unselected.pk, [item.pk for item in displayed])
        self.assertEqual(page.context["recount_codes_count"], 3)
        self.assertContains(page, "ordenados como el PDF de reconteo")

    def test_general_save_rejects_all_changes_when_one_quantity_is_invalid(self):
        second = DetalleConteoInventario.objects.create(
            carga=self.load,
            bodega="Bodega Central",
            ubicacion="A-02",
            codigo="ABC-002",
            cantidad=Decimal("3"),
        )
        response = self.client.post(reverse("archivos:inventory_count_bulk_quantity"), {
            "detalle_ids": [self.detail.pk, second.pk],
            f"original_{self.detail.pk}": "5.0000",
            f"cantidad_{self.detail.pk}": "8",
            f"original_{second.pk}": "3.0000",
            f"cantidad_{second.pk}": "-2.5",
        })
        self.assertRedirects(response, reverse("archivos:inventory_count_index"))
        self.detail.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("5"))
        self.assertEqual(second.cantidad, Decimal("3"))
        self.assertFalse(AjusteCantidadConteoInventario.objects.exists())

    def test_can_add_an_existing_code_in_each_allowed_warehouse(self):
        add_url = reverse("archivos:inventory_count_location")
        index_url = (
            reverse("archivos:inventory_count_index")
            + f"?editar=1&reconteo={self.recount.pk}&q=ABC-001"
        )

        for warehouse, location in (
            ("CENTRAL", "C-02"),
            ("RANSA", "R-02"),
            ("PEDREGAL", "P-02"),
        ):
            response = self.client.post(add_url, {
                "carga_id": self.load.pk,
                "codigo": "abc-001",
                "bodega": warehouse,
                "ubicacion": location.lower(),
                "cantidad": "6",
                "return_to": index_url,
            })
            self.assertRedirects(response, index_url)

        added = DetalleConteoInventario.objects.filter(carga=self.load).exclude(pk=self.detail.pk)
        self.assertEqual(added.count(), 3)
        self.assertEqual(set(added.values_list("bodega", flat=True)), {"CENTRAL", "RANSA", "PEDREGAL"})
        self.assertTrue(all(item.cantidad == Decimal("6") for item in added))
        self.load.refresh_from_db()
        self.assertEqual(self.load.total_registros, 4)
        self.assertEqual(self.load.total_codigos, 1)
        self.assertEqual(self.load.total_bodegas, 4)
        self.assertEqual(AjusteCantidadConteoInventario.objects.count(), 3)

    def test_can_add_selected_recount_code_missing_from_current_count(self):
        missing_difference = Diferencia.objects.create(
            carga=self.recount,
            sublinea="01",
            codigo="SIN-CONTAR",
            ultimo_conteo=0,
            existencia=4,
            diferencia=-4,
            fecha_generacion=timezone.localdate(),
            hora_generacion=timezone.localtime().time(),
            seleccionada=True,
        )
        index_url = reverse("archivos:inventory_count_index")
        return_to = (
            index_url
            + f"?editar=1&reconteo={self.recount.pk}&q=SIN-CONTAR"
        )

        search_page = self.client.get(index_url, {
            "editar": "1",
            "reconteo": self.recount.pk,
            "q": "SIN-CONTAR",
        })
        self.assertEqual(search_page.context["details"].paginator.count, 0)
        self.assertContains(search_page, 'value="SIN-CONTAR"')

        response = self.client.post(reverse("archivos:inventory_count_location"), {
            "carga_id": self.load.pk,
            "reconteo_id": self.recount.pk,
            "codigo": "sin-contar",
            "bodega": "CENTRAL",
            "ubicacion": "NUEVA-99",
            "cantidad": "4",
            "return_to": return_to,
        })
        self.assertRedirects(response, return_to)

        added = DetalleConteoInventario.objects.get(
            carga=self.load,
            codigo=missing_difference.codigo,
            ubicacion="NUEVA-99",
        )
        self.assertEqual(added.cantidad, Decimal("4"))
        self.assertTrue(AjusteCantidadConteoInventario.objects.filter(
            detalle=added,
            cantidad_anterior=0,
            cantidad_nueva=4,
            modificado_por=self.supervisor,
        ).exists())
        self.assertContains(self.client.get(return_to), "NUEVA-99")
        self.assertContains(self.client.get(index_url, {"q": "SIN-CONTAR"}), "NUEVA-99")

    def test_cannot_add_uncounted_code_not_selected_for_recount(self):
        response = self.client.post(reverse("archivos:inventory_count_location"), {
            "carga_id": self.load.pk,
            "reconteo_id": self.recount.pk,
            "codigo": "NO-SELECCIONADO",
            "bodega": "CENTRAL",
            "ubicacion": "NUEVA-98",
            "cantidad": "1",
        })

        self.assertRedirects(response, reverse("archivos:inventory_count_index"))
        self.assertFalse(DetalleConteoInventario.objects.filter(
            carga=self.load,
            codigo="NO-SELECCIONADO",
        ).exists())

    def test_add_location_rejects_invalid_data_duplicates_and_historical_load(self):
        add_url = reverse("archivos:inventory_count_location")
        valid = {
            "carga_id": self.load.pk,
            "codigo": "ABC-001",
            "bodega": "CENTRAL",
            "ubicacion": "NUEVA-01",
            "cantidad": "2",
        }
        invalid_cases = (
            {**valid, "bodega": "OTRA"},
            {**valid, "cantidad": "1.5"},
            {**valid, "codigo": "NO-EXISTE"},
        )
        for data in invalid_cases:
            with self.subTest(data=data):
                self.client.post(add_url, data)
        self.assertEqual(DetalleConteoInventario.objects.count(), 1)

        self.client.post(add_url, valid)
        self.client.post(add_url, {**valid, "ubicacion": "nueva-01"})
        self.assertEqual(DetalleConteoInventario.objects.count(), 2)

        CargaConteoInventario.objects.create(
            archivo_nombre="nuevo.xlsx", cargado_por=self.supervisor
        )
        response = self.client.post(add_url, {**valid, "ubicacion": "NUEVA-02"})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(DetalleConteoInventario.objects.count(), 2)

    def test_add_location_form_uses_restricted_warehouses_and_integer_quantity(self):
        page = self.client.get(
            reverse("archivos:inventory_count_index"),
            {"editar": "1", "reconteo": self.recount.pk, "q": "ABC-001"},
        )
        self.assertContains(page, "Agregar ubicación adicional")
        self.assertContains(page, 'value="ABC-001"')
        for warehouse in ("CENTRAL", "RANSA", "PEDREGAL"):
            self.assertContains(page, f'<option value="{warehouse}">{warehouse}</option>')
        self.assertContains(page, 'name="cantidad" inputmode="numeric" min="-99999999999999" max="99999999999999" step="1"')

    def test_editor_can_add_location_and_delete_its_row(self):
        editor = get_user_model().objects.create_user(username="editor-filas", password="test-2026")
        editor.groups.add(Group.objects.get(name="Editores"))
        self.client.force_login(editor)
        index_url = (
            reverse("archivos:inventory_count_index")
            + f"?editar=1&reconteo={self.recount.pk}&q=ABC-001"
        )

        response = self.client.post(reverse("archivos:inventory_count_location"), {
            "carga_id": self.load.pk,
            "codigo": "ABC-001",
            "bodega": "RANSA",
            "ubicacion": "R-02",
            "cantidad": "6",
            "return_to": index_url,
        })
        self.assertRedirects(response, index_url)
        added = DetalleConteoInventario.objects.get(ubicacion="R-02")
        page = self.client.get(index_url)
        delete_url = reverse("archivos:inventory_count_detail_delete", args=[added.pk])
        self.assertContains(page, f'formaction="{delete_url}"')

        response = self.client.post(delete_url, {
            f"original_{added.pk}": "6.0000",
            "return_to": index_url,
        })
        self.assertRedirects(response, index_url)
        self.assertFalse(DetalleConteoInventario.objects.filter(pk=added.pk).exists())
        self.assertTrue(DetalleConteoInventario.objects.filter(pk=self.detail.pk).exists())
        self.load.refresh_from_db()
        self.assertEqual((self.load.total_registros, self.load.total_codigos, self.load.total_bodegas), (1, 1, 1))

    def test_delete_row_rejects_stale_and_historical_count(self):
        delete_url = reverse("archivos:inventory_count_detail_delete", args=[self.detail.pk])
        self.assertEqual(self.client.get(delete_url).status_code, 404)
        response = self.client.post(delete_url, {f"original_{self.detail.pk}": "4.0000"})
        self.assertRedirects(response, reverse("archivos:inventory_count_index"))
        self.assertTrue(DetalleConteoInventario.objects.filter(pk=self.detail.pk).exists())

        CargaConteoInventario.objects.create(archivo_nombre="nuevo.xlsx", cargado_por=self.supervisor)
        response = self.client.post(delete_url, {f"original_{self.detail.pk}": "5.0000"})
        self.assertEqual(response.status_code, 403)
        self.assertTrue(DetalleConteoInventario.objects.filter(pk=self.detail.pk).exists())

    def test_general_edit_cannot_update_historical_count_for_any_role(self):
        bulk_url = reverse("archivos:inventory_count_bulk_quantity")
        data = {
            "detalle_ids": [self.detail.pk],
            f"original_{self.detail.pk}": "5.0000",
            f"cantidad_{self.detail.pk}": "8",
        }
        CargaConteoInventario.objects.create(archivo_nombre="nuevo.xlsx", cargado_por=self.supervisor)
        self.assertEqual(self.client.post(bulk_url, data).status_code, 403)
        viewer = get_user_model().objects.create_user(username="viewer-otros", password="test-2026")
        self.client.force_login(viewer)
        self.assertEqual(self.client.post(bulk_url, data).status_code, 403)
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("5"))

    def test_manager_and_other_authenticated_users_can_edit_current_count(self):
        self.assertEqual(self.client.get(self.url).status_code, 404)
        manager = get_user_model().objects.create_user(username="gerente-conteo", password="test-2026")
        manager.groups.add(Group.objects.get(name="Gerente"))
        self.client.force_login(manager)
        page = self.client.get(reverse("archivos:inventory_count_index"))
        self.assertTrue(page.context["can_edit_count"])
        self.assertContains(page, "Editar cantidades")
        self.assertEqual(self.post_quantity("9").status_code, 302)
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("9"))

        viewer = get_user_model().objects.create_user(username="viewer-conteo", password="test-2026")
        self.client.force_login(viewer)
        page = self.client.get(reverse("archivos:inventory_count_index"), {"editar": "1"})
        self.assertTrue(page.context["edit_mode"])
        self.assertEqual(self.client.post(reverse("archivos:inventory_count_bulk_quantity"), {
            "detalle_ids": [self.detail.pk],
            f"original_{self.detail.pk}": "9.0000",
            f"cantidad_{self.detail.pk}": "3",
        }).status_code, 302)
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("3"))
        self.assertEqual(
            list(AjusteCantidadConteoInventario.objects.values_list("modificado_por_id", flat=True)),
            [viewer.pk, manager.pk],
        )

    def test_editor_only_accesses_inventory_count_and_last_change_is_visible(self):
        editor = get_user_model().objects.create_user(
            username="editor-conteo", password="test-2026"
        )
        editor.groups.add(Group.objects.get(name="Editores"))
        self.client.force_login(editor)
        index_url = reverse("archivos:inventory_count_index")

        page = self.client.get(index_url, {"editar": "1", "reconteo": self.recount.pk})
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Conteo de inventario")
        self.assertNotContains(page, 'aria-label="Resumen del conteo"')
        self.assertNotContains(page, "Conteo seleccionado")
        self.assertNotContains(page, "inventory-load-info")
        self.assertContains(page, 'action="' + index_url + '#conteo-vigente"')
        self.assertContains(
            page,
            'value="' + index_url + f'?editar=1&amp;reconteo={self.recount.pk}#conteo-vigente"',
        )
        self.assertNotContains(page, "Cargar conteo")
        self.assertNotContains(page, f'href="{reverse("usuarios:inicio")}"')

        for forbidden_url in (
            reverse("usuarios:inicio"),
            reverse("archivos:index"),
            reverse("archivos:inventory_count_upload"),
            reverse("conteo:differences_index"),
        ):
            with self.subTest(url=forbidden_url):
                self.assertEqual(self.client.get(forbidden_url).status_code, 403)

        return_to = index_url + f"?editar=1&reconteo={self.recount.pk}#conteo-vigente"
        response = self.client.post(
            reverse("archivos:inventory_count_bulk_quantity"),
            {
                "detalle_ids": [self.detail.pk],
                f"original_{self.detail.pk}": "5.0000",
                f"cantidad_{self.detail.pk}": "8",
                "return_to": return_to,
            },
        )
        self.assertRedirects(response, return_to)
        self.detail.refresh_from_db()
        self.assertEqual(self.detail.cantidad, Decimal("8"))

        page = self.client.get(index_url)
        displayed_detail = page.context["details"].object_list[0]
        self.assertEqual(displayed_detail.last_modified_by, editor.username)
        self.assertIsNotNone(displayed_detail.last_modified_at)
        self.assertContains(page, editor.username)
        self.assertNotContains(page, "Sin modificaciones")


class BECountImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="admin-be", email="be@empresa.com", password="test-2026"
        )
        self.client.force_login(self.admin)

    def test_valid_file_consolidates_code_and_warehouse_and_is_searchable(self):
        workbook = make_be_count_workbook([
            ("agdosa", "abc-001", "1,5"),
            ("AGDOSA", "ABC-001", 2),
            ("Lourdes", "ABC-001", 4),
            ("Miramar", "XYZ-999", 0),
        ], title_rows=2)
        response = self.client.post(reverse("archivos:be_count_upload"), {"archivo": workbook})
        self.assertRedirects(response, reverse("archivos:be_count_index"))
        load = CargaConteoBE.objects.get()
        self.assertEqual(load.total_registros, 3)
        self.assertEqual(load.total_codigos, 2)
        self.assertEqual(
            DetalleConteoBE.objects.get(bodega="AGDOSA", codigo="ABC-001").cantidad,
            Decimal("3.5"),
        )
        page = self.client.get(reverse("archivos:be_count_index"), {"q": "ABC", "bodega": "Lourdes"})
        self.assertContains(page, "ABC-001")
        self.assertNotContains(page, "XYZ-999")
        self.assertContains(page, "Conteo BE")

    def test_invalid_replacement_preserves_previous_count(self):
        self.client.post(reverse("archivos:be_count_upload"), {
            "archivo": make_be_count_workbook([("AGDOSA", "ABC-001", 3)])
        })
        for invalid in (
            make_be_count_workbook([("Otra", "ABC-001", 5)]),
            make_be_count_workbook([("Lourdes", "ABC-001", -1)]),
            make_be_count_workbook([("Miramar", "ABC-001", "1.12345")]),
            make_be_count_workbook([("Miramar", "ABC-001", 5)], headers=("Bodega", "Codigo", "Saldo")),
        ):
            response = self.client.post(reverse("archivos:be_count_upload"), {"archivo": invalid})
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.context["form"].errors)
            self.assertEqual(CargaConteoBE.objects.count(), 1)
            self.assertEqual(DetalleConteoBE.objects.get().cantidad, Decimal("3"))

    def test_valid_replacement_and_template(self):
        self.client.post(reverse("archivos:be_count_upload"), {
            "archivo": make_be_count_workbook([("AGDOSA", "OLD-001", 3)])
        })
        self.client.post(reverse("archivos:be_count_upload"), {
            "archivo": make_be_count_workbook([("Miramar", "NEW-001", 7)])
        })
        self.assertEqual(CargaConteoBE.objects.count(), 1)
        self.assertFalse(DetalleConteoBE.objects.filter(codigo="OLD-001").exists())
        self.assertTrue(DetalleConteoBE.objects.filter(codigo="NEW-001").exists())
        response = self.client.get(reverse("archivos:be_count_template"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(list(workbook.active.values)[0], ("Bodega", "Codigo", "Cantidad"))
        workbook.close()


class RGEImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="admin-rge",
            email="rge@empresa.com",
            password="admin-seguro-2026",
        )
        self.client.force_login(self.admin)

    def test_import_keeps_only_code_and_final_location(self):
        workbook = make_rge_workbook(
            [
                ("Producto uno", "abc-001", "a01i-b01", 25),
                ("Producto dos", "ABC-002", "A01I-B02", 10),
            ],
            title_rows=5,
        )

        response = self.client.post(reverse("archivos:rge_upload"), {"archivo": workbook})

        self.assertRedirects(response, reverse("archivos:rge_index"))
        load = CargaRGE.objects.get()
        self.assertEqual(load.total_registros, 2)
        self.assertEqual(load.total_codigos, 2)
        self.assertEqual(load.total_ubicaciones, 2)
        self.assertEqual(
            list(DetalleRGE.objects.values_list("codigo", "ubicacion")),
            [("ABC-001", "A01I-B01"), ("ABC-002", "A01I-B02")],
        )
        self.assertEqual(DetalleRGE.objects.get(codigo="ABC-001").cantidad, Decimal("25"))

    def test_new_rge_replaces_previous_load(self):
        first = make_rge_workbook([("Anterior", "OLD-001", "A01I-B01", 5)])
        second = make_rge_workbook([("Nuevo", "NEW-001", "B01I-C01", 7)])
        self.client.post(reverse("archivos:rge_upload"), {"archivo": first})

        response = self.client.post(reverse("archivos:rge_upload"), {"archivo": second})

        self.assertRedirects(response, reverse("archivos:rge_index"))
        self.assertEqual(CargaRGE.objects.count(), 1)
        self.assertEqual(DetalleRGE.objects.count(), 1)
        self.assertTrue(DetalleRGE.objects.filter(codigo="NEW-001").exists())
        self.assertFalse(DetalleRGE.objects.filter(codigo="OLD-001").exists())

    def test_invalid_replacement_preserves_current_rge(self):
        valid = make_rge_workbook([("Anterior", "OLD-001", "A01I-B01", 5)])
        invalid = make_rge_workbook(
            [("Anterior", "OLD-001", 5)],
            headers=("Descripción", "Código", "Existencia"),
        )
        self.client.post(reverse("archivos:rge_upload"), {"archivo": valid})

        response = self.client.post(reverse("archivos:rge_upload"), {"archivo": invalid})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ubicacion")
        self.assertEqual(CargaRGE.objects.count(), 1)
        self.assertTrue(DetalleRGE.objects.filter(codigo="OLD-001").exists())

    def test_admin_can_download_rge_template(self):
        response = self.client.get(reverse("archivos:rge_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])


class CostBaseImportTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="admin-costos",
            email="costos@empresa.com",
            password="admin-seguro-2026",
        )
        self.client.force_login(self.admin)

    def test_import_keeps_only_code_and_mcosuc(self):
        workbook = make_cost_base_workbook(
            [
                ("Producto uno", "abc-001", Decimal("12.3456"), 25),
                ("Producto dos", "ABC-002", "1,50", 10),
            ],
            title_rows=4,
        )

        response = self.client.post(reverse("archivos:cost_base_upload"), {"archivo": workbook})

        self.assertRedirects(response, reverse("archivos:cost_base_index"))
        load = CargaBaseCosto.objects.get()
        self.assertEqual(load.total_registros, 2)
        self.assertEqual(CostoProducto.objects.get(pk="ABC-001").costo, Decimal("12.3456"))
        self.assertEqual(CostoProducto.objects.get(pk="ABC-002").costo, Decimal("1.5000"))

    def test_new_cost_base_replaces_previous_load(self):
        first = make_cost_base_workbook([("Anterior", "OLD-001", 5, 1)])
        second = make_cost_base_workbook([("Nuevo", "NEW-001", 7, 1)])
        self.client.post(reverse("archivos:cost_base_upload"), {"archivo": first})

        response = self.client.post(reverse("archivos:cost_base_upload"), {"archivo": second})

        self.assertRedirects(response, reverse("archivos:cost_base_index"))
        self.assertEqual(CargaBaseCosto.objects.count(), 1)
        self.assertEqual(CostoProducto.objects.count(), 1)
        self.assertTrue(CostoProducto.objects.filter(pk="NEW-001").exists())
        self.assertFalse(CostoProducto.objects.filter(pk="OLD-001").exists())

    def test_cost_is_displayed_with_decimal_point(self):
        workbook = make_cost_base_workbook([("Producto", "ABC-001", "12.3456", 1)])
        self.client.post(reverse("archivos:cost_base_upload"), {"archivo": workbook})

        response = self.client.get(reverse("archivos:cost_base_index"))

        self.assertContains(response, "12.3456")
        self.assertNotContains(response, "12,3456")

    def test_invalid_replacement_preserves_current_cost_base(self):
        valid = make_cost_base_workbook([("Anterior", "OLD-001", 5, 1)])
        invalid = make_cost_base_workbook(
            [("Nuevo", "NEW-001", 1)],
            headers=("Descripción", "Código", "Existencia"),
        )
        self.client.post(reverse("archivos:cost_base_upload"), {"archivo": valid})

        response = self.client.post(reverse("archivos:cost_base_upload"), {"archivo": invalid})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "mcosuc")
        self.assertEqual(CargaBaseCosto.objects.count(), 1)
        self.assertTrue(CostoProducto.objects.filter(pk="OLD-001").exists())

    def test_conflicting_duplicate_code_rejects_complete_file(self):
        workbook = make_cost_base_workbook(
            [
                ("Producto", "ABC-001", 5, 1),
                ("Producto", "ABC-001", 7, 1),
            ]
        )

        response = self.client.post(reverse("archivos:cost_base_upload"), {"archivo": workbook})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "costos diferentes")
        self.assertEqual(CargaBaseCosto.objects.count(), 0)

    def test_admin_can_download_cost_base_template(self):
        response = self.client.get(reverse("archivos:cost_base_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
