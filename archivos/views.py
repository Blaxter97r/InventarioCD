from functools import wraps
from io import BytesIO
from tempfile import TemporaryFile
from urllib.parse import urlsplit
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import (
    Case,
    Count,
    IntegerField,
    OuterRef,
    Prefetch,
    Q,
    Subquery,
    Value,
    When,
)
from django.db.models.functions import Cast, Lower, Upper
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.http import url_has_allowed_host_and_scheme
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from .forms import BECountUploadForm, CostBaseUploadForm, ExcelUploadForm, InventoryCountDetailForm, InventoryCountLocationForm, InventoryCountQuantityForm, InventoryCountUploadForm, RGEUploadForm
from .models import (
    AjusteCantidadConteoInventario,
    CargaBaseCosto,
    CargaConteoBE,
    CargaConteoInventario,
    CargaRGE,
    CostoProducto,
    DetalleConteoBE,
    DetalleConteoInventario,
    DetalleRGE,
    Linea,
    Sublinea,
    Ubicacion,
)
from .services import (
    ExcelImportError,
    import_be_count_workbook,
    import_catalog_workbook,
    import_cost_base_workbook,
    import_inventory_count_workbook,
    import_location_workbook,
    import_rge_workbook,
)


def admin_required(view):
    @login_required
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not (
            request.user.is_superuser
            or request.user.groups.filter(name="Admin").exists()
        ):
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return wrapped


@login_required
def index_view(request):
    query = request.GET.get("q", "").strip()
    ordered_sublineas = Sublinea.objects.annotate(
        numero_orden=Cast("numero", IntegerField()),
    ).order_by("numero_orden", "numero", "id_sublinea")
    lineas = Linea.objects.filter(activa=True)
    if query:
        lineas = lineas.filter(
            Q(id_linea__icontains=query)
            | Q(nlinea__icontains=query)
            | Q(descripcion__icontains=query)
            | Q(sublineas__id_sublinea__icontains=query)
            | Q(sublineas__numero__icontains=query)
            | Q(sublineas__descripcion__icontains=query)
        )
    lineas = (
        lineas.distinct()
        .annotate(
            total_sublineas=Count("sublineas", distinct=True),
            nlinea_orden=Cast("nlinea", IntegerField()),
        )
        .prefetch_related(Prefetch("sublineas", queryset=ordered_sublineas))
        .order_by("nlinea_orden", "nlinea", "id_linea")
    )

    context = {
        "query": query,
        "lineas": Paginator(lineas, 20).get_page(request.GET.get("lineas_page")),
        "total_lineas": Linea.objects.filter(activa=True).count(),
        "total_sublineas": Sublinea.objects.filter(linea__activa=True).count(),
    }
    return render(request, "archivos/index.html", context)


@admin_required
def upload_view(request):
    form = ExcelUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            result = import_catalog_workbook(form.cleaned_data["archivo"])
        except ExcelImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(
                request,
                "Carga completada: "
                f"{result.lineas_creadas} líneas creadas, "
                f"{result.lineas_actualizadas} actualizadas, "
                f"{result.lineas_retiradas} retiradas, "
                f"{result.sublineas_creadas} sublíneas creadas y "
                f"{result.sublineas_actualizadas} actualizadas; "
                f"{result.sublineas_retiradas} sublíneas retiradas.",
            )
            return redirect("archivos:index")
    return render(request, "archivos/upload.html", {"form": form})


@admin_required
def template_view(request):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instrucciones"
    instructions.append(["FORMATO PARA IMPORTAR LÍNEAS Y SUBLÍNEAS"])
    instructions.append(["1. No cambies los nombres de las hojas ni de las columnas."])
    instructions.append(["2. Guarda los códigos como texto si contienen ceros a la izquierda."])
    instructions.append(["3. Cada id_sublinea debe ser único."])
    instructions.append(["4. Cada id_linea usado en Sublineas debe existir en la hoja Lineas."])
    instructions.append(["5. La carga sustituye por completo el catálogo actual."])
    instructions.append([])
    instructions.append(["Ejemplo Lineas", "L01", "01", "Repuestos de motor"])
    instructions.append(["Ejemplo Sublineas", "SL001", "L01", "001", "Filtros de aceite"])
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 20
    instructions.column_dimensions["C"].width = 16
    instructions.column_dimensions["D"].width = 34

    for title, headers, widths in (
        ("Lineas", ("id_linea", "nlinea", "descripcion"), (18, 18, 46)),
        ("Sublineas", ("id_sublinea", "id_linea", "numero", "descripcion"), (20, 18, 18, 46)),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{chr(64 + len(headers))}1"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F76")
            cell.alignment = Alignment(horizontal="center")
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_lineas_sublineas.xlsx"'
    return response


@login_required
def location_index_view(request):
    query = request.GET.get("q", "").strip()
    selected_warehouse = request.GET.get("bodega", "").strip()
    selected_type = request.GET.get("tipo", "").strip()
    locations = Ubicacion.objects.all()

    if query:
        locations = locations.filter(
            Q(ubicacion__icontains=query)
            | Q(bodega__icontains=query)
            | Q(sector__icontains=query)
            | Q(tipo__icontains=query)
        )
    if selected_warehouse:
        locations = locations.filter(bodega=selected_warehouse)
    if selected_type:
        locations = locations.filter(tipo=selected_type)

    locations = locations.order_by(Lower("bodega"), Lower("sector"), Lower("ubicacion"))
    context = {
        "query": query,
        "selected_warehouse": selected_warehouse,
        "selected_type": selected_type,
        "locations": Paginator(locations, 30).get_page(request.GET.get("page")),
        "total_locations": Ubicacion.objects.count(),
        "total_warehouses": Ubicacion.objects.values("bodega").distinct().count(),
        "total_sectors": Ubicacion.objects.values("bodega", "sector").distinct().count(),
        "warehouses": Ubicacion.objects.order_by("bodega").values_list("bodega", flat=True).distinct(),
        "types": Ubicacion.objects.order_by("tipo").values_list("tipo", flat=True).distinct(),
    }
    return render(request, "archivos/location_index.html", context)


@admin_required
def location_upload_view(request):
    form = ExcelUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            result = import_location_workbook(form.cleaned_data["archivo"])
        except ExcelImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(
                request,
                f"Mapa sustituido: {result.registros} ubicaciones vigentes. "
                f"Se reemplazaron {result.reemplazadas} ubicaciones del mapa anterior.",
            )
            return redirect("archivos:location_index")
    return render(request, "archivos/location_upload.html", {"form": form})


@admin_required
def location_template_view(request):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instrucciones"
    instructions.append(["FORMATO PARA IMPORTAR EL MAPA DE UBICACIONES"])
    instructions.append(["1. No cambies el nombre de la hoja Ubicaciones."])
    instructions.append(["2. No cambies los encabezados de las columnas."])
    instructions.append(["3. Cada ubicación debe ser única."])
    instructions.append(["4. Usa formato texto para conservar ceros iniciales."])
    instructions.append([])
    instructions.append(["Ejemplo", "A-01-01", "Bodega Central", "Sector A", "Picking"])
    for column, width in zip(("A", "B", "C", "D", "E"), (20, 20, 24, 20, 20)):
        instructions.column_dimensions[column].width = width

    sheet = workbook.create_sheet("Ubicaciones")
    headers = ("ubicacion", "bodega", "sector", "tipo")
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:D1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip(("A", "B", "C", "D"), (22, 28, 24, 22)):
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_mapa_ubicaciones.xlsx"'
    return response


@login_required
def inventory_count_index_view(request):
    query = request.GET.get("q", "").strip()
    selected_warehouse = request.GET.get("bodega", "").strip()
    selected_load = request.GET.get("carga", "").strip()
    loads = CargaConteoInventario.objects.select_related("cargado_por")
    latest_load = loads.first()
    if selected_load:
        current_load = get_object_or_404(loads, pk=selected_load)
    else:
        current_load = latest_load
    details = DetalleConteoInventario.objects.none()
    warehouses = []
    new_location_code = ""
    recount_choices = []
    selected_recount = None
    recount_codes = []
    can_edit_count = bool(
        current_load
        and latest_load
        and current_load.pk == latest_load.pk
    )
    edit_mode = can_edit_count and request.GET.get("editar") == "1"
    edit_all = edit_mode and request.GET.get("todos") == "1"

    if edit_mode:
        from conteo.models import CargaDiferencias

        latest_recount_id = (
            CargaDiferencias.objects.filter(
                linea_id=OuterRef("linea_id"),
                diferencias__seleccionada=True,
            )
            .order_by("-fecha_carga", "-pk")
            .values("pk")[:1]
        )
        recount_choices = list(
            CargaDiferencias.objects.filter(
                pk=Subquery(latest_recount_id),
                diferencias__seleccionada=True,
            )
            .select_related("linea", "seleccionado_por")
            .annotate(
                selected_count=Count(
                    "diferencias",
                    filter=Q(diferencias__seleccionada=True),
                    distinct=True,
                )
            )
            .distinct()
        )
        recount_choices.sort(key=lambda item: (
            int(item.linea.nlinea) if item.linea.nlinea.isdigit() else 10**9,
            item.linea.nlinea,
        ))
        selected_recount_id = request.GET.get("reconteo", "").strip()
        if selected_recount_id:
            selected_recount = next(
                (item for item in recount_choices if str(item.pk) == selected_recount_id),
                None,
            )
            if selected_recount is None:
                raise Http404("El reconteo seleccionado no está disponible para edición.")
            edit_all = False
            selected_differences = list(
                selected_recount.diferencias.filter(seleccionada=True)
                .order_by("sublinea", "codigo", "pk")
            )
            seen_codes = set()
            for difference in selected_differences:
                normalized_code = difference.codigo.strip().upper()
                if normalized_code and normalized_code not in seen_codes:
                    seen_codes.add(normalized_code)
                    recount_codes.append(normalized_code)

    if current_load:
        latest_adjustment = (
            AjusteCantidadConteoInventario.objects.filter(detalle_id=OuterRef("pk"))
            .order_by("-fecha", "-pk")
        )
        details = current_load.detalles.annotate(
            last_modified_by=Subquery(
                latest_adjustment.values("modificado_por__username")[:1]
            ),
            last_modified_at=Subquery(latest_adjustment.values("fecha")[:1]),
        )
        if edit_mode:
            if selected_recount:
                recount_order = Case(
                    *[
                        When(codigo__iexact=code, then=Value(position))
                        for position, code in enumerate(recount_codes)
                    ],
                    default=Value(len(recount_codes)),
                    output_field=IntegerField(),
                )
                details = (
                    details.annotate(codigo_normalizado=Upper("codigo"))
                    .filter(codigo_normalizado__in=recount_codes)
                    .annotate(recount_order=recount_order)
                    .order_by("recount_order", "bodega", "ubicacion", "pk")
                )
            elif not edit_all:
                details = details.none()
        warehouses = details.order_by("bodega").values_list("bodega", flat=True).distinct()
        if query:
            details = details.filter(
                Q(codigo__icontains=query)
                | Q(ubicacion__icontains=query)
                | Q(bodega__icontains=query)
            )
        if selected_warehouse:
            details = details.filter(bodega=selected_warehouse)
        if request.GET.get("editar") == "1" and query:
            new_location_code = (
                current_load.detalles.filter(codigo__iexact=query)
                .order_by("pk")
                .values_list("codigo", flat=True)
                .first()
                or ""
            )
            if not new_location_code and selected_recount:
                normalized_query = query.upper()
                new_location_code = next(
                    (code for code in recount_codes if code == normalized_query),
                    "",
                )

    return render(
        request,
        "archivos/inventory_count_index.html",
        {
            "query": query,
            "selected_warehouse": selected_warehouse,
            "selected_load": selected_load,
            "current_load": current_load,
            "can_edit_count": can_edit_count,
            "edit_mode": edit_mode,
            "edit_all": edit_all,
            "recount_choices": recount_choices,
            "selected_recount": selected_recount,
            "recount_codes_count": len(recount_codes),
            "new_location_code": new_location_code,
            "details": Paginator(details, 50).get_page(request.GET.get("page")),
            "warehouses": warehouses,
            "load_history": Paginator(loads, 10).get_page(request.GET.get("cargas_page")),
        },
    )


@login_required
def inventory_count_download_view(request):
    current_load = CargaConteoInventario.objects.first()
    if current_load is None:
        raise Http404("No hay un conteo vigente para descargar.")

    # A normal openpyxl workbook keeps every cell in memory.  The inventory can
    # contain hundreds of thousands of rows, so build it in write-only mode and
    # let Django stream the finished file from a temporary file instead.
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Conteo")
    sheet.freeze_panes = "A2"
    for column, width in zip(("A", "B", "C", "D"), (24, 24, 24, 16)):
        sheet.column_dimensions[column].width = width
    header = [
        WriteOnlyCell(sheet, value=value)
        for value in ("Bodega", "Ubicacion", "Codigo", "Cantidad")
    ]
    for cell in header:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    sheet.append(header)

    details = current_load.detalles.order_by(
        "bodega", "ubicacion", "codigo", "pk"
    ).values_list("bodega", "ubicacion", "codigo", "cantidad")
    row_count = 1
    for bodega, ubicacion, codigo, cantidad in details.iterator(chunk_size=2000):
        row = []
        for value in (bodega, ubicacion, codigo):
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"
            row.append(cell)
        quantity_cell = WriteOnlyCell(sheet, value=cantidad)
        quantity_cell.number_format = "0.####"
        row.append(quantity_cell)
        sheet.append(row)
        row_count += 1

    sheet.auto_filter.ref = f"A1:D{row_count}"

    output = TemporaryFile()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename="conteo_inventario_vigente.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@login_required
def inventory_count_location_view(request):
    if request.method != "POST":
        raise Http404

    index_url = reverse("archivos:inventory_count_index")
    return_to = request.POST.get("return_to", "")
    if not (
        url_has_allowed_host_and_scheme(return_to, allowed_hosts={request.get_host()})
        and urlsplit(return_to).path == index_url
    ):
        return_to = index_url

    form = InventoryCountLocationForm(request.POST)
    if not form.is_valid():
        errors = " ".join(
            str(error)
            for field_errors in form.errors.values()
            for error in field_errors
        )
        messages.error(request, "No se agregó la ubicación. " + errors)
        return redirect(return_to)

    try:
        submitted_load_id = int(request.POST.get("carga_id", ""))
    except ValueError:
        raise Http404

    with transaction.atomic():
        latest_load = CargaConteoInventario.objects.select_for_update().first()
        if not latest_load or submitted_load_id != latest_load.pk:
            raise PermissionDenied("Solo se puede modificar el conteo vigente.")

        code = form.cleaned_data["codigo"]
        existing_code = (
            latest_load.detalles.filter(codigo__iexact=code)
            .order_by("pk")
            .values_list("codigo", flat=True)
            .first()
        )
        if not existing_code:
            from conteo.models import CargaDiferencias

            try:
                recount_id = int(request.POST.get("reconteo_id", ""))
            except (TypeError, ValueError):
                recount_id = None

            selected_code = None
            if recount_id:
                recount = (
                    CargaDiferencias.objects.filter(
                        pk=recount_id,
                        diferencias__seleccionada=True,
                    )
                    .distinct()
                    .first()
                )
                if recount:
                    latest_recount_id = (
                        CargaDiferencias.objects.filter(
                            linea_id=recount.linea_id,
                            diferencias__seleccionada=True,
                        )
                        .order_by("-fecha_carga", "-pk")
                        .values_list("pk", flat=True)
                        .first()
                    )
                    if recount.pk == latest_recount_id:
                        selected_code = (
                            recount.diferencias.filter(
                                seleccionada=True,
                                codigo__iexact=code,
                            )
                            .order_by("pk")
                            .values_list("codigo", flat=True)
                            .first()
                        )

            if not selected_code:
                messages.error(
                    request,
                    f"No se agregó la ubicación. El código {code} no pertenece al conteo vigente ni al reconteo seleccionado.",
                )
                return redirect(return_to)
            existing_code = selected_code

        warehouse = form.cleaned_data["bodega"]
        location = form.cleaned_data["ubicacion"]
        duplicate = latest_load.detalles.filter(
            codigo__iexact=existing_code,
            bodega__iexact=warehouse,
            ubicacion__iexact=location,
        ).exists()
        if duplicate:
            messages.error(
                request,
                f"No se agregó la ubicación. {existing_code} ya existe en {warehouse} / {location}.",
            )
            return redirect(return_to)

        quantity = form.cleaned_data["cantidad"]
        detail = DetalleConteoInventario.objects.create(
            carga=latest_load,
            codigo=existing_code,
            bodega=warehouse,
            ubicacion=location,
            cantidad=quantity,
        )
        AjusteCantidadConteoInventario.objects.create(
            detalle=detail,
            cantidad_anterior=0,
            cantidad_nueva=quantity,
            bodega_anterior="",
            bodega_nueva=warehouse,
            ubicacion_anterior="",
            ubicacion_nueva=location,
            modificado_por=request.user,
        )
        latest_load.total_registros = latest_load.detalles.count()
        latest_load.total_codigos = latest_load.detalles.values("codigo").distinct().count()
        latest_load.total_bodegas = latest_load.detalles.values("bodega").distinct().count()
        latest_load.save(
            update_fields=("total_registros", "total_codigos", "total_bodegas")
        )

    messages.success(
        request,
        f"Se agregó {existing_code} en {warehouse} / {location} con cantidad {quantity}.",
    )
    return redirect(return_to)


@login_required
def inventory_count_bulk_quantity_view(request):
    if request.method != "POST":
        raise Http404

    index_url = reverse("archivos:inventory_count_index")
    return_to = request.POST.get("return_to", "")
    if not (
        url_has_allowed_host_and_scheme(return_to, allowed_hosts={request.get_host()})
        and urlsplit(return_to).path == index_url
    ):
        return_to = index_url

    raw_ids = request.POST.getlist("detalle_ids")
    try:
        ids = [int(raw_id) for raw_id in raw_ids]
    except ValueError:
        raise Http404
    if not ids or len(ids) > 50 or len(set(ids)) != len(ids):
        raise Http404

    with transaction.atomic():
        latest_load = CargaConteoInventario.objects.first()
        details = list(DetalleConteoInventario.objects.select_for_update().filter(pk__in=ids))
        if len(details) != len(ids):
            raise Http404
        if not latest_load or any(detail.carga_id != latest_load.pk for detail in details):
            raise PermissionDenied("Solo se puede modificar el conteo vigente.")

        changes = []
        for detail in details:
            original = request.POST.get(f"original_{detail.pk}", "")
            try:
                original_quantity = Decimal(original)
            except InvalidOperation:
                original_quantity = None
            if original_quantity != detail.cantidad:
                messages.error(request, "Una cantidad cambió desde que abriste la página. Actualiza antes de guardar.")
                return redirect(return_to)

            original_warehouse = request.POST.get(
                f"original_bodega_{detail.pk}", detail.bodega
            )
            original_location = request.POST.get(
                f"original_ubicacion_{detail.pk}", detail.ubicacion
            )
            if (
                original_warehouse != detail.bodega
                or original_location != detail.ubicacion
            ):
                messages.error(
                    request,
                    "Una bodega o ubicación cambió desde que abriste la página. Actualiza antes de guardar.",
                )
                return redirect(return_to)

            raw_quantity = request.POST.get(f"cantidad_{detail.pk}", "")
            raw_warehouse = request.POST.get(f"bodega_{detail.pk}", detail.bodega)
            raw_location = request.POST.get(f"ubicacion_{detail.pk}", detail.ubicacion)
            form = InventoryCountDetailForm(
                {
                    "cantidad": raw_quantity,
                    "bodega": raw_warehouse,
                    "ubicacion": raw_location,
                },
                original_bodega=detail.bodega,
            )
            if not form.is_valid():
                errors = " ".join(str(error) for field_errors in form.errors.values() for error in field_errors)
                messages.error(request, f"No se guardaron los cambios. {detail.codigo} en {detail.ubicacion}: {errors}")
                return redirect(return_to)
            new_values = (
                form.cleaned_data["bodega"],
                form.cleaned_data["ubicacion"],
                form.cleaned_data["cantidad"],
            )
            if new_values != (detail.bodega, detail.ubicacion, detail.cantidad):
                changes.append((detail, *new_values))

        changed_locations = {
            detail.pk
            for detail, warehouse, location, quantity in changes
            if (warehouse, location) != (detail.bodega, detail.ubicacion)
        }
        proposed_values = {
            detail.pk: (warehouse, location)
            for detail, warehouse, location, quantity in changes
        }
        selected_location_owners = {}
        duplicate_changed_detail = None
        for detail in details:
            warehouse, location = proposed_values.get(
                detail.pk, (detail.bodega, detail.ubicacion)
            )
            location_key = (
                detail.codigo.casefold(),
                warehouse.casefold(),
                location.casefold(),
            )
            existing_owner = selected_location_owners.get(location_key)
            if existing_owner is not None and (
                detail.pk in changed_locations or existing_owner.pk in changed_locations
            ):
                duplicate_changed_detail = detail
                break
            selected_location_owners[location_key] = detail

        if duplicate_changed_detail is not None:
            warehouse, location = proposed_values.get(
                duplicate_changed_detail.pk,
                (duplicate_changed_detail.bodega, duplicate_changed_detail.ubicacion),
            )
            messages.error(
                request,
                f"No se guardaron los cambios. {duplicate_changed_detail.codigo} ya existe en {warehouse} / {location}.",
            )
            return redirect(return_to)

        for detail, warehouse, location, quantity in changes:
            if detail.pk not in changed_locations:
                continue
            if latest_load.detalles.filter(
                codigo__iexact=detail.codigo,
                bodega__iexact=warehouse,
                ubicacion__iexact=location,
            ).exclude(pk__in=ids).exists():
                messages.error(
                    request,
                    f"No se guardaron los cambios. {detail.codigo} ya existe en {warehouse} / {location}.",
                )
                return redirect(return_to)

        warehouse_changed = False
        for detail, warehouse, location, quantity in changes:
            AjusteCantidadConteoInventario.objects.create(
                detalle=detail,
                cantidad_anterior=detail.cantidad,
                cantidad_nueva=quantity,
                bodega_anterior=detail.bodega,
                bodega_nueva=warehouse,
                ubicacion_anterior=detail.ubicacion,
                ubicacion_nueva=location,
                modificado_por=request.user,
            )
            warehouse_changed = warehouse_changed or warehouse != detail.bodega
            detail.bodega = warehouse
            detail.ubicacion = location
            detail.cantidad = quantity
            detail.save(update_fields=["bodega", "ubicacion", "cantidad"])

        if warehouse_changed:
            latest_load.total_bodegas = latest_load.detalles.values("bodega").distinct().count()
            latest_load.save(update_fields=["total_bodegas"])

    if changes:
        messages.success(request, f"Se guardaron {len(changes)} registro(s) del conteo vigente.")
    else:
        messages.info(request, "No había cambios nuevos para guardar.")
    return redirect(return_to)


@login_required
def inventory_count_detail_delete_view(request, pk):
    if request.method != "POST":
        raise Http404

    index_url = reverse("archivos:inventory_count_index")
    return_to = request.POST.get("return_to", "")
    if not (
        url_has_allowed_host_and_scheme(return_to, allowed_hosts={request.get_host()})
        and urlsplit(return_to).path == index_url
    ):
        return_to = index_url

    with transaction.atomic():
        detail = get_object_or_404(
            DetalleConteoInventario.objects.select_for_update(), pk=pk
        )
        latest_load = CargaConteoInventario.objects.select_for_update().first()
        if not latest_load or detail.carga_id != latest_load.pk:
            raise PermissionDenied("Solo se puede modificar el conteo vigente.")

        try:
            original_quantity = Decimal(request.POST.get(f"original_{pk}", ""))
        except InvalidOperation:
            original_quantity = None
        if original_quantity != detail.cantidad:
            messages.error(request, "La fila cambió desde que abriste la página. Actualiza antes de eliminar.")
            return redirect(return_to)

        code, location = detail.codigo, detail.ubicacion
        detail.delete()
        latest_load.total_registros = latest_load.detalles.count()
        latest_load.total_codigos = latest_load.detalles.values("codigo").distinct().count()
        latest_load.total_bodegas = latest_load.detalles.values("bodega").distinct().count()
        latest_load.save(update_fields=("total_registros", "total_codigos", "total_bodegas"))

    messages.success(request, f"Se eliminó {code} en {location} del conteo vigente.")
    return redirect(return_to)


@login_required
def inventory_count_quantity_view(request, pk):
    if request.method != "POST":
        raise Http404

    index_url = reverse("archivos:inventory_count_index")
    return_to = request.POST.get("return_to", "")
    if not (
        url_has_allowed_host_and_scheme(return_to, allowed_hosts={request.get_host()})
        and urlsplit(return_to).path == index_url
    ):
        return_to = index_url

    form = InventoryCountQuantityForm(request.POST)
    if not form.is_valid():
        messages.error(request, "No se guardó la cantidad: " + " ".join(
            str(error) for errors in form.errors.values() for error in errors
        ))
        return redirect(return_to)

    with transaction.atomic():
        detail = get_object_or_404(
            DetalleConteoInventario.objects.select_for_update(), pk=pk
        )
        latest_load = CargaConteoInventario.objects.first()
        if not latest_load or detail.carga_id != latest_load.pk:
            raise PermissionDenied("Solo se puede modificar el conteo vigente.")

        original = request.POST.get("cantidad_original", "")
        try:
            original_quantity = Decimal(original)
        except InvalidOperation:
            original_quantity = None
        if original_quantity != detail.cantidad:
            messages.error(request, "La cantidad cambió desde que abriste la página. Actualiza antes de guardar.")
            return redirect(return_to)

        new_quantity = form.cleaned_data["cantidad"]
        if new_quantity != detail.cantidad:
            AjusteCantidadConteoInventario.objects.create(
                detalle=detail,
                cantidad_anterior=detail.cantidad,
                cantidad_nueva=new_quantity,
                modificado_por=request.user,
            )
            detail.cantidad = new_quantity
            detail.save(update_fields=["cantidad"])
            messages.success(request, f"Cantidad de {detail.codigo} en {detail.ubicacion} actualizada a {new_quantity}.")
        else:
            messages.info(request, "La cantidad ya tenía ese valor.")
    return redirect(return_to)


@admin_required
def inventory_count_upload_view(request):
    form = InventoryCountUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            result = import_inventory_count_workbook(
                form.cleaned_data["archivo"],
                request.user,
            )
        except ExcelImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(
                request,
                f"Conteo cargado: {result.registros} registros conservados.",
            )
            return redirect("archivos:inventory_count_index")
    return render(request, "archivos/inventory_count_upload.html", {"form": form})


@admin_required
def inventory_count_delete_view(request, pk):
    if request.method != "POST":
        raise Http404
    load = get_object_or_404(CargaConteoInventario, pk=pk)
    file_name = load.archivo_nombre
    record_count = load.total_registros
    load.delete()
    messages.success(
        request,
        f"Conteo {file_name} eliminado junto con sus {record_count} registros.",
    )
    return redirect("archivos:inventory_count_index")


@admin_required
def inventory_count_template_view(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conteo"
    headers = ("Bodega", "Ubicacion", "Codigo", "Cantidad")
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:D1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip(("A", "B", "C", "D"), (24, 24, 24, 16)):
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_conteo_inventario.xlsx"'
    return response


@login_required
def be_count_index_view(request):
    query = request.GET.get("q", "").strip()
    warehouse = request.GET.get("bodega", "").strip()
    current_load = CargaConteoBE.objects.select_related("cargado_por").first()
    details = current_load.detalles.all() if current_load else DetalleConteoBE.objects.none()
    if query:
        details = details.filter(codigo__icontains=query)
    if warehouse in ("AGDOSA", "Lourdes", "Miramar"):
        details = details.filter(bodega=warehouse)
    return render(request, "archivos/be_count_index.html", {
        "query": query,
        "selected_warehouse": warehouse,
        "current_load": current_load,
        "details": Paginator(details, 50).get_page(request.GET.get("page")),
        "warehouses": ("AGDOSA", "Lourdes", "Miramar"),
    })


@admin_required
def be_count_upload_view(request):
    form = BECountUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            result = import_be_count_workbook(form.cleaned_data["archivo"], request.user)
        except ExcelImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(request, f"Conteo BE actualizado: {result.registros} relaciones de bodega y código vigentes.")
            return redirect("archivos:be_count_index")
    return render(request, "archivos/be_count_upload.html", {
        "form": form,
        "current_load": CargaConteoBE.objects.first(),
    })


@admin_required
def be_count_template_view(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conteo BE"
    sheet.append(("Bodega", "Codigo", "Cantidad"))
    sheet.append(("AGDOSA", "ABC-001", 5))
    sheet.append(("Lourdes", "ABC-001", 2))
    sheet.append(("Miramar", "ABC-002", 1))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:C1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    for column, width in zip(("A", "B", "C"), (22, 26, 16)):
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_conteo_be.xlsx"'
    return response


@login_required
def rge_index_view(request):
    query = request.GET.get("q", "").strip()
    current_load = CargaRGE.objects.select_related("cargado_por").first()
    details = current_load.detalles.all() if current_load else DetalleRGE.objects.none()
    if current_load and query:
        details = details.filter(
            Q(codigo__icontains=query) | Q(ubicacion__icontains=query)
        )
    return render(
        request,
        "archivos/rge_index.html",
        {
            "query": query,
            "current_load": current_load,
            "details": Paginator(details, 50).get_page(request.GET.get("page")),
        },
    )


@admin_required
def rge_upload_view(request):
    form = RGEUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            result = import_rge_workbook(form.cleaned_data["archivo"], request.user)
        except ExcelImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(
                request,
                f"RGE actualizado: {result.registros} relaciones de código y ubicación vigentes.",
            )
            return redirect("archivos:rge_index")
    return render(
        request,
        "archivos/rge_upload.html",
        {"form": form, "current_load": CargaRGE.objects.first()},
    )


@admin_required
def rge_template_view(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RGE"
    sheet.append(("Codigo", "Ubicacion", "Cantidad"))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:C1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 16

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_rge.xlsx"'
    return response


@login_required
def cost_base_index_view(request):
    query = request.GET.get("q", "").strip()
    current_load = CargaBaseCosto.objects.select_related("cargado_por").first()
    costs = current_load.costos.all() if current_load else CostoProducto.objects.none()
    if query:
        costs = costs.filter(codigo__icontains=query)
    return render(
        request,
        "archivos/cost_base_index.html",
        {
            "query": query,
            "current_load": current_load,
            "costs": Paginator(costs, 50).get_page(request.GET.get("page")),
        },
    )


@admin_required
def cost_base_upload_view(request):
    form = CostBaseUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            result = import_cost_base_workbook(form.cleaned_data["archivo"], request.user)
        except ExcelImportError as error:
            form.add_error("archivo", str(error))
        else:
            messages.success(
                request,
                f"Base de costo actualizada: {result.registros} códigos vigentes.",
            )
            return redirect("archivos:cost_base_index")
    return render(
        request,
        "archivos/cost_base_upload.html",
        {"form": form, "current_load": CargaBaseCosto.objects.first()},
    )


@admin_required
def cost_base_template_view(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Costos"
    sheet.append(("Codigo", "MCOSUC"))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:B1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F76")
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 18
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_base_costo.xlsx"'
    return response
