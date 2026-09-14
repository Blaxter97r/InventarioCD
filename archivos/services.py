from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile
import re
import unicodedata

from django.db import IntegrityError, transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import (
    CargaBaseCosto,
    CargaConteoBE,
    CargaConteoInventario,
    CargaRGE,
    CostoProducto,
    DetalleConteoBE,
    DetalleConteoInventario,
    DetalleRGE,
    Linea,
    Sublinea,
    Ubicacion,
)


LINEA_HEADERS = ("id_linea", "nlinea", "descripcion")
SUBLINEA_HEADERS = ("id_sublinea", "id_linea", "numero", "descripcion")
UBICACION_HEADERS = ("ubicacion", "bodega", "sector", "tipo")
INVENTORY_COUNT_HEADERS = ("bodega", "ubicacion", "codigo", "cantidad")
BE_COUNT_HEADERS = ("bodega", "codigo", "cantidad")
BE_WAREHOUSES = {"AGDOSA": "AGDOSA", "LOURDES": "Lourdes", "MIRAMAR": "Miramar"}
RGE_HEADERS = ("codigo", "ubicacion")
RGE_QUANTITY_HEADERS = ("cantidad", "existencia", "saldo", "cantidad_rge", "existencia_rge")
COST_BASE_HEADERS = ("codigo", "mcosuc")


class ExcelImportError(Exception):
    pass


@dataclass(frozen=True)
class ImportResult:
    lineas_creadas: int = 0
    lineas_actualizadas: int = 0
    sublineas_creadas: int = 0
    sublineas_actualizadas: int = 0


@dataclass(frozen=True)
class LocationImportResult:
    registros: int = 0
    reemplazadas: int = 0


@dataclass(frozen=True)
class InventoryCountImportResult:
    carga: CargaConteoInventario
    registros: int = 0


@dataclass(frozen=True)
class BECountImportResult:
    carga: CargaConteoBE
    registros: int = 0


@dataclass(frozen=True)
class RGEImportResult:
    carga: CargaRGE
    registros: int = 0


@dataclass(frozen=True)
class CostBaseImportResult:
    carga: CargaBaseCosto
    registros: int = 0


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"_x(?:000a|000d)_", " ", text, flags=re.IGNORECASE)
    return re.sub(r"[^a-z0-9]+", "_", text.casefold()).strip("_")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_sheet(workbook, expected_name):
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().casefold() == expected_name.casefold():
            return workbook[sheet_name]
    raise ExcelImportError(f"Falta la hoja obligatoria “{expected_name}”.")


def _read_rows(sheet, required_headers):
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header_positions = {
        _normalize_header(value): position
        for position, value in enumerate(first_row)
        if _normalize_header(value)
    }
    missing = [header for header in required_headers if header not in header_positions]
    if missing:
        raise ExcelImportError(
            f"La hoja “{sheet.title}” no contiene las columnas: {', '.join(missing)}."
        )

    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        record = {
            header: _cell_text(row[position] if position < len(row) else None)
            for header, position in header_positions.items()
            if header in required_headers
        }
        if not any(record.values()):
            continue
        empty_fields = [header for header in required_headers if not record.get(header)]
        if empty_fields:
            raise ExcelImportError(
                f"Hoja “{sheet.title}”, fila {row_number}: faltan valores en "
                f"{', '.join(empty_fields)}."
            )
        record["_row"] = row_number
        records.append(record)
    return records


def _validate_records(lineas, sublineas):
    line_ids = set()
    for record in lineas:
        line_id = record["id_linea"]
        if line_id in line_ids:
            raise ExcelImportError(f"El id_linea {line_id} está repetido en la hoja Lineas.")
        line_ids.add(line_id)
        if len(line_id) > 30 or len(record["nlinea"]) > 30:
            raise ExcelImportError(f"Hoja Lineas, fila {record['_row']}: un código supera 30 caracteres.")
        if len(record["descripcion"]) > 255:
            raise ExcelImportError(f"Hoja Lineas, fila {record['_row']}: la descripción supera 255 caracteres.")

    available_line_ids = line_ids | set(Linea.objects.values_list("id_linea", flat=True))
    subline_ids = set()
    line_number_pairs = set()
    for record in sublineas:
        subline_id = record["id_sublinea"]
        pair = (record["id_linea"], record["numero"])
        if subline_id in subline_ids:
            raise ExcelImportError(f"El id_sublinea {subline_id} está repetido en la hoja Sublineas.")
        if pair in line_number_pairs:
            raise ExcelImportError(
                f"La sublínea número {record['numero']} está repetida para la línea {record['id_linea']}."
            )
        if record["id_linea"] not in available_line_ids:
            raise ExcelImportError(
                f"Hoja Sublineas, fila {record['_row']}: la línea {record['id_linea']} no existe."
            )
        if len(subline_id) > 30 or len(record["id_linea"]) > 30 or len(record["numero"]) > 30:
            raise ExcelImportError(f"Hoja Sublineas, fila {record['_row']}: un código supera 30 caracteres.")
        if len(record["descripcion"]) > 255:
            raise ExcelImportError(f"Hoja Sublineas, fila {record['_row']}: la descripción supera 255 caracteres.")
        subline_ids.add(subline_id)
        line_number_pairs.add(pair)


def import_catalog_workbook(uploaded_file):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ExcelImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        lineas = _read_rows(_find_sheet(workbook, "Lineas"), LINEA_HEADERS)
        sublineas = _read_rows(_find_sheet(workbook, "Sublineas"), SUBLINEA_HEADERS)
        if not lineas and not sublineas:
            raise ExcelImportError("El archivo no contiene registros para importar.")
        _validate_records(lineas, sublineas)

        counters = {"lineas_creadas": 0, "lineas_actualizadas": 0, "sublineas_creadas": 0, "sublineas_actualizadas": 0}
        try:
            with transaction.atomic():
                for record in lineas:
                    _, created = Linea.objects.update_or_create(
                        id_linea=record["id_linea"],
                        defaults={"nlinea": record["nlinea"], "descripcion": record["descripcion"]},
                    )
                    counters["lineas_creadas" if created else "lineas_actualizadas"] += 1

                for record in sublineas:
                    _, created = Sublinea.objects.update_or_create(
                        id_sublinea=record["id_sublinea"],
                        defaults={
                            "linea_id": record["id_linea"],
                            "numero": record["numero"],
                            "descripcion": record["descripcion"],
                        },
                    )
                    counters["sublineas_creadas" if created else "sublineas_actualizadas"] += 1
        except IntegrityError as error:
            raise ExcelImportError(
                "La carga produciría códigos duplicados. Revisa los números de sublínea dentro de cada línea."
            ) from error
        return ImportResult(**counters)
    finally:
        workbook.close()


def import_location_workbook(uploaded_file):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ExcelImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        records = _read_rows(_find_sheet(workbook, "Ubicaciones"), UBICACION_HEADERS)
        if not records:
            raise ExcelImportError("La hoja Ubicaciones no contiene registros para importar.")

        location_ids = set()
        for record in records:
            location_id = record["ubicacion"]
            if location_id in location_ids:
                raise ExcelImportError(f"La ubicación {location_id} está repetida en el archivo.")
            if len(location_id) > 50:
                raise ExcelImportError(
                    f"Hoja Ubicaciones, fila {record['_row']}: ubicación supera 50 caracteres."
                )
            for field in ("bodega", "sector", "tipo"):
                if len(record[field]) > 100:
                    raise ExcelImportError(
                        f"Hoja Ubicaciones, fila {record['_row']}: {field} supera 100 caracteres."
                    )
            location_ids.add(location_id)

        with transaction.atomic():
            replaced_count = Ubicacion.objects.count()
            Ubicacion.objects.all().delete()
            Ubicacion.objects.bulk_create(
                [
                    Ubicacion(
                        ubicacion=record["ubicacion"],
                        bodega=record["bodega"],
                        sector=record["sector"],
                        tipo=record["tipo"],
                    )
                    for record in records
                ],
                batch_size=2000,
            )
        return LocationImportResult(registros=len(records), reemplazadas=replaced_count)
    finally:
        workbook.close()


def _find_inventory_count_header(workbook):
    best_sheet = workbook.active
    best_positions = {}
    best_row_number = 1
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=2, values_only=True),
            start=1,
        ):
            positions = {
                _normalize_header(value): position
                for position, value in enumerate(row)
                if _normalize_header(value)
            }
            matches = sum(header in positions for header in INVENTORY_COUNT_HEADERS)
            best_matches = sum(header in best_positions for header in INVENTORY_COUNT_HEADERS)
            if matches > best_matches:
                best_sheet = sheet
                best_positions = positions
                best_row_number = row_number
            if matches == len(INVENTORY_COUNT_HEADERS):
                return sheet, positions, row_number
    missing = [header for header in INVENTORY_COUNT_HEADERS if header not in best_positions]
    raise ExcelImportError(
        "Faltan las columnas obligatorias: "
        + ", ".join(header.replace("_", " ") for header in missing)
        + "."
    )


def _find_be_count_header(workbook):
    best_positions = {}
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1
        ):
            positions = {
                _normalize_header(value): position
                for position, value in enumerate(row)
                if _normalize_header(value)
            }
            if sum(header in positions for header in BE_COUNT_HEADERS) > sum(
                header in best_positions for header in BE_COUNT_HEADERS
            ):
                best_positions = positions
            if all(header in positions for header in BE_COUNT_HEADERS):
                return sheet, positions, row_number
    missing = [header for header in BE_COUNT_HEADERS if header not in best_positions]
    raise ExcelImportError("Faltan las columnas obligatorias: " + ", ".join(missing) + ".")


def _find_rge_header(workbook):
    best_sheet = workbook.active
    best_positions = {}
    best_row_number = 1
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=20, values_only=True),
            start=1,
        ):
            positions = {
                _normalize_header(value): position
                for position, value in enumerate(row)
                if _normalize_header(value)
            }
            matches = sum(header in positions for header in RGE_HEADERS)
            best_matches = sum(header in best_positions for header in RGE_HEADERS)
            if matches > best_matches:
                best_sheet = sheet
                best_positions = positions
                best_row_number = row_number
            if matches == len(RGE_HEADERS):
                return sheet, positions, row_number
    missing = [header for header in RGE_HEADERS if header not in best_positions]
    raise ExcelImportError(
        "Faltan las columnas obligatorias: "
        + ", ".join(header.replace("_", " ") for header in missing)
        + "."
    )


def _find_cost_base_header(workbook):
    best_sheet = workbook.active
    best_positions = {}
    best_row_number = 1
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=20, values_only=True),
            start=1,
        ):
            positions = {
                _normalize_header(value): position
                for position, value in enumerate(row)
                if _normalize_header(value)
            }
            matches = sum(header in positions for header in COST_BASE_HEADERS)
            best_matches = sum(header in best_positions for header in COST_BASE_HEADERS)
            if matches > best_matches:
                best_sheet = sheet
                best_positions = positions
                best_row_number = row_number
            if matches == len(COST_BASE_HEADERS):
                return sheet, positions, row_number
    missing = [header for header in COST_BASE_HEADERS if header not in best_positions]
    raise ExcelImportError(
        "Faltan las columnas obligatorias: "
        + ", ".join(missing)
        + "."
    )


def _inventory_quantity(value, row_number):
    cleaned = str(value if value is not None else "").strip().replace(" ", "")
    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    try:
        quantity = Decimal(cleaned)
    except InvalidOperation as error:
        raise ExcelImportError(
            f"Fila {row_number}: Cantidad debe contener un número válido."
        ) from error
    if abs(quantity) >= Decimal("100000000000000"):
        raise ExcelImportError(f"Fila {row_number}: Cantidad excede el límite permitido.")
    return quantity


def _cost_value(value, row_number):
    if isinstance(value, Decimal):
        cost = value
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cost = Decimal(str(value))
    else:
        cleaned = str(value if value is not None else "").strip().replace(" ", "")
        if "," in cleaned and "." in cleaned:
            if cleaned.rfind(",") > cleaned.rfind("."):
                cleaned = cleaned.replace(".", "").replace(",", ".")
            else:
                cleaned = cleaned.replace(",", "")
        elif "," in cleaned:
            cleaned = cleaned.replace(",", ".")
        try:
            cost = Decimal(cleaned)
        except InvalidOperation as error:
            raise ExcelImportError(
                f"Fila {row_number}: MCOSUC debe contener un número válido."
            ) from error
    if not cost.is_finite() or abs(cost) >= Decimal("100000000000000"):
        raise ExcelImportError(f"Fila {row_number}: MCOSUC excede el límite permitido.")
    if cost < 0:
        raise ExcelImportError(f"Fila {row_number}: MCOSUC no puede ser negativo.")
    return cost.quantize(Decimal("0.0001"))


def import_inventory_count_workbook(uploaded_file, user):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ExcelImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        sheet, positions, header_row_number = _find_inventory_count_header(workbook)
        records = []
        warehouses = set()
        codes = set()
        first_data_row = header_row_number + 1
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=first_data_row, values_only=True),
            start=first_data_row,
        ):
            values = {
                header: row[positions[header]] if positions[header] < len(row) else None
                for header in INVENTORY_COUNT_HEADERS
            }
            if not any(value not in (None, "") for value in values.values()):
                continue
            if len(records) >= 200000:
                raise ExcelImportError("El archivo supera el máximo de 200,000 registros.")

            bodega = _cell_text(values["bodega"])
            ubicacion = _cell_text(values["ubicacion"])
            codigo = _cell_text(values["codigo"])
            if not bodega or not ubicacion or not codigo or values["cantidad"] in (None, ""):
                raise ExcelImportError(
                    f"Fila {row_number}: Bodega, Ubicacion, Codigo y Cantidad son obligatorios."
                )
            if len(bodega) > 100 or len(ubicacion) > 100 or len(codigo) > 100:
                raise ExcelImportError(
                    f"Fila {row_number}: Bodega, Ubicacion o Codigo supera 100 caracteres."
                )
            records.append(
                DetalleConteoInventario(
                    bodega=bodega,
                    ubicacion=ubicacion,
                    codigo=codigo,
                    cantidad=_inventory_quantity(values["cantidad"], row_number),
                )
            )
            warehouses.add(bodega)
            codes.add(codigo)

        if not records:
            raise ExcelImportError("El archivo no contiene registros de conteo para importar.")

        with transaction.atomic():
            carga = CargaConteoInventario.objects.create(
                archivo_nombre=uploaded_file.name[:255],
                total_registros=len(records),
                total_codigos=len(codes),
                total_bodegas=len(warehouses),
                cargado_por=user,
            )
            for record in records:
                record.carga = carga
            DetalleConteoInventario.objects.bulk_create(records, batch_size=2000)
        return InventoryCountImportResult(carga=carga, registros=len(records))
    finally:
        workbook.close()


def import_be_count_workbook(uploaded_file, user):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ExcelImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        sheet, positions, header_row_number = _find_be_count_header(workbook)
        quantities = {}
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row_number + 1, values_only=True),
            start=header_row_number + 1,
        ):
            values = {
                header: row[positions[header]] if positions[header] < len(row) else None
                for header in BE_COUNT_HEADERS
            }
            if all(value in (None, "") for value in values.values()):
                continue
            raw_warehouse = _cell_text(values["bodega"])
            warehouse = BE_WAREHOUSES.get(raw_warehouse.upper())
            code = _cell_text(values["codigo"]).upper()
            if not warehouse:
                raise ExcelImportError(
                    f"Fila {row_number}: Bodega debe ser AGDOSA, Lourdes o Miramar."
                )
            if not code or values["cantidad"] in (None, ""):
                raise ExcelImportError(f"Fila {row_number}: Codigo y Cantidad son obligatorios.")
            if len(code) > 100:
                raise ExcelImportError(f"Fila {row_number}: Codigo supera 100 caracteres.")
            key = (warehouse, code)
            if key not in quantities and len(quantities) >= 200000:
                raise ExcelImportError("El archivo supera el máximo de 200,000 relaciones de bodega y código.")

            raw_quantity = _cell_text(values["cantidad"]).replace(" ", "").replace(",", ".")
            try:
                quantity = Decimal(raw_quantity)
            except InvalidOperation as error:
                raise ExcelImportError(f"Fila {row_number}: Cantidad debe ser un número válido.") from error
            if not quantity.is_finite() or quantity < 0 or quantity >= Decimal("100000000000000"):
                raise ExcelImportError(f"Fila {row_number}: Cantidad debe ser positiva o cero y estar dentro del límite permitido.")
            if quantity != quantity.quantize(Decimal("0.0001")):
                raise ExcelImportError(f"Fila {row_number}: Cantidad admite hasta cuatro decimales.")
            total = quantities.get(key, Decimal(0)) + quantity
            if total >= Decimal("100000000000000"):
                raise ExcelImportError(f"Fila {row_number}: el total de {code} en {warehouse} excede el límite permitido.")
            quantities[key] = total

        if not quantities:
            raise ExcelImportError("El archivo no contiene registros de conteo BE para importar.")

        with transaction.atomic():
            CargaConteoBE.objects.all().delete()
            carga = CargaConteoBE.objects.create(
                archivo_nombre=uploaded_file.name[:255],
                total_registros=len(quantities),
                total_codigos=len({code for _, code in quantities}),
                cargado_por=user,
            )
            DetalleConteoBE.objects.bulk_create(
                [
                    DetalleConteoBE(carga=carga, bodega=warehouse, codigo=code, cantidad=quantities[(warehouse, code)])
                    for warehouse, code in sorted(quantities)
                ],
                batch_size=2000,
            )
        return BECountImportResult(carga=carga, registros=len(quantities))
    finally:
        workbook.close()


def import_rge_workbook(uploaded_file, user):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ExcelImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        sheet, positions, header_row_number = _find_rge_header(workbook)
        records = {}
        quantity_position = next(
            (positions[header] for header in RGE_QUANTITY_HEADERS if header in positions),
            None,
        )
        first_data_row = header_row_number + 1
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=first_data_row, values_only=True),
            start=first_data_row,
        ):
            values = {
                header: row[positions[header]] if positions[header] < len(row) else None
                for header in RGE_HEADERS
            }
            if not any(value not in (None, "") for value in values.values()):
                continue
            if len(records) >= 300000:
                raise ExcelImportError("El archivo supera el máximo de 300,000 registros únicos.")

            codigo = _cell_text(values["codigo"]).upper()
            ubicacion = _cell_text(values["ubicacion"]).upper()
            if not codigo or not ubicacion:
                raise ExcelImportError(
                    f"Fila {row_number}: Código y Ubicación son obligatorios."
                )
            if len(codigo) > 100 or len(ubicacion) > 100:
                raise ExcelImportError(
                    f"Fila {row_number}: Código o Ubicación supera 100 caracteres."
                )
            raw_quantity = (
                row[quantity_position]
                if quantity_position is not None and quantity_position < len(row)
                else None
            )
            quantity = (
                _inventory_quantity(raw_quantity, row_number)
                if raw_quantity not in (None, "")
                else None
            )
            key = (codigo, ubicacion)
            previous = records.get(key)
            if previous is None:
                records[key] = quantity
            elif quantity is not None:
                records[key] = previous + quantity

        if not records:
            raise ExcelImportError("El archivo no contiene registros RGE para importar.")

        codes = {codigo for codigo, _ in records}
        locations = {location for _, location in records}
        with transaction.atomic():
            CargaRGE.objects.all().delete()
            carga = CargaRGE.objects.create(
                archivo_nombre=uploaded_file.name[:255],
                total_registros=len(records),
                total_codigos=len(codes),
                total_ubicaciones=len(locations),
                cargado_por=user,
            )
            DetalleRGE.objects.bulk_create(
                [
                    DetalleRGE(
                        carga=carga,
                        codigo=codigo,
                        ubicacion=ubicacion,
                        cantidad=records[(codigo, ubicacion)],
                    )
                    for codigo, ubicacion in sorted(records)
                ],
                batch_size=2000,
            )
        return RGEImportResult(carga=carga, registros=len(records))
    finally:
        workbook.close()


def import_cost_base_workbook(uploaded_file, user):
    try:
        uploaded_file.seek(0)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ExcelImportError("El archivo no es un libro de Excel .xlsx válido.") from error

    try:
        sheet, positions, header_row_number = _find_cost_base_header(workbook)
        costs = {}
        first_data_row = header_row_number + 1
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=first_data_row, values_only=True),
            start=first_data_row,
        ):
            raw_codigo = row[positions["codigo"]] if positions["codigo"] < len(row) else None
            raw_cost = row[positions["mcosuc"]] if positions["mcosuc"] < len(row) else None
            if raw_codigo in (None, "") and raw_cost in (None, ""):
                continue
            if len(costs) >= 300000:
                raise ExcelImportError("El archivo supera el máximo de 300,000 códigos únicos.")

            codigo = _cell_text(raw_codigo).upper()
            if not codigo or raw_cost in (None, ""):
                raise ExcelImportError(f"Fila {row_number}: Código y MCOSUC son obligatorios.")
            if len(codigo) > 100:
                raise ExcelImportError(f"Fila {row_number}: Código supera 100 caracteres.")
            cost = _cost_value(raw_cost, row_number)
            if codigo in costs and costs[codigo] != cost:
                raise ExcelImportError(
                    f"Fila {row_number}: el código {codigo} aparece con costos diferentes."
                )
            costs[codigo] = cost

        if not costs:
            raise ExcelImportError("El archivo no contiene costos para importar.")

        with transaction.atomic():
            CargaBaseCosto.objects.all().delete()
            carga = CargaBaseCosto.objects.create(
                archivo_nombre=uploaded_file.name[:255],
                total_registros=len(costs),
                cargado_por=user,
            )
            CostoProducto.objects.bulk_create(
                [
                    CostoProducto(carga=carga, codigo=codigo, costo=costo)
                    for codigo, costo in sorted(costs.items())
                ],
                batch_size=2000,
            )
        return CostBaseImportResult(carga=carga, registros=len(costs))
    finally:
        workbook.close()
